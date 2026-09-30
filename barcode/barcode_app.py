# ╔══════════════════════════════════════════════════════╗
# ║         바코드 라벨 생성기 - Streamlit                ║
# ╚══════════════════════════════════════════════════════╝
import os, io, re, urllib.request, csv
from collections import OrderedDict
from functools import lru_cache
from datetime import datetime, timedelta
import streamlit as st
import streamlit.components.v1 as components_v1
import time
from PIL import Image, ImageDraw, ImageFont
import barcode
from barcode.writer import ImageWriter
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, HRFlowable)
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfWriter, PdfReader
import pdfplumber
import pypdfium2 as pdfium
import label_sheet   # 폼텍 3102/3218 라벨지 PDF (배대지 부착용)
import kit_config    # 수강생 키트 설정 (../settings.json)
import coupang_po_change   # 쿠팡 발주서 변경(입고예정일·납품센터) 요청 엑셀

# ── 폰트 준비 ──────────────────────────────────────────
# 한글 폰트는 저장소 fonts/ 에 함께 넣어두고 그대로 쓴다.
#
# 예전에는 앱이 켜질 때마다 GitHub에서 폰트를 내려받았는데, 그 요청이
# 한 번이라도 실패하면(403 / 429 / 타임아웃) 스크립트가 첫 줄에서 죽어
# "Oh no. Error running app." 화면만 뜨고, 재시작해도 같은 자리에서
# 계속 죽어 앱이 영영 안 켜졌다. 서버가 재시작되면(재배포 / 절전 해제)
# 받아둔 폰트 파일도 함께 사라지므로 매번 이 위험을 다시 겪는 구조였다.
#
# 그래서 (1) 폰트를 저장소에 포함시켜 네트워크를 아예 안 타게 하고,
#        (2) 그래도 파일이 없으면 내려받되 실패해도 앱은 뜨게 했다.
_APP_DIR = (os.path.dirname(os.path.abspath(__file__))
            if '__file__' in globals() else os.getcwd())
_FONT_MIN_BYTES = 100_000          # HTML 에러 페이지가 폰트로 둔갑하는 것 방지
_FONT_FALLBACK_URLS = {            # 저장소에 파일이 없을 때만 쓰는 예비 경로
    'NanumGothicBold.ttf': 'https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/NanumGothic-Bold.ttf',
    'NanumGothic.ttf':     'https://raw.githubusercontent.com/google/fonts/main/ofl/nanumgothic/NanumGothic-Regular.ttf',
}


def _resolve_font(filename):
    """폰트 파일 경로를 돌려준다. 못 찾아도 예외를 올리지 않고 None을 준다."""
    for cand in (os.path.join(_APP_DIR, 'fonts', filename),
                 os.path.join(_APP_DIR, filename),
                 filename):
        try:
            if os.path.exists(cand) and os.path.getsize(cand) >= _FONT_MIN_BYTES:
                return cand
        except OSError:
            pass
    # 여기까지 왔다면 저장소에 폰트가 없는 상태 — 예비로 한 번만 받아본다.
    url = _FONT_FALLBACK_URLS.get(filename)
    if not url:
        return None
    target = os.path.join(_APP_DIR, filename)
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'barcode-app'})
        with urllib.request.urlopen(req, timeout=15) as resp:   # 타임아웃 필수
            data = resp.read()
        if len(data) < _FONT_MIN_BYTES:
            return None
        with open(target, 'wb') as f:
            f.write(data)
        return target
    except Exception:
        return None       # 못 받아도 앱은 계속 뜬다 (PDF만 영향)


_FONT_BOLD = _resolve_font('NanumGothicBold.ttf')
_FONT_REG = _resolve_font('NanumGothic.ttf')
# 한쪽만 있으면 서로 대신 쓴다 — 굵기만 손해, 기능은 그대로.
FONT_PATH = _FONT_BOLD or _FONT_REG
FONT_REG_PATH = _FONT_REG or _FONT_BOLD
FONT_READY = FONT_PATH is not None
FONT_WARNING = ''

# ── 중국어(간체) 폰트 ──────────────────────────────────
# 1688 소싱 상품의 옵션값은 '均码-黑色-左' 처럼 한자로 들어온다.
# 나눔고딕에는 한자 글리프가 없어서, 그대로 그리면 글자가 통째로 사라지고
# 괄호·숫자만 남는다(예: '米色圆8.5厘米【2片】' → '8.5 【2 】').
# 그래서 GB2312 상용한자만 추려낸 노토산스 SC 를 따로 싣고,
# 나눔고딕에 없는 글자는 이 폰트로 대신 그린다. (label_sheet 가 글자별로 고른다)
_FONT_CN_BOLD = _resolve_font('NotoSansSC-Bold.ttf')
_FONT_CN_REG = _resolve_font('NotoSansSC-Regular.ttf')
FONT_CN_READY = bool(_FONT_CN_BOLD or _FONT_CN_REG)

# reportlab 한국어 폰트 등록
try:
    if not FONT_READY:
        raise RuntimeError('한글 폰트 파일을 찾지 못했습니다')
    pdfmetrics.registerFont(TTFont('NanumBold', FONT_PATH))
    pdfmetrics.registerFont(TTFont('NanumReg', FONT_REG_PATH))
    # 중국어 폰트는 없어도 앱이 죽지 않게 따로 감싼다 (한자만 안 나올 뿐)
    try:
        if _FONT_CN_BOLD or _FONT_CN_REG:
            pdfmetrics.registerFont(TTFont('NotoSCBold', _FONT_CN_BOLD or _FONT_CN_REG))
            pdfmetrics.registerFont(TTFont('NotoSCReg', _FONT_CN_REG or _FONT_CN_BOLD))
    except Exception:
        FONT_CN_READY = False
except Exception as _font_err:
    FONT_READY = False
    FONT_WARNING = (
        '한글 폰트를 불러오지 못했습니다 — 라벨·PDF 만들기가 안 될 수 있습니다. '
        f'(fonts/NanumGothicBold.ttf 확인 필요 · {_font_err}) '
        '피킹/스캔 등 나머지 기능은 그대로 쓸 수 있습니다.'
    )


def _pil_font(size, path=None):
    """PIL 폰트 로더 — 폰트가 없어도 죽지 않고 기본 폰트로 대체한다."""
    p = path or FONT_PATH
    if p:
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            pass
    try:
        return ImageFont.load_default(size)
    except TypeError:      # 구버전 Pillow는 size 인자를 받지 않는다
        return ImageFont.load_default()

# ══════════════════════════════════════════════════════
# ── 공통 헬퍼 (바코드 라벨용) ──────────────────────────
# ══════════════════════════════════════════════════════
def wrap_text(text, font, max_w, draw):
    lines, cur = [], []
    for word in text.split(' '):
        test = ' '.join(cur + [word])
        bbox = draw.textbbox((0,0), test, font=font)
        if bbox[2]-bbox[0] <= max_w: cur.append(word)
        else:
            if cur: lines.append(' '.join(cur)); cur=[word]
            else:
                tmp=''
                for ch in word:
                    tb=draw.textbbox((0,0),tmp+ch,font=font)
                    if tb[2]-tb[0]<=max_w: tmp+=ch
                    else:
                        if tmp: lines.append(tmp)
                        tmp=ch
                if tmp: lines.append(tmp)
                cur=[]
    if cur: lines.append(' '.join(cur))
    return lines

def get_barcode_img(barcode_number, write_text=False):
    bc = barcode.get('code128', barcode_number, writer=ImageWriter())
    bc.writer.set_options({'module_height':80,'module_width':1.6,
                           'quiet_zone':6,'write_text':write_text})
    raw = bc.render()
    g=raw.convert('L'); w_r,h_r=g.size; p=g.load()
    l =next(x for x in range(w_r)          if any(p[x,yy]<255 for yy in range(h_r)))
    rr=next(x for x in range(w_r-1,-1,-1)   if any(p[x,yy]<255 for yy in range(h_r)))+1
    t =next(yy for yy in range(h_r)         if any(p[xx,yy]<255 for xx in range(w_r)))
    if write_text:
        bar_end=t
        for yy in range(t,h_r):
            if sum(1 for xx in range(l,rr) if p[xx,yy]<100)/(rr-l)>=0.25: bar_end=yy
        return raw.crop((l,t,rr,bar_end+2))
    else:
        b=next(yy for yy in range(h_r-1,-1,-1) if any(p[xx,yy]<255 for xx in range(w_r)))+1
        return raw.crop((l,t,rr,b))

# ── 소형 라벨 생성 ─────────────────────────────────────
def create_small(product_name, barcode_number, material, fixed_origin, fixed_age,
                 canvas_size=None):
    # canvas_size 를 주면 그 크기로 그린다 (라벨지 칸 비율에 딱 맞춰 뽑을 때).
    # 여백/글자 크기는 세로 길이에 비례해 줄여서, 작은 칸에서도 균형이 유지된다.
    CANVAS_W, CANVAS_H = canvas_size if canvas_size else (650, 450)
    sc = CANVAS_H / 450.0
    PAD = max(6, int(30 * sc))
    img=Image.new('RGB',(CANVAS_W,CANVAS_H),'white'); draw=ImageDraw.Draw(img)
    font_big=_pil_font(max(9, int(26 * sc)))
    font_mid=_pil_font(max(7, int(20 * sc)))

    y=PAD
    for line in wrap_text(product_name,font_big,CANVAS_W-PAD*2,draw)[:2]:
        bb=draw.textbbox((0,0),line,font=font_big)
        draw.text(((CANVAS_W-(bb[2]-bb[0]))//2,y),line,font=font_big,fill='black')
        y+=bb[3]-bb[1]+6

    bc_img=get_barcode_img(barcode_number,write_text=False)
    c1b=draw.textbbox((0,0),fixed_origin,font=font_mid)
    c2b=draw.textbbox((0,0),fixed_age,font=font_mid)
    mat_h=font_mid.size+8 if material else 0
    fixed_h=(c1b[3]-c1b[1])+(c2b[3]-c2b[1])+mat_h+18
    bc_y=y+10; cue_y=CANVAS_H-fixed_h-PAD
    BAR_W=CANVAS_W-PAD*2; BAR_H=cue_y-bc_y-6
    if BAR_H<int(60*sc): BAR_H=int(60*sc)
    img.paste(bc_img.resize((BAR_W,BAR_H),Image.LANCZOS),(PAD,bc_y))
    cur_y=bc_y+BAR_H+8

    if material:
        mt=f'재질 : {material}'; mb=draw.textbbox((0,0),mt,font=font_mid)
        draw.text(((CANVAS_W-(mb[2]-mb[0]))//2,cur_y),mt,font=font_mid,fill='black')
        cur_y+=mb[3]-mb[1]+6
    for txt in (fixed_origin,fixed_age):
        tb=draw.textbbox((0,0),txt,font=font_mid)
        draw.text(((CANVAS_W-(tb[2]-tb[0]))//2,cur_y),txt,font=font_mid,fill='black')
        cur_y+=tb[3]-tb[1]+5
    return img

# ── 대형 라벨 생성 ─────────────────────────────────────
def fit_font(text, max_w, draw, max_size=42, min_size=8):
    for size in range(max_size,min_size-1,-1):
        f=_pil_font(size)
        bb=draw.textbbox((0,0),text,font=f)
        if bb[2]-bb[0]<=max_w: return f
    return _pil_font(min_size)

def fit_wrapped_text(text, max_w, max_lines, draw, max_size=33, min_size=18):
    """max_lines 이하로 wrap 되는 가장 큰 폰트 크기 찾기. (긴 상품명 잘림 방지)"""
    for size in range(max_size, min_size - 1, -1):
        f = _pil_font(size)
        lines = wrap_text(text, f, max_w, draw)
        if len(lines) <= max_lines:
            return f, lines
    f = _pil_font(min_size)
    lines = wrap_text(text, f, max_w, draw)[:max_lines]
    return f, lines

def create_large(product_name, barcode_number, material, fix_list, canvas_size=None):
    # canvas_size 는 '회전 전' 크기 (가로, 세로). 마지막에 90도 돌리므로
    # 최종 결과는 (세로, 가로)가 된다. 라벨지 칸에 맞출 때 이 점을 주의.
    CANVAS_W,CANVAS_H = canvas_size if canvas_size else (450, 640)
    sc = CANVAS_W / 450.0
    PAD = max(5, int(22 * sc))
    img=Image.new('RGB',(CANVAS_W,CANVAS_H),'white'); draw=ImageDraw.Draw(img)
    # 자동 폰트 축소: 33pt부터 시도하다 4줄 이하로 들어오면 사용 (긴 상품명 전체 표시)
    fn, name_lines = fit_wrapped_text(product_name, CANVAS_W - PAD*2, 4, draw,
                                      max_size=max(10, int(33*sc)), min_size=max(6, int(18*sc)))
    fm=_pil_font(max(7, int(20 * sc)))
    ff=_pil_font(max(6, int(16 * sc)))

    y=PAD
    for line in name_lines:
        bb=draw.textbbox((0,0),line,font=fn)
        draw.text(((CANVAS_W-(bb[2]-bb[0]))//2,y),line,font=fn,fill='black')
        y+=bb[3]-bb[1]+4
    y+=8

    bc_img=get_barcode_img(barcode_number,write_text=False)
    fix_h=0
    for txt in fix_list:
        for ln in wrap_text(txt,ff,CANVAS_W-PAD*2,draw):
            bb=draw.textbbox((0,0),ln,font=ff); fix_h+=bb[3]-bb[1]+3
        fix_h+=6

    BAR_W=CANVAS_W-PAD*2
    BAR_H=CANVAS_H-y-int(30*sc)-int(14*sc)-fix_h-PAD
    if BAR_H<int(60*sc): BAR_H=int(60*sc)
    img.paste(bc_img.resize((BAR_W,BAR_H),Image.LANCZOS),(PAD,y)); y+=BAR_H+6

    if material:
        mt=f'재질 : {material}'; mb=draw.textbbox((0,0),mt,font=fm)
        draw.text(((CANVAS_W-(mb[2]-mb[0]))//2,y),mt,font=fm,fill='black')
        y+=mb[3]-mb[1]+8

    draw.line([(PAD,y),(CANVAS_W-PAD,y)],fill=(160,160,160),width=1); y+=10
    for txt in fix_list:
        for ln in wrap_text(txt,ff,CANVAS_W-PAD*2,draw):
            bb=draw.textbbox((0,0),ln,font=ff)
            draw.text((PAD,y),ln,font=ff,fill='black'); y+=bb[3]-bb[1]+3
        y+=6

    draw.rectangle([1,1,CANVAS_W-2,CANVAS_H-2],outline=(180,180,180),width=1)
    return img.rotate(90,expand=True)

# ── 엑셀 처리 공통 ─────────────────────────────────────
def process_excel(uploaded_file, mode, settings):
    wb=load_workbook(uploaded_file); ws=wb.active
    col_insert=settings['col_insert']
    ws.column_dimensions[get_column_letter(col_insert)].width=settings['col_width']

    temp_dir='_temp'; os.makedirs(temp_dir,exist_ok=True)
    ok=0; errors=[]
    progress=st.progress(0)
    status=st.empty()

    total=sum(1 for r in range(settings['start_row'],ws.max_row+1)
              if ws.cell(r,settings['col_barcode']).value)

    for r in range(settings['start_row'],ws.max_row+1):
        bv=ws.cell(r,settings['col_barcode']).value
        if not bv: continue
        bv=str(bv).strip()
        nm=str(ws.cell(r,settings['col_name']).value or '').strip()
        mt=str(ws.cell(r,settings['col_material']).value or '').strip()
        img_path=f'{temp_dir}/label_{r}.png'

        try:
            if mode=='소형':
                img=create_small(nm,bv,mt,settings['origin'],settings['age'])
            else:
                img=create_large(nm,bv,mt,settings['fix_list'])
            img.save(img_path)
        except Exception as e:
            errors.append(f'{r}행 실패: {e}'); continue

        xl=XLImage(img_path)
        xl.width=settings['insert_w']; xl.height=settings['insert_h']
        ws.add_image(xl,f'{get_column_letter(col_insert)}{r}')
        ws.row_dimensions[r].height=settings['row_height']
        ok+=1
        progress.progress(ok/max(total,1))
        status.text(f'✅ {r}행 처리 중... ({ok}/{total})')

    output=io.BytesIO(); wb.save(output); output.seek(0)
    progress.progress(1.0); status.text(f'🎉 완료! {ok}개 생성')
    return output, ok, errors


# ── 폼텍 라벨지 만들기 (배대지 부착용) — 소형/대형 탭 공용 UI ──────────
def render_label_sheet_section(prefix, form_key, mode, uploaded, base_cols,
                               fixed_lines, defaults, make_label_image):
    """엑셀 업로드 아래에 붙는 '배대지용 바코드 라벨지 만들기' 섹션.

    만들어지는 라벨지 구조 (상품 한 줄 기준):
        [ 상품사진 + 옵션1/옵션2 ]  ← 구분용, 1칸
        [ 바코드 라벨 ][ 바코드 라벨 ] ...  ← J열 수량만큼
    바코드 라벨은 위 '소형/대형 라벨 생성'이 만드는 것과 똑같은 디자인을 쓴다.

    prefix          : 위젯 key 접두어 — 소형/대형 탭이 key 충돌 없이 공존
    form_key        : '3102'(소형) 또는 '3218'(대형)
    mode            : '소형' / '대형' — 안내 문구용
    uploaded        : 그 탭의 file_uploader 결과 (같은 엑셀을 그대로 재사용)
    base_cols       : dict(col_name, col_barcode, col_material, start_row)
    fixed_lines     : 라벨 아래 고정 문구 (제조국·연령·취급주의 …)
    defaults        : dict(image, opt1, opt2, qty) 열 번호 기본값
    make_label_image: (item, px_w, px_h) → PIL 이미지. 칸 채우기를 켜면
                      px 크기가 들어오고, 끄면 None 이 들어와 원본 디자인 크기로 그린다.
    """
    form = label_sheet.FORMS[form_key]
    per_page = form['cols'] * form['rows']
    st.divider()
    st.subheader(f'🏷️ 배대지용 바코드 라벨지 만들기 — {form["title"]}')
    # 배포 확인용 — 이 줄이 안 보이거나 버전이 다르면 아직 예전 코드가 도는 것이다
    st.caption(f'🔖 라벨지 기능 버전 **{label_sheet.VERSION}**')
    st.caption(
        '상품마다 **새 줄에서 시작**해 **줄 맨 왼쪽에 상품사진+옵션**, '
        '**나머지 칸에 수량만큼 바코드 라벨**이 들어갑니다. '
        '수량이 남으면 다음 줄을 통째로 바코드로 이어 채웁니다. '
        f'바코드 라벨 디자인은 위 「{mode} 라벨 생성」과 똑같습니다. '
        '위에서 업로드한 엑셀을 그대로 쓰며, 인쇄할 때 배율 100%(실제 크기)로 출력하세요.'
    )

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        col_image = st.number_input('상품사진 열', min_value=1, max_value=50,
                                    value=defaults['image'], key=f'{prefix}_ls_img')
    with c2:
        col_opt1 = st.number_input('옵션1 열', min_value=0, max_value=50,
                                   value=defaults['opt1'], key=f'{prefix}_ls_opt1',
                                   help='0 이면 표시하지 않음. A=1, B=2 ... F=6, G=7')
    with c3:
        col_opt2 = st.number_input('옵션2 열', min_value=0, max_value=50,
                                   value=defaults['opt2'], key=f'{prefix}_ls_opt2',
                                   help='0 이면 표시하지 않음')
    with c4:
        col_qty = st.number_input('수량 열', min_value=0, max_value=50,
                                  value=defaults['qty'], key=f'{prefix}_ls_qty',
                                  help='0 이면 상품마다 1장')

    c5, c6, c7, c8 = st.columns(4)
    with c5:
        start_row = st.number_input('시작 줄 번호', min_value=1, max_value=form['rows'], value=1,
                                    key=f'{prefix}_ls_start',
                                    help=f'쓰다 남은 라벨지에 이어서 뽑을 때. 맨 윗줄이 1번, '
                                         f'한 장 {form["rows"]}줄 × {form["cols"]}칸')
    with c6:
        photo_cell = st.checkbox('상품사진 칸 넣기', value=True, key=f'{prefix}_ls_photo',
                                 help='켜면 줄 맨 왼쪽에 상품사진이 들어가고 나머지 칸이 바코드로 채워집니다. '
                                      '끄면 사진 없이 바코드 라벨만 빈틈없이 채웁니다.')
    with c7:
        stretch = st.checkbox('칸 꽉 채우기', value=False, key=f'{prefix}_ls_stretch',
                              help='켜면 라벨을 칸 비율로 다시 그려 여백을 없앱니다. '
                                   '끄면 기존 라벨 디자인 비율 그대로 칸 안에 맞춥니다.')
    with c8:
        border = st.checkbox('칸 테두리 인쇄', value=False, key=f'{prefix}_ls_border',
                             help='일반 용지에 시험 인쇄해 칸 위치를 맞춰볼 때만 켜세요')

    if st.button('🏷️ 바코드 라벨지 만들기', type='primary', key=f'{prefix}_ls_btn'):
        if not uploaded:
            st.warning('⚠️ 엑셀 파일을 먼저 업로드해주세요!')
            return
        if not FONT_READY:
            st.error(FONT_WARNING or '한글 폰트가 없어 라벨지를 만들 수 없습니다.')
            return

        settings = dict(base_cols)
        settings.update(col_image=col_image, col_opt1=col_opt1,
                        col_opt2=col_opt2, col_qty=col_qty)
        try:
            with st.spinner('엑셀에서 사진·옵션·수량을 읽는 중...'):
                items, warns = label_sheet.collect_label_items(uploaded.getvalue(), settings)
        except Exception as e:
            st.error(f'엑셀을 읽지 못했습니다: {e}')
            return
        if not items:
            st.warning('만들 라벨이 없습니다. 바코드 열/시작 행 설정을 확인해주세요.')
            return

        # ── 읽은 값 확인표 ───────────────────────────────
        # 열 번호를 잘못 넣으면 옵션이 통째로 비어 나온다. 눈으로 바로 확인하라고 보여준다.
        # 단, 옵션2는 원래 없는 상품이 많으므로(단일 옵션 상품) 비어 있어도 정상으로 본다.
        st.markdown('**📄 엑셀에서 읽은 값 (앞 10개)**')
        st.dataframe(
            [{'행': it.row, '상품명': it.name[:24], '옵션1': it.opt1, '옵션2': it.opt2,
              '수량': it.qty, '바코드': it.barcode,
              '사진': '있음' if it.image is not None else '없음'} for it in items[:10]],
            use_container_width=True, hide_index=True,
        )
        # 옵션1 열을 쓰기로 해놓고(0이 아님) 전 행이 비었다면 열 번호가 틀렸을 가능성이 크다.
        if col_opt1 and all(not it.opt1 for it in items):
            st.warning(f'옵션1 열({get_column_letter(col_opt1)}열)이 모든 행에서 비어 있습니다. '
                       '열 번호가 맞는지 확인해주세요. (A=1, B=2, C=3 … F=6, G=7) '
                       '옵션이 원래 없는 상품이면 그대로 두셔도 됩니다.')

        try:
            with st.spinner('바코드 라벨을 그리는 중...'):
                opts = dict(layout='separate', fixed_lines=fixed_lines, show_name=True,
                            show_material=True, border=border, stretch=stretch,
                            photo_cell=photo_cell)
                # 줄 번호 → 칸 번호 (맨 윗줄 1번 = 1번 칸)
                pdf_buf, total, pages = label_sheet.create_label_sheet_pdf(
                    items, form_key, opts,
                    start_slot=(int(start_row) - 1) * form['cols'] + 1,
                    label_image_fn=make_label_image)
        except Exception as e:
            st.error(f'라벨지 생성 실패: {e}')
            return

        label_cnt = sum(int(it.qty) for it in items)
        no_img = sum(1 for it in items if it.image is None)
        st.success(f'🎉 상품 {len(items)}종 · 바코드 라벨 {label_cnt}장'
                   f'{" + 상품사진 " + str(len(items)) + "칸" if photo_cell else ""}'
                   f' · 전체 {total}칸 · A4 {pages}페이지 (한 장 {per_page}칸)')
        if no_img:
            st.warning(f'상품사진을 못 찾은 상품이 {no_img}종 있습니다 — 그 칸은 "(이미지 없음)"으로 나옵니다.')
        for w in warns[:30]:
            st.caption(f'· {w}')
        if len(warns) > 30:
            st.caption(f'· … 외 {len(warns) - 30}건')

        base = re.sub(r'\.xlsx$', '', uploaded.name, flags=re.I)
        st.download_button('⬇️ 라벨지 PDF 다운로드', pdf_buf,
                           file_name=f'{base}_라벨지_{form_key}.pdf',
                           mime='application/pdf', key=f'{prefix}_ls_dl')
        try:
            preview = pdfium.PdfDocument(pdf_buf.getvalue())[0].render(scale=1.6).to_pil()
            st.image(preview, caption='1페이지 미리보기', use_container_width=True)
        except Exception:
            pass



# ══════════════════════════════════════════════════════
# ── 출고 작업 지시서 PDF 생성 함수들 ──────────────────
# ══════════════════════════════════════════════════════

def parse_date(date_str):
    """날짜 문자열을 파싱해서 datetime 반환. 실패하면 None."""
    if not date_str or not date_str.strip():
        return None
    s = date_str.strip()
    nums = re.findall(r'\d+', s)
    try:
        if len(nums) == 1 and len(nums[0]) == 8:
            ymd = nums[0]
            return datetime(int(ymd[:4]), int(ymd[4:6]), int(ymd[6:8]))
        elif len(nums) >= 3:
            y, m, d = int(nums[0]), int(nums[1]), int(nums[2])
            if y < 100: y += 2000
            return datetime(y, m, d)
    except Exception:
        pass
    return None


def calc_deadline(date_str):
    """입고예정일 + 20일 → 입고마감일 문자열 반환"""
    dt = parse_date(date_str)
    if dt is None:
        return '날짜 없음'
    deadline = dt + timedelta(days=20)
    return deadline.strftime('%Y-%m-%d')


# ── 부족(결품) 수량 계산 헬퍼 ───────────────────────────
# K열(박스번호)에는 '▲M7(1),부족(-1)' 처럼
#   · 실제로 배대지 박스에 담기는 양(▲M7(1))
#   · 물건이 모자라서 못 보내는 양(부족(-1))
# 이 한 칸에 같이 들어온다.
# 반면 H열(수량)은 '주문 수량'이라 부족분까지 포함된 숫자다.
# 그래서 작업자가 실제로 담아야 할 개수는 아래 공식으로 구한다.
#
#   최종(실출고) 수량 = 총 수량(H열) − 부족 수량
#
# 예) 수량 2 / 박스번호 '▲M7(1),부족(-1)'  →  총 2, 부족 1, 최종 1

# '부족(-1)', '국내부족(-2)', 개수 없는 '부족' 까지 인식
_RE_SHORTAGE_TOKEN = re.compile(r'^((?:국내)?부족)\s*(?:\(\s*(-?\d+)\s*\))?$')
# 토큰 끝의 '(숫자)' 추출 — '▲M7(1)', '국내재고(3)' 공통
_RE_TAIL_QTY = re.compile(r'\((-?\d+)\)\s*$')


def split_box_tokens(bn_raw):
    """박스번호 문자열을 토큰 리스트로 분리.
    '▲M7(1),부족(-1)' → ['▲M7(1)', '부족(-1)']"""
    s = str(bn_raw or '').strip()
    if not s or s.lower() == 'nan':
        return []
    return [t.strip() for t in re.split(r'[,;/]+', s) if t.strip()]


def _qty_int(item):
    """item['quantity'] 를 안전하게 int 로. 빈 값이나 숫자가 아닌 문자는 0."""
    try:
        return int(item.get('quantity', 0) or 0)
    except (ValueError, TypeError):
        return 0


def item_shortage_qty(item):
    """상품 1행의 '부족 수량'을 계산해 반환 (부족 없으면 0).

    - '부족(-1)' / '국내부족(-2)' → 괄호 안 숫자의 절댓값을 부족으로 합산
    - 개수가 안 적힌 '부족' → 총 수량에서 나머지 토큰 수량을 빼서 역산
    - 결과는 0 ~ 총 수량 사이로 잘라서 음수/과다 계산을 막는다
    """
    tokens = split_box_tokens(item.get('boxNumber', ''))
    if not tokens:
        return 0
    total = _qty_int(item)

    shortage = 0        # 개수가 명시된 부족 합계
    bare_shortage = False   # 개수 없이 '부족'만 적힌 토큰이 있었는지
    known_other = 0     # 부족이 아닌 토큰들의 수량 합 (실제 담기는 양)

    for tok in tokens:
        m = _RE_SHORTAGE_TOKEN.match(tok)
        if m:
            if m.group(2) is not None:
                shortage += abs(int(m.group(2)))
            else:
                bare_shortage = True
            continue
        m2 = _RE_TAIL_QTY.search(tok)
        if m2:
            known_other += abs(int(m2.group(1)))

    if bare_shortage:
        # 개수 미상 '부족' → 주문 수량 − 박스에 담기는 양 으로 추정
        shortage = max(shortage, total - known_other)

    if total > 0:
        return max(0, min(shortage, total))
    return max(0, shortage)


def item_box_quantities(item):
    """상품 1행 → {배대지 박스키: 담기는 수량} 딕셔너리.

    부족/국내재고/RAW 토큰은 박스에 담기지 않으므로 제외한다.
    '▲M7' 처럼 개수가 안 적힌 경우엔 (총 수량 − 부족 수량)을 담는 양으로 본다.
    """
    tokens = split_box_tokens(item.get('boxNumber', ''))
    if not tokens:
        return {}
    net_qty = max(0, _qty_int(item) - item_shortage_qty(item))

    result = {}
    pending = []   # 개수가 안 적힌 박스 토큰
    for tok in tokens:
        if _RE_SHORTAGE_TOKEN.match(tok):
            continue
        if any(kw in tok for kw in ('국내', '재고', 'RAW')):
            continue
        m = re.search(r'([A-Za-z]+\d+)', tok)
        if not m:
            m = re.search(r'(\d+)', tok.split('(')[0])
        if not m:
            continue
        key = m.group(1).upper()
        qm = _RE_TAIL_QTY.search(tok)
        if qm:
            result[key] = result.get(key, 0) + abs(int(qm.group(1)))
        else:
            pending.append(key)

    # 개수 미표기 박스는 남은 수량(= 최종 수량 − 이미 배분된 양)을 나눠 가진다.
    # 나누어떨어지지 않는 나머지는 앞 박스부터 1개씩 더 줘서 합계가 보존되게 한다
    # (예전에는 몫만 줘서 rest=5, 박스 2개면 2+2=4 로 1개가 사라졌다).
    if pending:
        rest = max(0, net_qty - sum(result.values()))
        share, extra = divmod(rest, len(pending))
        for i, key in enumerate(pending):
            result[key] = result.get(key, 0) + share + (1 if i < extra else 0)
    return result


def summarize_shortage(items):
    """items 전체의 (총 수량, 부족 수량, 최종 수량) 튜플 반환"""
    total = sum(_qty_int(i) for i in items)
    short = sum(item_shortage_qty(i) for i in items)
    short = min(short, total)
    return total, short, max(0, total - short)


def create_work_order_pdf(group_key, items, shipment_id=None, box_number=None):
    """reportlab으로 출고 작업 지시서 PDF 생성 → BytesIO 반환
    shipment_id/box_number가 있으면 상단 우측에 표시"""
    # ReportLab Paragraph는 XML 파서를 쓴다. 시트에서 온 문자열(물류센터명·상품명 등)에
    # '<b' '<br' '<font' 처럼 태그로 읽히는 조각이 섞이면 파싱이 깨져 그 송장의 PDF 생성이
    # 통째로 죽는다 ('&' 단독은 통과한다). 사용자 데이터가 들어가는 모든 자리는 이 함수로
    # 감싼다 — 상품 행은 원래 감싸고 있었고, 헤더 카드/그룹키만 빠져 있었다.
    from xml.sax.saxutils import escape as _xml_escape
    buf = io.BytesIO()
    PAGE_W, PAGE_H = A4
    MARGIN = 18 * mm

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN, bottomMargin=MARGIN
    )

    # 스타일 정의
    s_title   = ParagraphStyle('title',   fontName='NanumBold', fontSize=18, leading=22, textColor=colors.HexColor('#111111'))
    s_sub     = ParagraphStyle('sub',     fontName='NanumBold', fontSize=9,  leading=12, textColor=colors.HexColor('#888888'), spaceAfter=2)
    s_card_lbl= ParagraphStyle('cardlbl', fontName='NanumBold', fontSize=8,  leading=10, textColor=colors.HexColor('#888888'))
    s_card_val= ParagraphStyle('cardval', fontName='NanumBold', fontSize=12, leading=15, textColor=colors.HexColor('#111111'))
    s_card_big= ParagraphStyle('cardbig', fontName='NanumBold', fontSize=22, leading=26, textColor=colors.HexColor('#1a56db'))
    s_th      = ParagraphStyle('th',      fontName='NanumBold', fontSize=9,  leading=11, textColor=colors.white)
    s_td      = ParagraphStyle('td',      fontName='NanumReg',  fontSize=9,  leading=12, textColor=colors.HexColor('#111111'), wordWrap='CJK')
    s_td_bold = ParagraphStyle('tdbold',  fontName='NanumBold', fontSize=10, leading=12, textColor=colors.HexColor('#1a56db'))
    s_mono    = ParagraphStyle('mono',    fontName='NanumReg',  fontSize=9,  leading=12, textColor=colors.HexColor('#111111'))
    s_footer  = ParagraphStyle('footer',  fontName='NanumReg',  fontSize=8,  leading=10, textColor=colors.HexColor('#888888'))
    s_shipment= ParagraphStyle('shipment',fontName='NanumBold', fontSize=14, leading=18, textColor=colors.HexColor('#1a56db'), alignment=2)
    s_boxnum_huge = ParagraphStyle('boxnumhuge', fontName='NanumBold', fontSize=48, leading=54, textColor=colors.HexColor('#dc2626'), alignment=0)
    # 부족/최종 수량 표기용 스타일
    s_card_short  = ParagraphStyle('cardshort', fontName='NanumBold', fontSize=13, leading=16, textColor=colors.HexColor('#dc2626'))
    s_td_short    = ParagraphStyle('tdshort',   fontName='NanumBold', fontSize=7,  leading=9,  textColor=colors.HexColor('#dc2626'), alignment=1)

    # 총 수량(주문) / 부족 / 최종(실제 담을 수량)
    total_qty, shortage_qty, final_qty = summarize_shortage(items)
    first       = items[0]
    deadline    = calc_deadline(first.get('expectedDate',''))
    created_at  = datetime.now().strftime('%Y-%m-%d %H:%M')
    usable_w    = PAGE_W - MARGIN * 2

    story = []

    # ── 송장번호 바코드 (피킹검증 스캔용) ───────────────
    # CSV에 있는 송장번호(shipmentNumber)를 바코드로 추가
    invoice_for_barcode = first.get('shipmentNumber', '') or ''
    barcode_flowable = None
    if invoice_for_barcode:
        try:
            # 벡터 기반 Code128 (PNG보다 훨씬 선명)
            from reportlab.graphics.barcode.code128 import Code128
            from reportlab.graphics.shapes import Drawing
            from reportlab.graphics import renderPDF
            from reportlab.platypus import Flowable

            class _BarcodeFlowable(Flowable):
                def __init__(self, value, width_mm=70, height_mm=14):
                    Flowable.__init__(self)
                    self.value = str(value)
                    self.width = width_mm * mm
                    self.height = height_mm * mm
                    self.hAlign = 'RIGHT'

                def draw(self):
                    # barWidth를 원하는 폭에 맞춰 계산 (여백 포함)
                    # Code128은 가변 길이라 렌더 후 스케일링
                    bc = Code128(self.value, barHeight=self.height, humanReadable=False)
                    bc_w = bc.width
                    scale = self.width / bc_w if bc_w > 0 else 1
                    self.canv.saveState()
                    self.canv.scale(scale, 1)
                    bc.drawOn(self.canv, 0, 0)
                    self.canv.restoreState()

            barcode_flowable = _BarcodeFlowable(invoice_for_barcode, width_mm=70, height_mm=14)
        except Exception:
            # Fallback: 기존 PNG 방식
            try:
                from reportlab.platypus import Image as RLImage
                bc_img = get_barcode_img(str(invoice_for_barcode), write_text=False)
                bc_buf = io.BytesIO()
                bc_img.save(bc_buf, format='PNG')
                bc_buf.seek(0)
                barcode_flowable = RLImage(bc_buf, width=70*mm, height=14*mm)
                barcode_flowable.hAlign = 'RIGHT'
            except Exception:
                barcode_flowable = None

    # ── 상단 큰 박스번호 표시 ──
    # 호출 시 box_number 인자로 전달된 값을 사용 (자동 부여된 송장별 박스번호)
    # 송장 전체가 국내재고인 경우 box_number=None이 전달되어 표시되지 않음
    # 배대지 박스별 수량 요약도 함께 표시 (예: "58번 M2(10),W11(93)")
    # 배대지 박스별 수량은 item_box_quantities()로 토큰 단위 집계한다.
    # (과거에는 문자열에 '부족'이 한 글자라도 있으면 그 행을 통째로 제외해서
    #  '▲M7(1),부족(-1)' 같은 혼합 행이 요약에서 사라졌고, 수량도 부족분이
    #  포함된 H열 값을 그대로 더해 실제 담을 개수보다 많게 나왔다.)
    dapae_summary_str = ''
    if box_number:
        from collections import Counter as _Counter
        dapae_counts = _Counter()  # {배대지박스키: 실제 담을 수량합}
        for it in items:
            for bkey, bqty in item_box_quantities(it).items():
                dapae_counts[bkey] += bqty
        dapae_counts = _Counter({k: v for k, v in dapae_counts.items() if v > 0})
        if dapae_counts:
            parts = [f'{k}({v})' for k, v in sorted(dapae_counts.items())]
            # '1번'은 48pt, 배대지 구성은 18pt(타이틀 크기)로 축소
            dapae_summary_str = (
                f' <font size="18" color="#111111">'
                f'{",".join(parts)}</font>'
            )
    big_box_label = f'{box_number}번{dapae_summary_str}' if box_number else ''

    # ── 헤더: 좌측 타이틀 + 우측 쉽먼트(+바코드) ──────────
    # 박스번호는 좌측 상단에 이미 표시되므로 우측에는 쉽먼트 ID만
    if shipment_id:
        shipment_text = f'쉽먼트 {shipment_id}'
    else:
        shipment_text = ''

    if shipment_text or barcode_flowable:
        right_cell = []
        if barcode_flowable:
            right_cell.append(barcode_flowable)
            right_cell.append(Spacer(1, 1*mm))
        if shipment_text:
            right_cell.append(Paragraph(shipment_text, s_shipment))
        right_tbl = Table([[c] for c in right_cell], colWidths=[usable_w * 0.5])
        right_tbl.setStyle(TableStyle([
            ('ALIGN', (0,0), (-1,-1), 'RIGHT'),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ]))
        # 좌측: 큰 박스번호 + 타이틀
        if big_box_label:
            left_cell = [
                Paragraph(f'📦 {big_box_label}', s_boxnum_huge),
                Paragraph('출고 작업 지시서', s_title),
            ]
            left_tbl = Table([[c] for c in left_cell], colWidths=[usable_w * 0.5])
            left_tbl.setStyle(TableStyle([
                ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ('TOPPADDING', (0,0), (-1,-1), 0),
                ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ]))
            left_content = left_tbl
        else:
            left_content = Paragraph('출고 작업 지시서', s_title)

        header_data = [[left_content, right_tbl]]
        header_tbl = Table(header_data, colWidths=[usable_w * 0.5, usable_w * 0.5])
        header_tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'BOTTOM'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_tbl)
    else:
        if big_box_label:
            story.append(Paragraph(f'📦 {big_box_label}', s_boxnum_huge))
        story.append(Paragraph('출고 작업 지시서', s_title))

    story.append(Paragraph(_xml_escape(str(group_key or '')), s_sub))
    story.append(Spacer(1, 1*mm))
    story.append(HRFlowable(width='100%', thickness=2, color=colors.HexColor('#1a56db')))
    story.append(Spacer(1, 4*mm))

    # ── 정보 카드 (3열 테이블) ─────────────────────────
    col_w = usable_w / 3

    def info_card(label, value, big=False, extra=None):
        """[라벨, 값] (+ 부가설명 extra) 문단 리스트 반환"""
        lbl = Paragraph(label, s_card_lbl)
        val = Paragraph(str(value), s_card_big if big else s_card_val)
        return [lbl, val] if extra is None else [lbl, val, extra]

    # 부족이 있으면 '총 수량' 카드에 '− 부족 N' 을 붙이고 최종 수량을 함께 보여준다
    if shortage_qty > 0:
        qty_value = (f'{total_qty} '
                     f'<font size="12" color="#dc2626">- 부족 {shortage_qty}</font>')
        qty_card = info_card(
            '총 수량', qty_value, big=True,
            extra=Paragraph(f'최종 담을 수량 {final_qty}개', s_card_short),
        )
    else:
        qty_card = info_card('총 수량', total_qty, big=True)

    # qty_card 는 <font> 마크업을 일부러 넣은 값이라 이스케이프하지 않는다
    card_data = [[
        info_card('물류센터',  _xml_escape(str(first.get('logisticsCenter','-') or '-'))),
        info_card('입고예정일', _xml_escape(str(first.get('expectedDate','-') or '-'))),
        qty_card,
    ],[
        info_card('송장번호',  _xml_escape(str(first.get('shipmentNumber','-') or '-'))),
        info_card('입고마감일', deadline),
        info_card('품목 수',   f'{len(items)}개'),
    ]]

    def make_card_cell(label_val_list):
        """[lbl_para, val_para, (extra_para)] → 테이블 셀용 nested table"""
        t = Table([[para] for para in label_val_list], colWidths=[col_w - 6*mm])
        t.setStyle(TableStyle([
            ('LEFTPADDING',  (0,0),(-1,-1), 0),
            ('RIGHTPADDING', (0,0),(-1,-1), 0),
            ('TOPPADDING',   (0,0),(-1,-1), 1),
            ('BOTTOMPADDING',(0,0),(-1,-1), 1),
        ]))
        return t

    card_table_data = []
    for row in card_data:
        card_table_data.append([make_card_cell(cell) for cell in row])

    card_bg = [colors.HexColor('#f8fafc'), colors.HexColor('#eff6ff')]
    card_tbl = Table(card_table_data, colWidths=[col_w]*3)
    card_style = [
        ('BACKGROUND', (0,0),(2,0), card_bg[0]),
        ('BACKGROUND', (0,1),(2,1), card_bg[1]),
        ('BOX',        (0,0),(2,1), 1, colors.HexColor('#e2e8f0')),
        ('INNERGRID',  (0,0),(2,1), 0.5, colors.HexColor('#e2e8f0')),
        ('LEFTPADDING',  (0,0),(-1,-1), 4*mm),
        ('RIGHTPADDING', (0,0),(-1,-1), 2*mm),
        ('TOPPADDING',   (0,0),(-1,-1), 3*mm),
        ('BOTTOMPADDING',(0,0),(-1,-1), 3*mm),
        ('VALIGN',     (0,0),(-1,-1), 'MIDDLE'),
        ('ROUNDEDCORNERS', [4]),
    ]
    card_tbl.setStyle(TableStyle(card_style))
    story.append(card_tbl)
    story.append(Spacer(1, 5*mm))

    # ── 상품 테이블 ────────────────────────────────────
    # 부족이 있으면 수량 칸에 '부족 N → 최종 M'을 함께 적어야 하므로
    # 수량 컬럼을 조금 넓히고 상품명 컬럼에서 그만큼 덜어낸다.
    if shortage_qty > 0:
        cw = [usable_w*p for p in [0.24, 0.31, 0.13, 0.16, 0.16]]
    else:
        cw = [usable_w*p for p in [0.25, 0.35, 0.08, 0.16, 0.16]]

    header_row = [
        Paragraph('바코드', s_th),
        Paragraph('상품명', s_th),
        Paragraph('수량', s_th),
        Paragraph('위치', s_th),
        Paragraph('박스', s_th),
    ]
    table_data = [header_row]

    for item in items:
        item_qty   = _qty_int(item)
        item_short = item_shortage_qty(item)
        if item_short > 0:
            # 주문 수량(위) + '부족 N → 최종 M'(아래, 빨강)
            qty_cell = [
                Paragraph(str(item_qty), s_td_bold),
                Paragraph(f'부족 {item_short} → {max(0, item_qty - item_short)}', s_td_short),
            ]
        else:
            qty_cell = Paragraph(str(item_qty), s_td_bold)
        row = [
            Paragraph(_xml_escape(item.get('productBarcode','')), s_mono),
            Paragraph(_xml_escape(item.get('productName','')),    s_td),
            qty_cell,
            Paragraph(_xml_escape(item.get('location','')),       s_td),
            Paragraph(_xml_escape(item.get('boxNumber','')),      s_td),
        ]
        table_data.append(row)

    # 합계 행 — 부족이 있으면 '부족 N개 제외 → 최종 M개'를 함께 표기
    s_sum_note = ParagraphStyle('sumnote', fontName='NanumBold', fontSize=9, leading=12,
                                textColor=colors.HexColor('#dc2626'))
    sum_note = (Paragraph(f'부족 {shortage_qty}개 제외 → 최종 {final_qty}개', s_sum_note)
                if shortage_qty > 0 else Paragraph('', s_td))
    table_data.append([
        Paragraph('합  계', ParagraphStyle('sum', fontName='NanumBold', fontSize=9, textColor=colors.HexColor('#111111'))),
        sum_note,
        Paragraph(str(total_qty), ParagraphStyle('sumqty', fontName='NanumBold', fontSize=12, textColor=colors.HexColor('#1a56db'))),
        Paragraph('', s_td),
        Paragraph('', s_td),
    ])

    tbl = Table(table_data, colWidths=cw, repeatRows=1)

    row_colors = []
    for i in range(1, len(table_data)-1):
        bg = colors.white if i % 2 == 1 else colors.HexColor('#f8fafc')
        row_colors.append(('BACKGROUND', (0,i),(4,i), bg))

    tbl_style = [
        # 헤더
        ('BACKGROUND', (0,0),(4,0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR',  (0,0),(4,0), colors.white),
        ('ALIGN',      (0,0),(4,0), 'CENTER'),
        # 합계 행
        ('BACKGROUND', (0,-1),(4,-1), colors.HexColor('#eff6ff')),
        ('LINEABOVE',  (0,-1),(4,-1), 1, colors.HexColor('#93c5fd')),
        # 전체
        ('FONTSIZE',   (0,0),(-1,-1), 9),
        ('TOPPADDING', (0,0),(-1,-1), 4),
        ('BOTTOMPADDING',(0,0),(-1,-1), 4),
        ('LEFTPADDING',(0,0),(-1,-1), 3*mm),
        ('RIGHTPADDING',(0,0),(-1,-1), 2*mm),
        ('VALIGN',     (0,0),(-1,-1), 'MIDDLE'),
        ('ALIGN',      (2,1),(2,-1), 'CENTER'),  # 수량 가운데
        ('GRID',       (0,0),(-1,-1), 0.4, colors.HexColor('#e2e8f0')),
        ('LINEBELOW',  (0,0),(4,0), 1, colors.HexColor('#1a56db')),
    ] + row_colors

    tbl.setStyle(TableStyle(tbl_style))
    story.append(tbl)
    story.append(Spacer(1, 5*mm))

    # ── 푸터 ──────────────────────────────────────────
    story.append(HRFlowable(width='100%', thickness=0.5, color=colors.HexColor('#e2e8f0')))
    story.append(Spacer(1, 2*mm))
    footer_txt = f'※ 바코드와 수량을 작업 전 반드시 대조해 주세요. (자동 생성 문서) · 생성일시: {created_at}'
    if shortage_qty > 0:
        # 부족분은 박스에 담기지 않으므로 실제 담을 개수를 한 번 더 안내
        footer_txt = (f'※ 부족 {shortage_qty}개는 박스에 담지 않습니다. '
                      f'총 {total_qty}개 - 부족 {shortage_qty}개 = 실제 담을 수량 {final_qty}개 ·' + footer_txt[1:])
    story.append(Paragraph(footer_txt, s_footer))

    doc.build(story)
    buf.seek(0)
    return buf


def _box_str_has_valid_box(bn_raw):
    """박스번호 문자열에 '해외 출고 박스 (▲M7 등)' 정보가 하나라도 있는지.

    - '▲M7(1)'                → True
    - '국내재고(3)'             → False (국내 재고는 송장박스 없음)
    - '부족(-1)'                → False (못 보내는 거)
    - '▲M7(1),부족(-1)'        → True   ← 핵심 수정: 부족 섞여도 ▲M7 살림
    - '국내재고(1),부족(-1)'    → False
    - '▲M7'                    → True   (수량 없어도 박스 정보 있음)

    과거: 문자열에 '부족'/'국내'/'재고'/'RAW' 키워드 하나라도 있으면 domestic
    처리 → 혼합 행도 잘못 제외돼 M열 박스번호 부여 안 됨.
    """
    if not bn_raw:
        return False
    s = str(bn_raw).strip()
    if not s:
        return False
    tokens = [t.strip() for t in re.split(r'[,;/]+', s) if t.strip()]
    for tok in tokens:
        parsed = pick_parse_box(tok)
        # 피킹가능 + 실제 박스 식별자(기호+박스) 있어야 출고 박스
        if parsed.get('상태') == '피킹가능' and parsed.get('박스'):
            return True
    return False


def assign_box_numbers(items):
    """items에서 쉽먼트(송장)별로 박스번호 자동 부여.
    정렬 기준: 입고예정일 > 물류센터 > 송장번호
    송장 전체가 국내재고/부족일 때만 제외 (일부만 국내재고면 부여).
    반환: {송장번호: 박스번호(int)} 딕셔너리
    """
    if not items:
        return {}
    ship_info = {}   # ship → (expectedDate, center, ship)
    ship_valid = {}  # ship → 하나라도 유효한 박스가 있으면 True
    for it in items:
        ship = str(it.get('shipmentNumber', '') or '').strip()
        if not ship:
            continue
        bn_raw = str(it.get('boxNumber', '') or '').strip()
        has_valid_box = _box_str_has_valid_box(bn_raw)
        if ship not in ship_info:
            center = str(it.get('logisticsCenter', '') or '').strip()
            edate = str(it.get('expectedDate', '') or '').strip()
            ship_info[ship] = (edate, center, ship)
            ship_valid[ship] = has_valid_box
        else:
            # 하나라도 유효하면 True로 유지
            if has_valid_box:
                ship_valid[ship] = True
    valid_ships = {s: info for s, info in ship_info.items() if ship_valid.get(s, False)}
    # 정렬: 입고예정일 → 물류센터 → 송장번호
    sorted_ships = sorted(valid_ships.values(), key=lambda x: (x[0], x[1], x[2]))
    return {s[2]: idx + 1 for idx, s in enumerate(sorted_ships)}


def assign_box_numbers_with_existing(items, existing_box_map):
    """기존 매핑(시트 M열)을 보존하고, 그 외 송장만 신규 번호 부여.
    - existing_box_map에 있는 송장은 기존 번호 그대로 유지 (발주 취소돼도 고정)
    - 신규 송장은 max(기존)+1 부터 (입고예정일>물류센터>송장번호) 순으로 부여
    반환: {송장번호: 박스번호(int)}
    """
    result = {}
    for s, v in (existing_box_map or {}).items():
        s = str(s).strip()
        try:
            n = int(str(v).strip())
            if s and n > 0:
                result[s] = n
        except (ValueError, TypeError):
            continue
    ship_info = {}
    ship_valid = {}
    for it in items or []:
        ship = str(it.get('shipmentNumber', '') or '').strip()
        if not ship or ship in result:
            continue
        bn_raw = str(it.get('boxNumber', '') or '').strip()
        has_valid_box = _box_str_has_valid_box(bn_raw)
        if ship not in ship_info:
            center = str(it.get('logisticsCenter', '') or '').strip()
            edate = str(it.get('expectedDate', '') or '').strip()
            ship_info[ship] = (edate, center, ship)
            ship_valid[ship] = has_valid_box
        else:
            if has_valid_box:
                ship_valid[ship] = True
    valid_new = {s: info for s, info in ship_info.items() if ship_valid.get(s, False)}
    sorted_new = sorted(valid_new.values(), key=lambda x: (x[0], x[1], x[2]))
    next_num = max(result.values(), default=0) + 1
    for s in sorted_new:
        result[s[2]] = next_num
        next_num += 1
    return result


def create_box_number_pdf(box_numbers, cols=5, rows=7, label_prefix='', label_suffix=''):
    """박스번호 리스트 → 오려 쓰는 번호표 PDF (A4, 격자 배치) → bytes 반환

    첨부해 주신 엑셀(5열 × 큰 숫자 + 얇은 테두리)과 같은 형태다.
    출고지시서 재출력으로 박스번호가 부여되면, 그 번호 수만큼 번호표를 뽑아
    실제 박스에 붙이는 용도.

    Parameters
    ----------
    box_numbers : iterable[int]
        찍을 번호들. 중복은 제거하고 오름차순 정렬한다.
    cols, rows : int
        A4 한 장에 들어갈 칸 수 (가로 × 세로). 기본 5 × 7 = 35칸/장.
    label_prefix, label_suffix : str
        번호 앞/뒤에 붙일 문구 (예: suffix='번'). 비우면 숫자만.
    """
    # 정렬 + 중복 제거 — 같은 번호표를 두 장 뽑는 사고를 막는다.
    # 시트 M열에서 온 값은 '12' 같은 문자열이거나 빈 칸일 수 있어 숫자만 걸러낸다.
    _clean = set()
    for n in box_numbers or []:
        try:
            _clean.add(int(str(n).strip()))
        except (ValueError, TypeError):
            continue
    nums = sorted(_clean)
    if not nums:
        return b''

    cols = max(1, int(cols))
    rows = max(1, int(rows))

    from reportlab.pdfgen import canvas as _canvas

    buf = io.BytesIO()
    PAGE_W, PAGE_H = A4
    MARGIN = 10 * mm
    cell_w = (PAGE_W - MARGIN * 2) / cols
    cell_h = (PAGE_H - MARGIN * 2) / rows

    c = _canvas.Canvas(buf, pagesize=A4)
    per_page = cols * rows

    # 글자 크기는 "가장 긴 번호" 기준으로 한 번만 정해서 전 칸에 똑같이 쓴다.
    # (칸마다 따로 맞추면 7과 4567의 크기가 달라져 번호표가 들쭉날쭉해진다)
    _texts = [f'{label_prefix}{n}{label_suffix}' for n in nums]
    _widest = max(_texts, key=lambda t: pdfmetrics.stringWidth(t, 'NanumBold', 100))
    _unit_w = pdfmetrics.stringWidth(_widest, 'NanumBold', 1.0) or 1.0
    # 칸 크기에 맞춰 가로/세로 둘 다 안 넘치는 최대 크기를 바로 계산한다.
    # (0.88 / 0.78 은 자를 여백 — 멀리서도 읽히게 칸을 꽉 채운다)
    font_size = min(cell_w * 0.88 / _unit_w, cell_h * 0.78)
    font_size = max(8.0, min(font_size, 200.0))
    c.setFont('NanumBold', font_size)

    for idx, num in enumerate(nums):
        pos = idx % per_page
        if pos == 0 and idx > 0:
            c.showPage()
            c.setFont('NanumBold', font_size)   # showPage 하면 폰트 설정이 풀린다

        r, col = divmod(pos, cols)
        x = MARGIN + col * cell_w
        # reportlab 좌표는 왼쪽 아래가 원점 — 위에서부터 채우려고 뒤집는다.
        y = PAGE_H - MARGIN - (r + 1) * cell_h

        # 가위로 자를 기준선 (엑셀의 얇은 테두리와 같은 역할)
        c.setStrokeColor(colors.HexColor('#999999'))
        c.setLineWidth(0.5)
        c.rect(x, y, cell_w, cell_h, stroke=1, fill=0)

        text = _texts[idx]
        c.setFillColor(colors.black)
        # 시각적 중앙: 대문자/숫자 높이의 절반만큼 아래로 내린다.
        c.drawCentredString(x + cell_w / 2, y + cell_h / 2 - font_size * 0.35, text)

    c.showPage()
    c.save()
    return buf.getvalue()


def create_shipment_barcodes_pdf(shipment_numbers):
    """송장번호 리스트 → 바코드 PDF (한 페이지에 여러 송장 배치)"""
    from reportlab.platypus import Image as RLImage
    buf = io.BytesIO()
    PAGE_W, PAGE_H = A4
    MARGIN = 15 * mm

    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN, bottomMargin=MARGIN
    )

    s_title = ParagraphStyle('btitle', fontName='NanumBold', fontSize=14, leading=18, alignment=1, spaceAfter=8)
    s_num = ParagraphStyle('bnum', fontName='NanumReg', fontSize=10, leading=12, alignment=1, spaceAfter=4)

    story = [Paragraph('📦 송장번호 바코드 (피킹검증 스캔용)', s_title), Spacer(1, 4*mm)]

    # 2열 그리드로 배치
    rows_data = []
    pair = []
    for sn in shipment_numbers:
        try:
            img = get_barcode_img(str(sn), write_text=False)
            img_buf = io.BytesIO()
            img.save(img_buf, format='PNG')
            img_buf.seek(0)
            rl_img = RLImage(img_buf, width=80*mm, height=22*mm)
            rl_img.hAlign = 'CENTER'
            cell = [Paragraph(f'<b>{sn}</b>', s_num), rl_img]
            cell_tbl = Table([[c] for c in cell], colWidths=[85*mm])
            cell_tbl.setStyle(TableStyle([
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('LEFTPADDING', (0,0), (-1,-1), 0),
                ('RIGHTPADDING', (0,0), (-1,-1), 0),
                ('TOPPADDING', (0,0), (-1,-1), 1),
                ('BOTTOMPADDING', (0,0), (-1,-1), 1),
            ]))
            pair.append(cell_tbl)
            if len(pair) == 2:
                rows_data.append(pair)
                pair = []
        except Exception:
            continue
    if pair:
        pair.append(Paragraph('', s_num))
        rows_data.append(pair)

    if rows_data:
        grid = Table(rows_data, colWidths=[90*mm, 90*mm])
        grid.setStyle(TableStyle([
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#cccccc')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cccccc')),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(grid)

    doc.build(story)
    buf.seek(0)
    return buf


def create_box_labels_pdf(box_entries):
    """폼텍 3100 (38.1 x 21.2mm, 5열 13행 = 65칸/페이지) 라벨 PDF 생성

    box_entries: [(box_num, total_qty, size_label), ...] 또는 [(box_num, info_text), ...]
                 각 라벨에 표시할 박스 정보 리스트
    """
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.lib.pagesizes import A4 as _A4
    from reportlab.lib.units import mm as _mm

    # 폼텍 3100 사양 (38.1 × 21.2mm, 5×13 = 65칸)
    LABEL_W = 38.1 * _mm
    LABEL_H = 21.2 * _mm
    X_GAP = 2.54 * _mm      # 라벨 사이 가로 간격 (폼텍 3100 표준)
    Y_GAP = 0 * _mm         # 세로 간격 없음 (라벨이 붙어있음)
    COLS = 5
    ROWS = 13
    PER_PAGE = COLS * ROWS  # 65
    # A4: 210 x 297mm
    # 좌우 여백: (210 - (5*38.1 + 4*2.54)) / 2 = (210 - 200.66) / 2 = 4.67mm
    # 상하 여백: (297 - 13*21.2) / 2 = 10.7mm
    total_w = COLS * LABEL_W + (COLS - 1) * X_GAP
    total_h = ROWS * LABEL_H + (ROWS - 1) * Y_GAP
    LEFT_MARGIN = (210 * _mm - total_w) / 2
    TOP_MARGIN = (297 * _mm - total_h) / 2

    from reportlab.lib.utils import ImageReader as _ImageReader

    buf = io.BytesIO()
    c = _canvas.Canvas(buf, pagesize=_A4)
    page_w, page_h = _A4

    for idx, entry in enumerate(box_entries):
        slot_idx = idx % PER_PAGE
        if slot_idx == 0 and idx > 0:
            c.showPage()

        col = slot_idx % COLS
        row = slot_idx // COLS
        # 프린터 오프셋 보정:
        # - 1, 2열만 왼쪽으로 10% 이동 (3~5열은 그대로)
        # - 전체 위로 25% 이동 (기존 10% + 추가 15%)
        x_offset = -LABEL_W * 0.1 if col < 2 else 0
        y_offset = LABEL_H * 0.25  # PDF는 위로 갈수록 y 증가
        x = LEFT_MARGIN + col * (LABEL_W + X_GAP) + x_offset
        # 좌표 변환: PDF는 좌하단 원점, 라벨은 좌상단부터 채움
        y_top = page_h - TOP_MARGIN - row * (LABEL_H + Y_GAP) + y_offset
        y = y_top - LABEL_H

        # 라벨 내용 파싱
        if isinstance(entry, (tuple, list)):
            box_num = str(entry[0])
            qty = entry[1] if len(entry) > 1 else None
            size_label = entry[2] if len(entry) > 2 else None
        else:
            box_num = str(entry)
            qty = None
            size_label = None

        # 사이즈 라벨에서 한글만 추출 (이모지 제거): 🟢대 → 대
        size_char = ''
        if size_label:
            for ch in size_label:
                if ch in ('대', '중', '소'):
                    size_char = ch
                    break

        # 라벨 레이아웃:
        # ┌─────────────────────┐
        # │ 대   1번    65개    │  ← 상단: 정보
        # │ |||||||||||||||||   │  ← 중하단: 바코드
        # └─────────────────────┘

        # 상단 좌측: 사이즈
        if size_char:
            c.setFont('NanumBold', 11)
            c.drawString(x + LABEL_W * 0.05, y + LABEL_H * 0.68, size_char)

        # 상단 중앙: 박스 번호 (가장 큰 글씨)
        c.setFont('NanumBold', 16)
        c.drawCentredString(x + LABEL_W * 0.48, y + LABEL_H * 0.66,
                            f'{box_num}번')

        # 상단 우측: 수량
        if qty is not None:
            c.setFont('NanumReg', 7)
            c.drawRightString(x + LABEL_W - LABEL_W * 0.05,
                              y + LABEL_H * 0.72, f'{qty}개')

        # 중하단: Code128 바코드 (#N 형식)
        try:
            barcode_text = f'#{box_num}'
            bc_img = get_barcode_img(barcode_text, write_text=False)
            bc_buf = io.BytesIO()
            bc_img.save(bc_buf, format='PNG')
            bc_buf.seek(0)
            bc_w = LABEL_W * 0.85
            bc_h = LABEL_H * 0.42
            c.drawImage(
                _ImageReader(bc_buf),
                x + (LABEL_W - bc_w) / 2,
                y + LABEL_H * 0.1,
                width=bc_w,
                height=bc_h,
                preserveAspectRatio=False,
                mask='auto',
            )
            # 바코드 아래 텍스트
            c.setFont('NanumReg', 6)
            c.drawCentredString(x + LABEL_W / 2, y + LABEL_H * 0.02, barcode_text)
        except Exception:
            pass

    c.save()
    buf.seek(0)
    return buf


def create_multi_trigger_label_pdf():
    """다량 입력 트리거 바코드(#MULTI) 폼텍 3100 형식 A4 1장 (65칸 동일 바코드)"""
    # create_box_labels_pdf 재활용: 65개 동일 엔트리
    entries = [('MULTI', '', '다량')] * 65
    return create_box_labels_pdf(entries)


@lru_cache(maxsize=1)
def _multi_trigger_label_pdf_bytes():
    """#MULTI 라벨 PDF 는 내용이 항상 같으므로 프로세스당 한 번만 만든다."""
    return create_multi_trigger_label_pdf().getvalue()


def _box_labels_pdf_bytes_cached(label_entries):
    """박스 라벨 PDF bytes — 같은 박스 구성이면 세션 캐시 재사용.

    입고분류 화면은 멀티셀렉트 변경, 버튼 클릭, 현황 새로고침마다 전체 rerun 되는데
    그때마다 라벨 PDF(박스 수만큼 바코드 이미지 + reportlab)를 다시 만들고 있었다.
    """
    sig = tuple(tuple(e) for e in label_entries)
    cache = st.session_state.get('_box_label_pdf_cache')
    if cache and cache[0] == sig:
        return cache[1]
    data = create_box_labels_pdf(label_entries).getvalue()
    st.session_state['_box_label_pdf_cache'] = (sig, data)
    return data


# ══════════════════════════════════════════════════════
# Streamlit UI
# ══════════════════════════════════════════════════════
st.set_page_config(page_title='로켓배송 운영 관리', page_icon='🚀', layout='centered')
st.markdown("""<style>
    .block-container { max-width: 58rem !important; }
    /* 피킹 스캔 결과 피드백 */
    .scan-ok {
        background: #d4edda; border-left: 6px solid #28a745;
        padding: 1.2rem 1.5rem; border-radius: 8px; margin: 0.5rem 0; color: #155724;
    }
    .scan-error {
        background: #f8d7da; border-left: 6px solid #dc3545;
        padding: 1.2rem 1.5rem; border-radius: 8px; margin: 0.5rem 0; color: #721c24;
        animation: shake 0.5s ease-in-out;
    }
    .scan-warning {
        background: #fff3cd; border-left: 6px solid #ffc107;
        padding: 1.2rem 1.5rem; border-radius: 8px; margin: 0.5rem 0; color: #856404;
    }
    .scan-complete {
        background: #cce5ff; border-left: 6px solid #007bff;
        padding: 1.2rem 1.5rem; border-radius: 8px; margin: 0.5rem 0; color: #004085;
    }
    .scan-shortage {
        background: #e2e3f1; border-left: 6px solid #6c63ff;
        padding: 1.2rem 1.5rem; border-radius: 8px; margin: 0.5rem 0; color: #383467;
    }
    @keyframes shake {
        0%, 100% { transform: translateX(0); }
        20% { transform: translateX(-10px); }
        40% { transform: translateX(10px); }
        60% { transform: translateX(-6px); }
        80% { transform: translateX(6px); }
    }
    .shipment-input {
        background: #f0f2f6; padding: 2rem; border-radius: 12px; text-align: center;
    }
</style>""", unsafe_allow_html=True)
import kit_ui   # 로그인 · 내 설정 · 관리자
kit_ui.require_login()

st.title('🚀 로켓배송 운영 관리')
st.caption('엑셀 파일을 업로드하면 바코드 이미지를 자동으로 삽입합니다')

# 폰트가 없으면 앱을 막지 않고 알려만 준다 (예전엔 여기서 앱 전체가 죽었다)
if FONT_WARNING:
    st.warning(f'⚠️ {FONT_WARNING}')

# ══════════════════════════════════════════════════════
# 피킹 검증 시스템 — 헬퍼 함수 & 설정
# ══════════════════════════════════════════════════════
PICKING_CONFIG = {
    "SERVICE_ACCOUNT_FILE": "service_account.json",
}

def _extract_sheet_id(url_or_id):
    """구글 시트 URL 또는 ID에서 스프레드시트 ID만 추출"""
    m = re.search(r'/d/([a-zA-Z0-9_-]+)', url_or_id)
    if m:
        return m.group(1)
    return url_or_id.strip()

# ══════════════════════════════════════════════════════
# 구글 시트 API 게이트웨이 — 호출량 절감 (다중 PC 429 방지)
# ══════════════════════════════════════════════════════
# 서비스 계정 1개의 할당량(사용자당 분당 읽기 60 / 쓰기 60)을 접속한 모든
# PC가 함께 나눠 쓴다. 개선 전에는 스캔 1회마다 API를 6번(읽기 3 / 쓰기 3)
# 호출해서 2대만 동시에 스캔해도 한도를 넘겼다.
#   → "APIError: [429] Resource has been exhausted"
#
# 줄인 방법:
#   1) 워크시트 핸들 캐시 — worksheet()가 매번 하던 메타데이터 조회 제거
#   2) 시트 스냅샷 캐시   — 스캔마다 하던 get_all_values() 제거
#   3) 쓰기 큐 + 플러셔   — 셀 쓰기를 모아 batch_update 1회로 전송
#
# 결과: 스캔당 즉시 API 호출 0회. 접속 PC가 몇 대든 시트 하나당
#       _GS_FLUSH_INTERVAL(2초)마다 최대 2회만 호출한다.
#
# ⚠️ 모든 PC가 같은 Streamlit 서버 프로세스에 접속하므로, 아래 모듈 전역
#    캐시는 곧 3대가 공유하는 상태다. 그래서 A PC가 스캔한 누적 수량을
#    B PC가 이어서 누적해도 값이 어긋나지 않는다.
import threading as _gs_threading

_GS_LOCK = _gs_threading.RLock()
_GS_WS_CACHE = {}        # (sheet_id, tab) -> (client, worksheet)
_GS_VALUES_CACHE = {}    # (sheet_id, tab) -> {'values': [[...]], 'ts': monotonic}
_GS_PENDING_CELLS = {}   # (sheet_id, tab) -> {(row1, col1): value}
_GS_PENDING_ROWS = {}    # (sheet_id, tab) -> [[...], ...]  (append용)
_GS_CLIENTS = {}         # (sheet_id, tab) -> client
_GS_ENSURE = {}          # (sheet_id, tab) -> {'rows':.., 'cols':.., 'header':[..]}
_GS_FLUSH_INTERVAL = 2.0 # 쓰기 큐 전송 주기(초)
_GS_VALUES_TTL = 120.0   # 스냅샷 유효 시간(초)
_GS_MAX_RETRY = 5        # 같은 쓰기를 재시도할 최대 횟수
_GS_RETRY_COUNT = {}     # (sheet_id, tab) -> 연속 실패 횟수
_GS_FLUSHER_STARTED = False
_GS_REFRESHING = set()   # 백그라운드 새로 읽기가 진행 중인 (sheet_id, tab)
# 플러시 직렬화용. 플러셔 스레드 / 백그라운드 새로읽기 워커 / '지금 저장' 버튼이
# 동시에 _gs_flush_once 를 부를 수 있는데, 각자 큐에서 서로 다른 값을 들고 나가면
# 네트워크가 느린 쪽(429 백오프 중)이 나중에 도착해 새 값을 옛 값으로 덮어쓴다.
# 한 번에 하나만 보내면 큐에서 꺼낸 순서 = 시트에 닿는 순서가 보장된다.
_GS_FLUSH_LOCK = _gs_threading.Lock()


def _gs_key(sheet_url, tab_name):
    return (_extract_sheet_id(sheet_url), str(tab_name or '').strip())


def _gs_worksheet(client, sheet_url, tab_name, ensure=None):
    """워크시트 핸들 캐시.

    spreadsheet.worksheet(title)은 호출할 때마다 시트 메타데이터 API를 1회
    쓴다. 스캔마다 부르면 그것만으로 할당량이 녹아서 한 번만 열고 재사용한다.
    ensure를 주면 탭이 없을 때 생성한다(피킹로그처럼 자동 생성이 필요한 탭).
    """
    key = _gs_key(sheet_url, tab_name)
    with _GS_LOCK:
        hit = _GS_WS_CACHE.get(key)
        if hit and hit[0] is client:
            return hit[1]
    # 네트워크 호출은 락 밖에서 — 잠깐 중복 조회가 생겨도 무해하다
    spreadsheet = client.open_by_key(key[0])   # 여기서는 API 호출 없음
    try:
        ws = spreadsheet.worksheet(key[1])     # API 1회
    except Exception:
        if not ensure:
            raise
        ws = spreadsheet.add_worksheet(
            title=key[1], rows=ensure.get('rows', 1000), cols=ensure.get('cols', 10))
        header = ensure.get('header')
        if header:
            ws.append_row(header)
    with _GS_LOCK:
        _GS_WS_CACHE[key] = (client, ws)
        _GS_CLIENTS[key] = client
    return ws


def _gs_apply_pending_to_values(key, values):
    """큐에 남은 셀 쓰기를 새로 읽은 스냅샷 위에 덧입힌다 (_GS_LOCK 안에서 호출).

    새로 읽는 사이 스캔이 큐에 넣은 값은 시트에 아직 없으므로, 읽은 값을
    그대로 저장하면 다음 스캔이 옛 값 위에 누적해 수량이 어긋난다.
    """
    pend = _GS_PENDING_CELLS.get(key)
    if not pend:
        return
    for (r, c), v in pend.items():
        while len(values) < r:
            values.append([])
        row = values[r - 1]
        while len(row) < c:
            row.append('')
        row[c - 1] = str(v)


def _gs_refresh_async(client, key):
    """스냅샷을 백그라운드에서 새로 읽는다 — 호출자(스캔)는 기다리지 않는다."""
    with _GS_LOCK:
        if key in _GS_REFRESHING:
            return
        _GS_REFRESHING.add(key)

    def _worker():
        try:
            # 대기 중인 쓰기를 먼저 보내야 새로 읽은 값에 최근 스캔분이 들어있다
            _gs_flush_once()
            ws = _gs_worksheet(client, key[0], key[1])
            values = ws.get_all_values()       # API 1회 (백그라운드)
            with _GS_LOCK:
                _gs_apply_pending_to_values(key, values)
                _GS_VALUES_CACHE[key] = {'values': values, 'ts': time.monotonic()}
        except Exception:
            pass
        finally:
            with _GS_LOCK:
                _GS_REFRESHING.discard(key)

    _gs_threading.Thread(target=_worker, daemon=True, name='gsheet-refresh').start()


def _gs_snapshot(client, sheet_url, tab_name, ttl=None, force=False):
    """시트 전체 값 스냅샷 (get_all_values 캐시).

    반환된 리스트를 호출부가 직접 수정하면 다음 스캔 계산에 바로 반영된다
    (재고 탭에서 이미 쓰던 방식과 동일). 실패 시 None.

    캐시가 있으면 TTL이 지났어도 즉시 반환하고 새로 읽기는 백그라운드로
    돌린다(stale-while-revalidate). TTL 만료 시점의 스캔이 시트 전체 읽기
    (+429 백오프 대기)를 기다리며 수 초씩 멈추던 문제를 없애기 위함이다.
    작업 중에는 이 서버 전역 캐시가 사실상 원본이라 즉시 반환해도 안전하고,
    백그라운드 새로 읽기는 시트를 직접 고친 외부 수정분만 끌어온다.
    force=True(연결/새로고침 버튼)만 동기로 새로 읽는다.
    """
    key = _gs_key(sheet_url, tab_name)
    ttl = _GS_VALUES_TTL if ttl is None else ttl
    now = time.monotonic()
    if not force:
        with _GS_LOCK:
            hit = _GS_VALUES_CACHE.get(key)
            if hit is not None:
                _GS_CLIENTS[key] = client
        if hit is not None:
            if (now - hit['ts']) >= ttl:
                _gs_refresh_async(client, key)
            return hit['values']
    # 여기 도달 = 캐시가 아예 없거나(첫 연결) force. 이때만 동기로 읽는다.
    # 새로 읽기 전에 대기 중인 쓰기를 먼저 반영해야 방금 스캔분이 사라지지 않는다
    if force:
        _gs_flush_once()
    try:
        ws = _gs_worksheet(client, sheet_url, tab_name)
        values = ws.get_all_values()           # API 1회
    except Exception:
        with _GS_LOCK:
            hit = _GS_VALUES_CACHE.get(key)
        return hit['values'] if hit else None
    with _GS_LOCK:
        _gs_apply_pending_to_values(key, values)
        _GS_VALUES_CACHE[key] = {'values': values, 'ts': time.monotonic()}
        _GS_CLIENTS[key] = client
    return values


def _gs_invalidate(sheet_url, tab_name):
    """스냅샷 무효화 — 다음 조회에서 시트를 다시 읽는다."""
    with _GS_LOCK:
        _GS_VALUES_CACHE.pop(_gs_key(sheet_url, tab_name), None)


def _gs_queue_cells(client, sheet_url, tab_name, cells):
    """셀 쓰기를 큐에 넣는다 (즉시 API 호출 없음).

    cells: {(row_1based, col_1based): 값}
    같은 셀을 여러 번 넣으면 마지막 값만 남는다 — 증분이 아니라 누적된
    최종값을 쓰기 때문에 중간 값이 사라져도 결과가 같다.
    """
    if not cells:
        return
    key = _gs_key(sheet_url, tab_name)
    with _GS_LOCK:
        _GS_CLIENTS[key] = client
        _GS_PENDING_CELLS.setdefault(key, {}).update(cells)
    _gs_start_flusher()


def _gs_queue_append(client, sheet_url, tab_name, row_values, ensure=None):
    """행 추가(로그)를 큐에 넣는다. 플러시 때 append_rows로 한 번에 보낸다."""
    key = _gs_key(sheet_url, tab_name)
    with _GS_LOCK:
        _GS_CLIENTS[key] = client
        if ensure:
            _GS_ENSURE[key] = ensure
        _GS_PENDING_ROWS.setdefault(key, []).append(list(row_values))
    _gs_start_flusher()


def _gs_flush_once():
    """큐에 쌓인 쓰기를 시트별로 묶어 전송. 셀은 batch_update 1회,
    로그는 append_rows 1회. 실패분은 다음 주기에 재시도한다.

    동시에 두 번 돌지 않도록 _GS_FLUSH_LOCK 으로 직렬화한다 — 호출자 중 하나가
    이미 보내는 중이면 끝날 때까지 기다렸다가 그 다음 큐를 가져간다."""
    with _GS_FLUSH_LOCK:
        _gs_flush_once_locked()


def _gs_flush_once_locked():
    from gspread.utils import rowcol_to_a1
    with _GS_LOCK:
        cell_jobs = {k: v for k, v in _GS_PENDING_CELLS.items() if v}
        row_jobs = {k: v for k, v in _GS_PENDING_ROWS.items() if v}
        _GS_PENDING_CELLS.clear()
        _GS_PENDING_ROWS.clear()
        clients = dict(_GS_CLIENTS)
        ensures = dict(_GS_ENSURE)

    for key, cells in cell_jobs.items():
        client = clients.get(key)
        if client is None:
            continue
        try:
            ws = _gs_worksheet(client, key[0], key[1])
            data = [{'range': rowcol_to_a1(r, c), 'values': [[v]]}
                    for (r, c), v in sorted(cells.items())]
            # RAW = 기존 update_cell의 기본 동작. USER_ENTERED로 바꾸면 위치값
            # "3-4" 같은 문자열이 날짜로 자동 변환될 수 있어 그대로 유지한다.
            ws.batch_update(data, value_input_option='RAW')            # API 1회
            with _GS_LOCK:
                _GS_RETRY_COUNT.pop(key, None)
        except Exception:
            with _GS_LOCK:
                n = _GS_RETRY_COUNT.get(key, 0) + 1
                _GS_RETRY_COUNT[key] = n
                if n <= _GS_MAX_RETRY:
                    pend = _GS_PENDING_CELLS.setdefault(key, {})
                    for cell, val in cells.items():
                        pend.setdefault(cell, val)   # 새로 쌓인 값이 우선
                # 한도를 넘으면 버린다 — 영구 오류를 2초마다 재시도하며
                # 할당량을 태우는 게 더 나쁘다. 다음 스캔이 다시 큐에 넣는다.

    for key, rows in row_jobs.items():
        client = clients.get(key)
        if client is None:
            continue
        try:
            ws = _gs_worksheet(client, key[0], key[1], ensure=ensures.get(key))
            ws.append_rows(rows, value_input_option='RAW')             # API 1회
            with _GS_LOCK:
                _GS_RETRY_COUNT.pop(key, None)
        except Exception:
            with _GS_LOCK:
                n = _GS_RETRY_COUNT.get(key, 0) + 1
                _GS_RETRY_COUNT[key] = n
                if n <= _GS_MAX_RETRY:
                    _GS_PENDING_ROWS.setdefault(key, [])[:0] = rows    # 순서 보존


def _gs_start_flusher():
    """쓰기 큐 플러셔 스레드를 한 번만 띄운다."""
    global _GS_FLUSHER_STARTED
    with _GS_LOCK:
        if _GS_FLUSHER_STARTED:
            return
        _GS_FLUSHER_STARTED = True

    def _loop():
        while True:
            time.sleep(_GS_FLUSH_INTERVAL)
            try:
                _gs_flush_once()
            except Exception:
                pass

    _gs_threading.Thread(target=_loop, daemon=True, name='gsheet-flusher').start()


def gs_pending_count():
    """아직 시트로 전송되지 않고 큐에 남아있는 쓰기 건수."""
    with _GS_LOCK:
        return (sum(len(v) for v in _GS_PENDING_CELLS.values())
                + sum(len(v) for v in _GS_PENDING_ROWS.values()))


def gs_flush_pending():
    """대기 중인 시트 쓰기를 지금 즉시 전송 (작업 종료/시트 확인 직전용)."""
    try:
        _gs_flush_once()
        return True
    except Exception:
        return False


def _gs_is_quota_error(exc):
    """429(할당량 초과)인지 판별."""
    return '429' in str(exc) or 'exhausted' in str(exc).lower()


@st.cache_resource(ttl=3600)
def get_gsheet_client():
    try:
        import gspread
        from google.oauth2.service_account import Credentials
        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        # 1순위: PC 모드의 서비스 계정 키 파일 (프로그램 폴더의 service_account.json)
        key_path = kit_config.service_account_path()
        if key_path.exists():
            creds = Credentials.from_service_account_file(str(key_path), scopes=scopes)
        # 2순위: Streamlit Secrets (secrets.toml 이 없으면 FileNotFoundError)
        elif "gcp_service_account" in st.secrets:
            creds = Credentials.from_service_account_info(
                dict(st.secrets["gcp_service_account"]), scopes=scopes
            )
        else:
            return None
        # BackOffHTTPClient: 429/5xx를 만나면 스스로 기다렸다 재시도한다.
        # 호출량을 줄이는 게 1차 방어, 이건 그래도 몰릴 때의 2차 방어.
        client = gspread.authorize(creds, http_client=gspread.BackOffHTTPClient)
        # 타임아웃이 없으면 응답이 끊긴 요청이 영원히 대기해서
        # "구글 시트 연결 중..." 스피너가 끝나지 않는다. (연결 10초 / 응답 60초)
        client.set_timeout((10, 60))
        return client
    except FileNotFoundError:
        return None
    except Exception as e:
        st.warning(f"구글 시트 연결 실패: {e}")
        return None

def pick_load_sheet_as_df(client, sheet_url, tab_name):
    try:
        import pandas as _pd
        # 사용자가 '연결'을 누른 시점이므로 캐시 무시하고 새로 읽는다.
        all_values = _gs_snapshot(client, sheet_url, tab_name, force=True)
        if all_values is None:
            raise RuntimeError('시트를 읽지 못했습니다')
        if len(all_values) < 2:
            return _pd.DataFrame()
        headers = all_values[0]
        # 중복 컬럼명 처리: 같은 이름이면 _2, _3 붙임
        seen = {}
        unique_headers = []
        for h in headers:
            h = str(h).strip()
            if h == '':
                h = f'_unnamed_{len(unique_headers)}'
            if h in seen:
                seen[h] += 1
                unique_headers.append(f"{h}_{seen[h]}")
            else:
                seen[h] = 1
                unique_headers.append(h)
        # 캐시된 원본을 나중에 수정해도 DataFrame이 영향받지 않도록 복사
        df = _pd.DataFrame([list(r) for r in all_values[1:]], columns=unique_headers)
        return df
    except Exception as e:
        if _gs_is_quota_error(e):
            st.error(
                f"시트 '{tab_name}' 로드 실패: 구글 API 사용량 한도(429)에 걸렸습니다. "
                "다른 PC의 스캔을 잠시 멈추고 30초 뒤 다시 연결해 주세요."
            )
        else:
            st.error(f"시트 '{tab_name}' 로드 실패: {e}")
        return None

def pick_append_log(client, sheet_url, log_entry):
    """피킹로그 추가 — 큐에 넣고 즉시 반환 (플러셔가 모아서 1회 전송)."""
    try:
        _gs_queue_append(
            client, sheet_url, "피킹로그", log_entry,
            ensure={'rows': 1000, 'cols': 10,
                    'header': ["시간","송장번호","바코드","상품명","결과",
                               "스캔수량","필요수량","회차기호","박스번호"]},
        )
        return True
    except Exception:
        return False

def pick_update_sheet_inventory(client, sheet_url, tab_name, barcode, decrement=1, box_number=None):
    """배대지 시트 스캔 수량 기록.
    매칭: E열(바코드, 4) 일치하는 행 중 A열(박스번호, 0) 일치 우선, 없으면 바코드만 매칭.
    기록:
      - U열(스캔수량, 20): 기존값에 decrement 만큼 누적
      - V열(남은수량, 21): G열 수량 - U열 스캔수량 (안 온 상품 체크용)
    G열(수량)은 건드리지 않음 — 원본 주문 수량으로 보존.

    박스번호 정규화: 영문숫자만 추출해서 비교 (예: "●M6(17)" → "M6").

    스캔마다 시트를 통째로 다시 읽지 않고 스냅샷 캐시를 쓰며, 계산한 누적값을
    캐시에 바로 반영한 뒤 쓰기는 큐에 넘긴다 (즉시 API 호출 0회).
    """
    import re as _re_box
    try:
        all_values = _gs_snapshot(client, sheet_url, tab_name)
        if all_values is None or len(all_values) < 2:
            return False
        _BOX_COL = 0   # A열
        _BC_COL = 4    # E열 (바코드)
        _QTY_COL = 6   # G열 (원본 수량)
        _SCAN_COL = 20 # U열 (누적 스캔)
        _REM_COL = 21  # V열 (남은 수량)

        def _norm(s):
            """박스번호 정규화 — 영문+숫자만 추출 대문자화 ('●M6(17)' → 'M6')"""
            return _re_box.sub(r'[^A-Z0-9]', '', str(s or '').upper())

        target_box = _norm(box_number)
        barcode = str(barcode or '').strip()
        if not barcode:
            return False

        # 읽기(prev_scan) → 계산 → 캐시 쓰기 → 큐 넣기 를 한 덩어리로 잠근다.
        # 캐시 리스트는 접속한 모든 PC가 공유하므로, 두 PC가 같은 바코드를 거의 동시에
        # 찍으면 둘 다 같은 prev 를 읽고 같은 값을 써서 스캔 1건이 사라진다.
        # _GS_LOCK 은 RLock 이라 안에서 부르는 _gs_queue_cells 가 다시 잡아도 된다.
        with _GS_LOCK:
            # 1차: 바코드 + 박스번호 정확 매칭
            # 2차: 바코드만 매칭 (박스 불일치일 때 fallback)
            primary_row = -1
            fallback_row = -1
            for row_idx in range(1, len(all_values)):
                row = all_values[row_idx]
                if len(row) <= _BC_COL:
                    continue
                if str(row[_BC_COL]).strip() != barcode:
                    continue
                row_box = _norm(row[_BOX_COL]) if len(row) > _BOX_COL else ''
                if target_box and row_box == target_box:
                    primary_row = row_idx
                    break
                if fallback_row < 0:
                    fallback_row = row_idx
            hit_row = primary_row if primary_row >= 0 else fallback_row
            if hit_row < 0:
                return False

            row = all_values[hit_row]
            # G열 수량 (원본)
            qty_raw = row[_QTY_COL] if len(row) > _QTY_COL else ''
            try:
                qty_orig = int(float(str(qty_raw).strip() or '0'))
            except (ValueError, TypeError):
                qty_orig = 0
            # U열 기존 스캔 수량
            prev_raw = row[_SCAN_COL] if len(row) > _SCAN_COL else ''
            try:
                prev_scan = int(float(str(prev_raw).strip() or '0'))
            except (ValueError, TypeError):
                prev_scan = 0
            new_scan = max(0, prev_scan + int(decrement))
            remaining = qty_orig - new_scan

            # 캐시 즉시 반영 — 다음 스캔이 이 값 위에 누적된다 (다른 PC 포함)
            while len(all_values[hit_row]) <= max(_SCAN_COL, _REM_COL):
                all_values[hit_row].append('')
            all_values[hit_row][_SCAN_COL] = str(new_scan)
            all_values[hit_row][_REM_COL] = str(remaining)

            _gs_queue_cells(client, sheet_url, tab_name, {
                (hit_row + 1, _SCAN_COL + 1): new_scan,
                (hit_row + 1, _REM_COL + 1): remaining,
            })
            return True
    except Exception:
        return False


def pick_update_check_qty(client, sheet_url, tab_name, barcode, ship_num, scanned_qty):
    """출고확인 시트의 해당 행 L열(확인 수량)에 스캔된 수량 기록.
    매칭: F열(바코드) + I열(송장번호)
    """
    try:
        all_values = _gs_snapshot(client, sheet_url, tab_name)
        if all_values is None or len(all_values) < 2:
            return False
        # 여러 PC 동시 스캔 시 lost update 방지 (pick_update_sheet_inventory 와 동일)
        with _GS_LOCK:
            # F열(5) = 바코드, I열(8) = 송장번호, L열(11) = 확인 수량 → 12번째 열
            for row_idx in range(1, len(all_values)):
                row = all_values[row_idx]
                row_bc = str(row[5]).strip() if len(row) > 5 else ''
                row_ship = str(row[8]).strip() if len(row) > 8 else ''
                if row_bc == barcode and (not ship_num or row_ship == ship_num):
                    while len(all_values[row_idx]) <= 11:
                        all_values[row_idx].append('')
                    all_values[row_idx][11] = str(scanned_qty)
                    _gs_queue_cells(client, sheet_url, tab_name,
                                    {(row_idx + 1, 12): scanned_qty})
                    return True
            return False
    except Exception:
        return False


_SHEET_SHIP_COL_IDX = 8   # I열 (송장번호), 0-based
_SHEET_QTY_COL_IDX = 11   # L열 (확인수량), 0-based → update_cell에는 +1=12
_SHEET_BOX_COL_IDX = 12   # M열 (출고박스번호), 0-based → update_cell에는 +1=13
_SHEET_BOX_COL_LETTER = 'M'


def pick_read_box_numbers(client, sheet_url, tab_name):
    """출고확인 시트에서 송장번호(I열) → 출고박스번호(M열) 매핑을 읽음.
    M열이 비어있거나 숫자가 아닌 행은 제외.
    반환: {송장번호: 박스번호(int)} (성공, 빈 dict 포함 가능)
         / None (API 호출 실패 — 읽기 자체가 안 됨)
    """
    try:
        # 박스번호는 정확해야 하므로 캐시를 무시하고 새로 읽는다
        all_values = _gs_snapshot(client, sheet_url, tab_name, force=True)
    except Exception:
        return None
    if all_values is None:
        return None
    if len(all_values) < 2:
        return {}
    mapping = {}
    for row in all_values[1:]:
        ship = str(row[_SHEET_SHIP_COL_IDX]).strip() if len(row) > _SHEET_SHIP_COL_IDX else ''
        box = str(row[_SHEET_BOX_COL_IDX]).strip() if len(row) > _SHEET_BOX_COL_IDX else ''
        if not ship or not box:
            continue
        try:
            # "36" / " 36 " / "36.0" (엑셀 float 표기) 모두 허용
            n = int(float(box))
            if n > 0:
                mapping[ship] = n
        except (ValueError, TypeError):
            continue
    return mapping


def pick_write_box_numbers(client, sheet_url, tab_name, ship_to_box, only_empty=True):
    """출고확인 시트 M열에 송장별 박스번호 기록 (batch).
    only_empty=True이면 M열이 비어있는 행만 쓰기(기존 값 보존).
    반환: 기록된 셀 수 / -1 (API 실패 시)
    """
    if not ship_to_box:
        return 0
    try:
        ws = _gs_worksheet(client, sheet_url, tab_name)
        all_values = _gs_snapshot(client, sheet_url, tab_name, force=True)
        if all_values is None or len(all_values) < 2:
            return 0
        updates = []
        written_rows = {}
        for row_idx in range(1, len(all_values)):
            row = all_values[row_idx]
            ship = str(row[_SHEET_SHIP_COL_IDX]).strip() if len(row) > _SHEET_SHIP_COL_IDX else ''
            if not ship or ship not in ship_to_box:
                continue
            current_m = str(row[_SHEET_BOX_COL_IDX]).strip() if len(row) > _SHEET_BOX_COL_IDX else ''
            if only_empty and current_m:
                continue
            updates.append({
                'range': f'{_SHEET_BOX_COL_LETTER}{row_idx + 1}',
                'values': [[str(ship_to_box[ship])]],
            })
            written_rows[row_idx] = str(ship_to_box[ship])
        if updates:
            ws.batch_update(updates)
            for row_idx in written_rows:
                while len(all_values[row_idx]) <= _SHEET_BOX_COL_IDX:
                    all_values[row_idx].append('')
                all_values[row_idx][_SHEET_BOX_COL_IDX] = written_rows[row_idx]
        return len(updates)
    except Exception:
        return -1


def stock_update_barcode(client, sheet_url, tab_name, barcode, qty, location):
    """등록상품정보 시트에서 바코드(D열, index 3) 매칭 후 재고/위치 업데이트.
    - X열(index 23): 기존 재고 + qty 누적
    - V열(index 21): 위치 추가 (중복이면 스킵, 여러 위치는 ", " 구분)
    - C열(index 2): 상품명 반환용
    반환: dict {ok: bool, name: str, new_stock: int, error: str, row_idx: int}
    """
    _BC_COL = 3       # D열
    _NAME_COL = 2     # C열
    _STOCK_COL = 23   # X열 (24번째)
    _LOC_COL = 21     # V열 (22번째)

    try:
        qty_int = int(qty)
    except (ValueError, TypeError):
        qty_int = 1
    if qty_int <= 0:
        return {'ok': False, 'name': '', 'new_stock': 0, 'error': '수량은 1 이상이어야 함'}

    # 시트 데이터 스냅샷 (세션별이 아니라 서버 전역 캐시 — 여러 PC가 함께 쓴다)
    try:
        all_values = _gs_snapshot(client, sheet_url, tab_name)
    except Exception as e:
        return {'ok': False, 'name': '', 'new_stock': 0, 'error': f'시트 열기 실패: {e}'}
    if all_values is None:
        return {'ok': False, 'name': '', 'new_stock': 0, 'error': '시트를 읽지 못했습니다'}
    if len(all_values) < 2:
        return {'ok': False, 'name': '', 'new_stock': 0, 'error': '시트가 비어있음'}

    # 여러 PC 동시 스캔 시 lost update 방지 (pick_update_sheet_inventory 와 동일)
    with _GS_LOCK:
        for row_idx in range(1, len(all_values)):
            row = all_values[row_idx]
            if len(row) <= _BC_COL:
                continue
            if str(row[_BC_COL]).strip() != str(barcode).strip():
                continue

            # X열 기존 재고 + qty
            prev_raw = row[_STOCK_COL] if len(row) > _STOCK_COL else ''
            try:
                prev = int(float(str(prev_raw).strip() or '0'))
            except (ValueError, TypeError):
                prev = 0
            new_stock = prev + qty_int

            # V열 위치 (중복 방지)
            existing_loc = str(row[_LOC_COL]).strip() if len(row) > _LOC_COL else ''
            loc_new = str(location or '').strip()
            updated_loc = existing_loc
            if loc_new:
                if not existing_loc:
                    updated_loc = loc_new
                else:
                    parts = [s.strip() for s in existing_loc.split(',') if s.strip()]
                    if loc_new not in parts:
                        parts.append(loc_new)
                        updated_loc = ', '.join(parts)

            # 캐시 즉시 반영 (다음 스캔에서 누적 정확하게)
            while len(all_values[row_idx]) <= max(_STOCK_COL, _LOC_COL):
                all_values[row_idx].append('')
            all_values[row_idx][_STOCK_COL] = str(new_stock)
            if updated_loc != existing_loc:
                all_values[row_idx][_LOC_COL] = updated_loc

            # 쓰기는 큐로 — 플러셔가 모아서 batch_update 1회로 보낸다 (즉시 API 0회)
            _cells = {(row_idx + 1, _STOCK_COL + 1): new_stock}
            if updated_loc != existing_loc:
                _cells[(row_idx + 1, _LOC_COL + 1)] = updated_loc
            _gs_queue_cells(client, sheet_url, tab_name, _cells)

            name = str(row[_NAME_COL]).strip() if len(row) > _NAME_COL else ''
            return {'ok': True, 'name': name, 'new_stock': new_stock, 'error': ''}
        return {'ok': False, 'name': '', 'new_stock': 0, 'error': f'미등록 바코드: {barcode}'}


def stock_change_location(client, sheet_url, tab_name, old_loc, new_loc, preview_only=False):
    """등록상품정보 시트 V열(위치)에서 기존 위치를 새 위치로 일괄 교체.
    - V열 형식: ", " 구분된 여러 위치 (예: "G박스, A-1")
    - old_loc 토큰만 정확히 매칭해서 new_loc으로 교체 (다른 위치는 보존)
    - new_loc이 빈 문자열이면 old_loc만 제거
    - 같은 행에 new_loc이 이미 있으면 중복 추가 안 함
    - preview_only=True면 시트 안 건드리고 변경 대상만 반환

    반환: {ok, total_changed, changes: [{row, name, barcode, old, new}]}
    """
    _BC_COL = 3       # D열
    _NAME_COL = 2     # C열
    _LOC_COL = 21     # V열 (22번째)

    old_loc = str(old_loc or '').strip()
    new_loc = str(new_loc or '').strip()
    if not old_loc:
        return {'ok': False, 'total_changed': 0, 'changes': [], 'error': '기존 위치를 입력하세요'}

    try:
        ws = _gs_worksheet(client, sheet_url, tab_name)
        all_values = _gs_snapshot(client, sheet_url, tab_name, force=True)
    except Exception as e:
        return {'ok': False, 'total_changed': 0, 'changes': [], 'error': f'시트 열기 실패: {e}'}
    if all_values is None or len(all_values) < 2:
        return {'ok': False, 'total_changed': 0, 'changes': [], 'error': '시트가 비어있음'}

    changes = []
    updates = []  # batch_update용
    for row_idx in range(1, len(all_values)):
        row = all_values[row_idx]
        if len(row) <= _LOC_COL:
            continue
        existing = str(row[_LOC_COL]).strip()
        if not existing:
            continue
        parts = [s.strip() for s in existing.split(',') if s.strip()]
        if old_loc not in parts:
            continue
        # 토큰 교체
        if new_loc:
            new_parts = []
            for p in parts:
                if p == old_loc:
                    # 같은 행에 new_loc 이 이미 있으면(다른 토큰으로든, 방금 넣었든) 중복 추가 안 함
                    if new_loc not in new_parts and new_loc not in parts:
                        new_parts.append(new_loc)
                else:
                    if p not in new_parts:
                        new_parts.append(p)
            updated = ', '.join(new_parts)
        else:
            # new_loc 비었으면 old_loc 제거만
            updated = ', '.join(p for p in parts if p != old_loc)

        if updated == existing:
            continue

        bc = str(row[_BC_COL]).strip() if len(row) > _BC_COL else ''
        nm = str(row[_NAME_COL]).strip() if len(row) > _NAME_COL else ''
        changes.append({
            'row': row_idx + 1,
            'barcode': bc,
            'name': nm[:40],
            'old': existing,
            'new': updated,
        })
        if not preview_only:
            # V열은 인덱스 21 → A1 표기 'V' (22번째 열)
            updates.append({
                'range': f'V{row_idx + 1}',
                'values': [[updated]],
            })

    if not preview_only and updates:
        try:
            ws.batch_update(updates)
            # 캐시 무효화 (다음 조회 시 신선한 데이터 반영)
            _gs_invalidate(sheet_url, tab_name)
        except Exception as e:
            return {'ok': False, 'total_changed': 0, 'changes': changes,
                    'error': f'시트 쓰기 실패: {e}'}

    return {'ok': True, 'total_changed': len(changes), 'changes': changes, 'error': ''}


def _tts_signature_tone_js():
    """프로필별 시그니처 비프음 JS — TTS 앞에 재생되어 양쪽 노트북 구분.
    pick_tts_tone_freq + pick_tts_tone_pattern 기반.
    """
    freq = int(st.session_state.get('pick_tts_tone_freq', 880))
    pattern = str(st.session_state.get('pick_tts_tone_pattern', 'ding'))
    # 패턴별 멜로디 정의 (시작주파수, [(시간ms, 주파수, 길이ms)...])
    if pattern == 'ding-ding':
        notes = [(0, freq, 100), (130, freq, 100)]
    elif pattern == 'do-re-mi':
        notes = [(0, int(freq*0.84), 80), (90, int(freq*0.95), 80), (180, freq, 120)]
    elif pattern == 'bzz-bzz':
        notes = [(0, freq, 150), (200, freq, 150)]
    elif pattern == 'low-low-high':
        notes = [(0, freq, 100), (130, freq, 100), (260, int(freq*1.6), 150)]
    else:
        notes = [(0, freq, 120)]
    # 각 음을 schedule (Web Audio API 사용)
    parts = ["try{var aC=new(window.AudioContext||window.webkitAudioContext)();"]
    for i, (delay, f, dur) in enumerate(notes):
        parts.append(
            f"setTimeout(function(){{"
            f"var o=aC.createOscillator();var g=aC.createGain();"
            f"o.type='square';"
            f"o.frequency.value={f};"
            f"g.gain.value=0.25;"
            f"o.connect(g);g.connect(aC.destination);"
            f"o.start();"
            f"setTimeout(function(){{g.gain.value=0;o.stop();}},{dur});"
            f"}},{delay});"
        )
    parts.append("}catch(e){}")
    return ''.join(parts)


# 한국어 → 영어 짧은 메시지 변환표 (영어 voice 사용 시 → 발음 정확)
_KO_TO_EN_TTS_MAP = {
    '재고완료': 'OK!',
    '다시 찍어주세요': 'Retry',
    '확인을 시작하세요': 'Start',
    '검증확인이 완료되었습니다. 출고하세요': 'Done. Ship it.',
    '수량을 입력하세요': 'Enter quantity',
    '바코드를 스캔해 주세요': 'Scan barcode',
    '분류 완료': 'Sort done',
    '보류': 'Hold',
    '다른 박스 상품': 'Wrong box',
    '미스캔 있음': 'Missing scan',
}


def _ko_message_to_en(msg):
    """한국어 메시지를 영어로 변환 (영어 voice 사용 시).
    - 정확 매칭 우선
    - "N번 완료. 포장하세요" → "Box N done, pack"
    - "N번" → "Box N"
    - "N번 K개" → "Box N, K items"
    - 못 찾으면 원문 반환
    """
    import re as _re_en
    s = str(msg).strip()
    if s in _KO_TO_EN_TTS_MAP:
        return _KO_TO_EN_TTS_MAP[s]
    # "N번 완료. 포장하세요" 패턴 (음성으로 된 한자 숫자)
    m = _re_en.match(r'^(\S+)번\s*완료\.?\s*포장하세요\.?$', s)
    if m:
        return f'Box {m.group(1)} done, pack'
    # "N번" 단독
    m = _re_en.match(r'^(\S+)번$', s)
    if m:
        return f'Box {m.group(1)}'
    # "N번 K개"
    m = _re_en.match(r'^(\S+)번\s*(\d+)개$', s)
    if m:
        return f'Box {m.group(1)}, {m.group(2)} items'
    # "N번 미스캔 있음"
    m = _re_en.match(r'^(\S+)번\s*미스캔\s*있음$', s)
    if m:
        return f'Box {m.group(1)}, missing scan'
    return s  # 변환 못 한 건 그대로


_EDGE_TTS_CACHE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.tts_cache')

# ── edge-tts 생성은 절대 스캔을 기다리게 하지 않는다 ──
# 예전에는 캐시에 없는 문구를 만나면 스캔 렌더 도중 Microsoft 서버에 동기로
# 요청하고 최대 6초를 기다렸다. 입고분류의 "삼번", "십이번 완료. 포장하세요"
# 처럼 박스마다 달라지는 문구는 워밍업에 없어서, 새 박스를 처음 찍을 때마다
# 화면이 멈춘 것처럼 보였고, 네트워크가 막힌 날은 문구마다 6초씩 멈췄다.
#
# 지금 구조:
#   1) 메모리 캐시(_EDGE_TTS_MEM) → 디스크 캐시 순으로 즉시 반환
#   2) 없으면 백그라운드 스레드에서 생성 시작 → 최대 _EDGE_TTS_WAIT_BUDGET 초만
#      기다렸다가 아직이면 브라우저 TTS로 폴백 (다음 스캔부터는 캐시 히트)
#   3) 생성이 실패하면 _EDGE_TTS_FAIL_COOLDOWN 초 동안은 시도 자체를 건너뛴다
_EDGE_TTS_MEM = {}                 # 캐시 키 → base64 mp3
_EDGE_TTS_INFLIGHT = {}            # 캐시 키 → threading.Event (생성 중)
_EDGE_TTS_LOCK = _gs_threading.Lock()
_EDGE_TTS_FAIL_UNTIL = 0.0         # 이 시각(monotonic)까지는 생성 시도 안 함
_EDGE_TTS_FAIL_COOLDOWN = 60.0     # 실패 후 재시도까지 대기(초)
_EDGE_TTS_WAIT_BUDGET = 0.8        # 스캔이 생성을 기다려 줄 최대 시간(초)
_EDGE_TTS_GEN_TIMEOUT = 8          # 백그라운드 생성 타임아웃(초)


def _edge_tts_cache_key(text, voice, rate_pct):
    import hashlib
    return hashlib.md5(f'{voice}|{rate_pct}|{text}'.encode()).hexdigest()


def _edge_tts_load_cached(key):
    """메모리 → 디스크 순으로 캐시 조회. 없으면 None (네트워크 없음)."""
    import base64
    hit = _EDGE_TTS_MEM.get(key)
    if hit:
        return hit
    path = os.path.join(_EDGE_TTS_CACHE_DIR, key + '.mp3')
    try:
        if os.path.exists(path):
            with open(path, 'rb') as f:
                b64 = base64.b64encode(f.read()).decode()
            if b64:
                _EDGE_TTS_MEM[key] = b64
                return b64
    except Exception:
        pass
    return None


def _edge_tts_generate_blocking(text, voice, rate_pct, key):
    """실제 생성(네트워크). 백그라운드 스레드에서만 부른다. 성공 시 캐시에 넣는다."""
    import base64
    global _EDGE_TTS_FAIL_UNTIL
    try:
        import asyncio
        import edge_tts

        async def _gen():
            c = edge_tts.Communicate(text, voice, rate=f'{rate_pct:+d}%')
            out = b''
            async for ch in c.stream():
                if ch['type'] == 'audio':
                    out += ch['data']
            return out

        data = asyncio.run(asyncio.wait_for(_gen(), timeout=_EDGE_TTS_GEN_TIMEOUT))
        if not data:
            raise RuntimeError('empty audio')
        try:
            os.makedirs(_EDGE_TTS_CACHE_DIR, exist_ok=True)
            with open(os.path.join(_EDGE_TTS_CACHE_DIR, key + '.mp3'), 'wb') as f:
                f.write(data)
        except Exception:
            pass   # 디스크 캐시 실패는 치명적이지 않다 (메모리 캐시는 남는다)
        b64 = base64.b64encode(data).decode()
        _EDGE_TTS_MEM[key] = b64
        return b64
    except Exception:
        # 네트워크/서비스 장애 — 한동안 시도하지 않는다 (스캔마다 기다리지 않도록)
        _EDGE_TTS_FAIL_UNTIL = time.monotonic() + _EDGE_TTS_FAIL_COOLDOWN
        return None


def _edge_tts_start_async(text, voice, rate_pct, key):
    """생성 스레드를 띄우고 완료 Event 를 돌려준다. 이미 생성 중이면 그 Event."""
    with _EDGE_TTS_LOCK:
        ev = _EDGE_TTS_INFLIGHT.get(key)
        if ev is not None:
            return ev
        ev = _gs_threading.Event()
        _EDGE_TTS_INFLIGHT[key] = ev

    def _run():
        try:
            _edge_tts_generate_blocking(text, voice, rate_pct, key)
        finally:
            with _EDGE_TTS_LOCK:
                _EDGE_TTS_INFLIGHT.pop(key, None)
            ev.set()

    _gs_threading.Thread(target=_run, daemon=True, name='edge-tts').start()
    return ev


def _edge_tts_b64(text, voice='ko-KR-InJoonNeural', rate_pct=0, wait=None):
    """edge-tts(Microsoft 신경망 TTS, 무료) mp3 → base64 문자열.

    캐시(메모리/디스크)에 있으면 즉시 반환. 없으면 백그라운드 생성을 시작하고
    최대 wait 초(기본 _EDGE_TTS_WAIT_BUDGET)만 기다린다. 그 안에 안 끝나면 None
    → 호출부가 브라우저 TTS로 폴백하고, 생성은 계속 진행되어 다음부터 캐시 히트.
    wait=0 이면 전혀 기다리지 않는다(워밍업용).
    """
    text = str(text).strip()
    if not text:
        return None
    key = _edge_tts_cache_key(text, voice, rate_pct)
    hit = _edge_tts_load_cached(key)
    if hit:
        return hit
    # 최근에 실패했으면 시도하지 않는다 — 문구마다 기다리며 멈추는 것을 막는다
    if time.monotonic() < _EDGE_TTS_FAIL_UNTIL:
        return None
    ev = _edge_tts_start_async(text, voice, rate_pct, key)
    wait = _EDGE_TTS_WAIT_BUDGET if wait is None else wait
    if wait > 0:
        ev.wait(timeout=wait)
        return _EDGE_TTS_MEM.get(key)
    return None


def _tts_num_to_kor(n):
    """1 → '일', 12 → '십이', 34 → '삼십사' (TTS 발음용 한글 숫자)."""
    _ones = {1: '일', 2: '이', 3: '삼', 4: '사', 5: '오', 6: '육', 7: '칠', 8: '팔', 9: '구', 10: '십'}
    n = int(n)
    if n in _ones:
        return _ones[n]
    if 10 < n <= 99:
        tens, ones = n // 10, n % 10
        t_str = ('이삼사오육칠팔구'[tens - 2] if tens >= 2 else '') + '십'
        return t_str + (_ones.get(ones, '') if ones else '')
    return str(n)


def _edge_tts_warmup_phrases():
    """워밍업할 고정 문구 — 피킹검증/입고분류/재고 탭이 실제로 말하는 문구 전부.

    입고분류는 출고박스 번호(1~60번)마다 "N번", "N번 완료. 포장하세요",
    "N번 미스캔 있음" 을 말하므로 그것까지 미리 만든다. 예전 워밍업은
    "N번박스"(피킹검증용)만 있어서 입고분류는 매 박스 첫 스캔이 느렸다.
    """
    phrases = [
        '입고완료', '확인', '없는 상품 입니다', '수량 초과', '재고 부족',
        '입고완료 재고 부족', '다른 박스 상품', '분류 완료', '수량을 입력하세요',
        '바코드를 스캔해 주세요', '보류',
        '재고완료', '다시 찍어주세요', '확인을 시작하세요',
        '검증확인이 완료되었습니다. 출고하세요',
    ]
    for n in range(1, 11):
        k = _tts_num_to_kor(n)
        phrases.append(f'{k}번박스')
        phrases.append(f'{k}번박스 재고 부족')
    for n in range(1, 61):
        k = _tts_num_to_kor(n)
        phrases.append(f'{k}번')
        phrases.append(f'{k}번 완료. 포장하세요')
        phrases.append(f'{k}번 미스캔 있음')
    return phrases


_EDGE_TTS_WARMED = set()   # (voice, rate_pct) — 프로세스당 한 번만 워밍업


def _edge_tts_warmup(voice, rate_pct=0):
    """자주 쓰는 문구를 백그라운드에서 미리 생성해 첫 스캔 지연 제거.
    서버 프로세스당 voice 별 한 번만 돈다 (세션마다 다시 돌면 디스크만 긁는다)."""
    with _EDGE_TTS_LOCK:
        if (voice, rate_pct) in _EDGE_TTS_WARMED:
            return
        _EDGE_TTS_WARMED.add((voice, rate_pct))
    phrases = _edge_tts_warmup_phrases()

    def _run():
        for p in phrases:
            key = _edge_tts_cache_key(p, voice, rate_pct)
            if _edge_tts_load_cached(key):
                continue
            if time.monotonic() < _EDGE_TTS_FAIL_UNTIL:
                # 서비스가 죽어있으면 나머지도 실패할 것이므로 잠시 쉬었다 이어간다
                time.sleep(_EDGE_TTS_FAIL_COOLDOWN)
            _edge_tts_generate_blocking(p, voice, rate_pct, key)

    _gs_threading.Thread(target=_run, daemon=True, name='edge-tts-warmup').start()


def _tts_ko_script(message, pitch=None, extra_rate=0.0):
    """TTS JS 코드 생성.
    - 한국어 프로필: 한국어 voice + 한국어 멘트 + 비프음
    - 영어 프로필 (남자): 영어 voice + 영어 짧은 멘트 (한→영 변환) + 비프음
      → 브라우저에 실제 남자 voice 있음 (Microsoft David/Mark, Daniel 등)
      → 영어 텍스트로 매칭되니까 "혀 꼬부라진" 문제 없음
    반환: <script> 태그 내부에 바로 넣을 JS 코드
    """
    lang = st.session_state.get('pick_tts_lang', 'ko')
    gender = st.session_state.get('pick_tts_gender', 'female')
    voice_hint = str(st.session_state.get('pick_tts_voice_hint', ''))
    base_rate = float(st.session_state.get('pick_tts_rate', 1.0))
    profile_extra_rate = float(st.session_state.get('pick_tts_extra_rate', 0.0))
    final_rate = max(0.5, min(2.0, base_rate + extra_rate + profile_extra_rate))
    if pitch is None:
        profile_pitch = st.session_state.get('pick_tts_pitch')
        pitch = float(profile_pitch) if profile_pitch is not None else (0.9 if gender == 'male' else 1.1)

    # 영어 voice 사용 시 한국어 → 영어 변환
    if lang == 'en':
        final_message = _ko_message_to_en(message)
        lang_code = 'en-US'
    else:
        final_message = str(message)
        lang_code = 'ko-KR'
    msg_esc = final_message.replace('\\', '\\\\').replace("'", "\\'").replace('\n', ' ')
    hint_esc = voice_hint.replace("'", "\\'")

    # 시그니처 비프음 (기본 OFF — 거슬린다는 피드백)
    # 켜고 싶으면 세션에 pick_tts_enable_beep=True 설정
    tone_js = _tts_signature_tone_js() if st.session_state.get('pick_tts_enable_beep', False) else ''
    tts_delay = 450 if tone_js else 80  # 비프 있으면 끝난 후, 없으면 바로

    # ── edge-tts 프로필 (한국어 남자 음성) ──
    # 브라우저 Web Speech에는 한국어 남자 voice가 없으므로
    # 서버에서 Microsoft 신경망 TTS(무료)로 mp3 생성 → 브라우저에서 재생
    edge_voice = st.session_state.get('pick_tts_edge_voice')
    if edge_voice:
        edge_rate_pct = int(round((final_rate - 1.0) * 100))
        b64 = _edge_tts_b64(str(message), voice=edge_voice, rate_pct=edge_rate_pct)
        if b64:
            return (
                tone_js +
                "try{window.speechSynthesis.cancel();"
                "var _a=new Audio('data:audio/mpeg;base64," + b64 + "');"
                "_a.volume=1.0;"
                + f"setTimeout(function(){{_a.play().catch(function(e){{}});}},{tts_delay});"
                "}catch(e){}"
            )
        # 생성 실패(네트워크 등) → 아래 브라우저 TTS로 폴백

    return (
        tone_js +
        "try{window.speechSynthesis.cancel();setTimeout(function(){"
        "var u=new SpeechSynthesisUtterance('" + msg_esc + "');"
        "u.lang='" + lang_code + "';"
        "u.rate=" + f"{final_rate:.2f}" + ";"
        "u.pitch=" + f"{pitch:.2f}" + ";"
        "u.volume=1.0;"
        "var lang='" + lang + "';"
        "var g='" + gender + "';"
        "var profileHint='" + hint_esc + "';"
        # 한국어 힌트
        "var fHko=['heami','yuna','sun-hi','sunhi','hyunjung','ji-min','jimin','seo-hyeon','seohyeon'];"
        "var mHko=['injoon','in-joon','minsu','min-su','jungmin','jung-min'];"
        # 영어 voice 힌트 (남자/여자)
        "var fHen=['zira','jenny','samantha','susan','victoria','female'];"
        "var mHen=['david','mark','guy','daniel','alex','fred','male','man'];"
        "var H,Vfilter;"
        "if(lang==='en'){"
        "  Vfilter=window.speechSynthesis.getVoices().filter(function(v){return (v.lang||'').toLowerCase().indexOf('en')===0;});"
        "  H=g==='male'?mHen:fHen;"
        "}else{"
        "  Vfilter=window.speechSynthesis.getVoices().filter(function(v){return (v.lang||'').toLowerCase().indexOf('ko')===0;});"
        "  H=g==='male'?mHko:fHko;"
        "}"
        "var V=window.speechSynthesis.getVoices();"
        "var k=null;"
        # 0차(최우선): 사용자가 고급 설정에서 직접 고른 voice
        "try{var pref=window.localStorage.getItem('__preferred_voice');"
        "if(pref){k=V.find(function(v){return v.name===pref;});}}catch(e){}"
        # 1차: 프로필에서 지정한 voice_hint 이름으로 해당 언어 voice 중 매칭
        "if(!k&&profileHint){"
        "for(var j=0;j<Vfilter.length;j++){"
        "if((Vfilter[j].name||'').toLowerCase().indexOf(profileHint)>=0){k=Vfilter[j];break;}"
        "}}"
        # 2차: 성별 힌트 리스트에서 매칭
        "if(!k){for(var i=0;i<H.length&&!k;i++){"
        "for(var j=0;j<Vfilter.length;j++){"
        "var v=Vfilter[j];"
        "if((v.name||'').toLowerCase().indexOf(H[i])>=0){k=v;break;}"
        "}}}"
        # 3차: 해당 언어 voice 아무거나
        "if(!k){k=Vfilter[0];}"
        "if(k)u.voice=k;"
        "window.speechSynthesis.speak(u);"
        + f"}},{tts_delay});}}catch(e){{}}"
    )


def render_scan_result_card(status, title, detail, barcode='', img_url='', img_size=120, height=210):
    """스캔 결과 카드를 components.html(iframe)로 렌더링.

    st.markdown은 DOMPurify가 HTML을 소독하면서 onerror / referrerpolicy 속성을
    제거함 → alicdn 등 Referer 차단 CDN 이미지가 X로 깨지고, 깨진 이미지를
    숨기지도 못함. iframe 안에서는 <meta name="referrer"> + onerror가 정상 동작.
    """
    from html import escape as _e
    from streamlit.components.v1 import html as _card_html
    palette = {
        'ok':       ('#d4edda', '#28a745', '#155724'),
        'over':     ('#fff3cd', '#ffc107', '#856404'),
        'warning':  ('#fff3cd', '#ffc107', '#856404'),
        'error':    ('#f8d7da', '#dc3545', '#721c24'),
        'shortage': ('#e2e3f1', '#6c63ff', '#383467'),
    }
    bg, border, fg = palette.get(status, palette['ok'])
    img_url = str(img_url or '').strip()
    if not img_url.lower().startswith(('http://', 'https://')):
        img_url = ''
    img_tag = (
        f'<img src="{_e(img_url, quote=True)}" loading="lazy" decoding="async" '
        f'referrerpolicy="no-referrer" '
        f'style="width:{img_size}px;height:{img_size}px;object-fit:contain;border-radius:8px;'
        f'background:rgba(255,255,255,0.95);margin-left:14px;flex-shrink:0;" '
        'onerror="this.style.display=\'none\'"/>'
    ) if img_url else ''
    bc_tag = (
        '<div style="margin-top:8px;">'
        '<span style="display:inline-block;background:#111827;color:#fff;'
        'font-family:ui-monospace,Consolas,monospace;font-size:1.5rem;font-weight:bold;'
        'letter-spacing:2px;padding:5px 14px;border-radius:8px;">'
        f'📊 {_e(str(barcode))}</span></div>'
    ) if barcode else ''
    _card_html(
        '<html><head><meta name="referrer" content="no-referrer">'
        '<style>body{margin:0;font-family:"Segoe UI",system-ui,sans-serif;}</style></head><body>'
        f'<div style="background:{bg};border-left:6px solid {border};color:{fg};'
        'padding:14px 18px;border-radius:8px;display:flex;align-items:center;'
        'justify-content:space-between;box-sizing:border-box;">'
        '<div style="flex:1;min-width:0;word-break:break-all;">'
        f'<strong style="font-size:1.25rem;">{_e(str(title))}</strong>'
        f'{bc_tag}'
        f'<div style="margin-top:5px;font-size:0.95rem;">{_e(str(detail))}</div>'
        '</div>'
        f'{img_tag}'
        '</div></body></html>',
        height=height,
    )


def pick_parse_box(box_str):
    import pandas as _pd
    if _pd.isna(box_str) or str(box_str).strip() == "":
        return {"기호": None, "박스": None, "수량": None, "상태": "알수없음"}
    box_str = str(box_str).strip()
    # 셀 전체가 부족 토큰일 때만 '부족' — '부족(-1)', '국내부족(-2)', 개수 없는 '부족'.
    # (개수 없는 '부족' 은 지시서 쪽 헬퍼(_RE_SHORTAGE_TOKEN)와 같은 기준으로 인식한다.)
    # '부족(-1),▲M7(1)' 처럼 박스와 섞인 셀은 아래에서 박스 토큰을 찾아 피킹가능으로 본다.
    # 예전에는 앞에서부터만(match) 봐서 부족이 먼저 적힌 셀은 통째로 부족 처리돼 피킹에서
    # 빠지고, 박스가 먼저 적힌 셀은 피킹가능이 되는 등 적힌 순서에 따라 결과가 달랐다.
    match = re.match(r"((?:국내)?부족)(?:\((-?\d+)\))?$", box_str)
    if match:
        qty = int(match.group(2)) if match.group(2) is not None else None
        return {"기호": match.group(1), "박스": None, "수량": qty, "상태": "부족"}
    # 영문+숫자 형식 (예: W1, M3) + 수량 — 셀 어디에 있든 첫 박스 토큰을 쓴다
    match = re.search(r"([●★■▲◆◇○□△▼♦♠♣♥☆※·]+)([A-Za-z]*\d+)\((\d+)\)", box_str)
    if match:
        return {"기호": match.group(1), "박스": match.group(2).upper(), "수량": int(match.group(3)), "상태": "피킹가능"}
    match = re.search(r"([●★■▲◆◇○□△▼♦♠♣♥☆※·]+)([A-Za-z]*\d+)", box_str)
    if match:
        return {"기호": match.group(1), "박스": match.group(2).upper(), "수량": None, "상태": "피킹가능"}
    match = re.match(r"(국내재고)\((\d+)\)", box_str)
    if match:
        return {"기호": match.group(1), "박스": None, "수량": int(match.group(2)), "상태": "피킹가능"}
    if box_str == "국내재고":
        return {"기호": "국내재고", "박스": None, "수량": None, "상태": "피킹가능"}
    return {"기호": box_str, "박스": None, "수량": None, "상태": "알수없음"}

def pick_clean_출고(df):
    import pandas as _pd
    df = df.copy()
    required = ["바코드", "상품명", "수량", "쉽먼트운송장번호"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        st.error(f"출고지시서 필수 컬럼 누락: {missing}")
        return None
    df["수량"] = _pd.to_numeric(df["수량"], errors="coerce").fillna(0).astype(int)

    # 쉽먼트운송장번호 정규화:
    # - 과학표기법(4.62E+11) → 462139010304
    # - float 표기(462139010304.0) → 462139010304
    # - 공백/None 제거
    def _norm_ship(v):
        s = str(v if v is not None else '').strip()
        if not s or s.lower() in ('nan', 'none'):
            return ''
        # 과학표기법 또는 소수점 포함 → int 변환
        if 'E' in s or 'e' in s or '.' in s:
            try:
                return str(int(float(s)))
            except (ValueError, TypeError):
                pass
        return s
    df["쉽먼트운송장번호"] = df["쉽먼트운송장번호"].apply(_norm_ship)
    df["바코드"] = df["바코드"].astype(str).str.strip()
    # 확인 수량(L열) 보존 - 시트 재로드 시 진행 상태 복원용
    check_col = None
    for c in df.columns:
        c_norm = str(c).strip().replace(' ', '')
        if c_norm in ('확인수량', '확인량'):
            check_col = c
            break
    if check_col is not None:
        df["확인수량"] = _pd.to_numeric(df[check_col], errors="coerce").fillna(0).astype(int)
    elif len(df.columns) >= 12:
        # 헤더가 비어있거나 다른 이름이면 L열(12번째) 위치로 시도
        df["확인수량"] = _pd.to_numeric(df.iloc[:, 11], errors="coerce").fillna(0).astype(int)
    else:
        df["확인수량"] = 0

    # 이미지 URL 추출: 헤더명 우선, 없으면 N열(14번째)로 폴백
    img_col = None
    for c in df.columns:
        c_norm = str(c).strip().replace(' ', '').lower()
        if c_norm in ('이미지', '이미지url', 'imageurl', 'image_url', '이미지링크', '썸네일', 'thumbnail'):
            img_col = c
            break
    if img_col is not None:
        df["이미지URL"] = df[img_col].astype(str).str.strip()
    elif len(df.columns) >= 14:
        df["이미지URL"] = df.iloc[:, 13].astype(str).str.strip()
    else:
        df["이미지URL"] = ""
    # 빈/유효하지 않은 값 정리
    df["이미지URL"] = df["이미지URL"].apply(
        lambda v: '' if str(v).strip().lower() in ('', 'nan', 'none', '-') else str(v).strip()
    )
    if "박스번호" in df.columns:
        parsed = df["박스번호"].apply(pick_parse_box)
        df["회차기호"] = parsed.apply(lambda x: x["기호"])
        df["박스넘버"] = parsed.apply(lambda x: x["박스"])
        df["박스내수량"] = parsed.apply(lambda x: x["수량"])
        df["피킹상태"] = parsed.apply(lambda x: x["상태"])
    return df

def pick_clean_배대지(df):
    import pandas as _pd
    df = df.copy()
    if "바코드" not in df.columns:
        st.error("배대지 시트에 '바코드' 컬럼이 없습니다")
        return None
    df["바코드"] = df["바코드"].astype(str).str.strip()
    if "수량" in df.columns:
        df["수량"] = _pd.to_numeric(df["수량"], errors="coerce").fillna(0).astype(int)
    if "배대지주문수량" in df.columns:
        df["배대지주문수량"] = _pd.to_numeric(df["배대지주문수량"], errors="coerce").fillna(0).astype(int)
    if "박스번호" in df.columns:
        parsed = df["박스번호"].apply(pick_parse_box)
        df["회차기호"] = parsed.apply(lambda x: x["기호"])
        df["박스넘버"] = parsed.apply(lambda x: x["박스"])
    return df

def pick_init_session():
    defaults = {
        "pick_df_출고": None, "pick_df_배대지": None,
        "pick_selected_shipment": None, "pick_selected_shipments": [], "pick_show_add_input": False, "pick_picking_state": {},
        "pick_inventory_state": {}, "pick_scan_log": [],
        "pick_last_scan_result": None, "pick_scan_counter": 0,
        "pick_completed_shipments": set(), "pick_shortage_items": [],
        "pick_bulk_complete_snapshot": {},
        "pick_data_loaded": False, "pick_gsheet_client": None,
        "pick_use_gsheet": False,
        "pick_sheet_url_출고": "", "pick_sheet_tab_출고": "",
        "pick_sheet_url_배대지": "", "pick_sheet_tab_배대지": "",
    }
    for key, val in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = val

pick_init_session()

def pick_load_all_data(url_출고, tab_출고, url_배대지="", tab_배대지=""):
    """구글 시트에서 출고지시서/배대지 데이터 로드.

    ⚠️ 새 탭 로드가 실패해도 이전 세션의 pick_df_출고가 남아있으면
    성공으로 처리되던 버그가 있었음 → 사용자가 탭을 바꿨는데 이전 탭
    데이터로 작업하게 됨 (예: 출고확인_2 연결했는데 출고확인 데이터 사용).
    이제 이번 로드가 실제 성공했을 때만 True + 상태 갱신, 실패 시 이전
    데이터를 지우고 False 반환.
    """
    import pandas as _pd
    client = get_gsheet_client()
    if not client:
        return False
    st.session_state.pick_gsheet_client = client
    st.session_state.pick_use_gsheet = True
    # 시트 재로드 시 M열 캐시 초기화 (발주 변동 반영)
    for _k in list(st.session_state.keys()):
        if _k.startswith("_pick_existing_box_") or _k.startswith("_pick_box_written_"):
            del st.session_state[_k]

    url_출고 = str(url_출고 or '').strip()
    tab_출고 = str(tab_출고 or '').strip()   # 탭명 끝 공백 → 로드 실패 방지
    url_배대지 = str(url_배대지 or '').strip()
    tab_배대지 = str(tab_배대지 or '').strip()

    if not (url_출고 and tab_출고):
        return False

    def _fail(msg=None):
        # 이전 탭 데이터로 오인 작업하지 않도록 로드 상태 초기화
        if msg:
            st.error(msg)
        st.session_state.pick_df_출고 = None
        st.session_state.pick_data_loaded = False
        st.session_state.pick_sheet_tab_출고 = ""
        return False

    df_출고 = pick_load_sheet_as_df(client, url_출고, tab_출고)
    if df_출고 is None:
        return _fail()  # pick_load_sheet_as_df가 이미 에러 표시
    if df_출고.empty:
        return _fail(f"'{tab_출고}' 탭에 데이터가 없습니다 (헤더+1행 이상 필요)")
    _cleaned = pick_clean_출고(df_출고)
    if _cleaned is None or _cleaned.empty:
        return _fail(f"'{tab_출고}' 탭에서 유효한 데이터를 읽지 못했습니다 (필수 컬럼 확인)")

    st.session_state.pick_df_출고 = _cleaned
    st.session_state.pick_sheet_url_출고 = url_출고
    st.session_state.pick_sheet_tab_출고 = tab_출고
    # 데이터 세대 번호 — 입고분류 sort_state 가 '데이터가 바뀌었는지' 를 이걸로 판단한다.
    # (예전에는 id(df) 로 비교했는데, id 는 GC 뒤 재사용될 수 있어 새 데이터를 옛 상태로
    #  오인하거나, CSV 모드에서는 rerun 마다 df 가 새로 만들어져 진행 상태가 매번 날아갔다.)
    st.session_state['pick_data_ver'] = st.session_state.get('pick_data_ver', 0) + 1

    # 배대지 입고 로드 (선택 — 실패해도 출고 로드에는 영향 없음)
    if url_배대지 and tab_배대지:
        df_배대지 = pick_load_sheet_as_df(client, url_배대지, tab_배대지)
        if df_배대지 is not None and not df_배대지.empty:
            st.session_state.pick_df_배대지 = pick_clean_배대지(df_배대지)
            st.session_state.pick_sheet_url_배대지 = url_배대지
            st.session_state.pick_sheet_tab_배대지 = tab_배대지

    st.session_state.pick_data_loaded = True
    return True

def pick_init_inventory():
    df = st.session_state.pick_df_배대지
    if df is None or df.empty:
        return
    inventory = {}
    for _, row in df.iterrows():
        barcode = row["바코드"]
        symbol = row.get("회차기호", "기타")
        qty = row.get("수량", 0)
        key = (symbol, barcode)
        inventory[key] = inventory.get(key, 0) + qty
    st.session_state.pick_inventory_state = inventory

def pick_init_picking(shipment_ids):
    """단일 또는 다중 쉽먼트 ID를 받아 피킹 초기화.
    다중일 경우 각 바코드별로 어느 쉽먼트(박스)에 속하는지 추적.
    박스 라벨은 시트 M열(출고박스번호) 값 우선 사용, 없으면 인덱스 기반 fallback."""
    if isinstance(shipment_ids, str):
        shipment_ids = [shipment_ids]

    df = st.session_state.pick_df_출고
    picking = {}
    shortage_items = []

    # 시트 M열(출고박스번호) 매핑 조회 — 세션 → 캐시 → 시트 직접 읽기 순
    _box_map = dict(st.session_state.get('pick_ship_to_box') or {})
    if not _box_map:
        _url = st.session_state.get('pick_sheet_url_출고', '')
        _tab = st.session_state.get('pick_sheet_tab_출고', '')
        if _url and _tab:
            _cache_key = f"_pick_existing_box_{_url}_{_tab}"
            _box_map = dict(st.session_state.get(_cache_key) or {})
    if (not _box_map
            and st.session_state.get('pick_use_gsheet')
            and st.session_state.get('pick_gsheet_client')):
        try:
            _read = pick_read_box_numbers(
                st.session_state.pick_gsheet_client,
                st.session_state.get('pick_sheet_url_출고', ''),
                st.session_state.get('pick_sheet_tab_출고', ''),
            )
            if _read:
                _box_map = _read
                # 세션에도 저장해서 다음부터 재사용
                st.session_state['pick_ship_to_box'] = _box_map
                _url = st.session_state.get('pick_sheet_url_출고', '')
                _tab = st.session_state.get('pick_sheet_tab_출고', '')
                if _url and _tab:
                    st.session_state[f"_pick_existing_box_{_url}_{_tab}"] = _box_map
        except Exception:
            pass

    for ship_idx, shipment_id in enumerate(shipment_ids, start=1):
        shipment_df = df[df["쉽먼트운송장번호"] == shipment_id]
        if shipment_df.empty:
            st.error(f"쉽먼트 {shipment_id}를 찾을 수 없습니다")
            continue
        # M열 출고박스번호 우선, 없으면 인덱스 기반 fallback
        _m_box = _box_map.get(str(shipment_id).strip())
        if _m_box:
            ship_label = f"{_m_box}번박스"
        else:
            ship_label = f"{ship_idx}번박스"
        for _, row in shipment_df.iterrows():
            bc = row["바코드"]
            symbol = row.get("회차기호", "")
            qty = row["수량"]
            pick_status = row.get("피킹상태", "피킹가능")
            if pick_status == "부족":
                shortage_items.append({
                    "바코드": bc, "상품명": row["상품명"],
                    "부족수량": abs(row.get("박스내수량", 0) or 0),
                    "박스번호": row.get("박스번호", ""),
                    "쉽먼트박스": ship_label,
                })
                continue
            # H열 수량은 주문 수량이라 '부족(-N)' 분량까지 들어있다. '▲M7(1),부족(-1)' 처럼
            # 박스와 부족이 섞인 행은 실제로 담기는 양만 스캔 대상으로 잡아야 검증 완료가
            # 가능하다 (출고지시서·입고분류와 같은 기준). 부족이 없는 행은 그대로 수량이다.
            _short = item_shortage_qty({'boxNumber': row.get('박스번호', ''), 'quantity': qty})
            if _short > 0:
                qty = max(0, int(qty) - _short)
                if qty <= 0:
                    # 박스 표기는 있지만 담길 양이 없는 행 → 부족 목록으로
                    shortage_items.append({
                        "바코드": bc, "상품명": row["상품명"],
                        "부족수량": _short,
                        "박스번호": row.get("박스번호", ""),
                        "쉽먼트박스": ship_label,
                    })
                    continue
            if bc in picking:
                picking[bc]["필요수량"] += qty
                # 같은 바코드가 여러 쉽먼트에 있으면 박스 라벨을 합침
                if ship_label not in picking[bc]["쉽먼트박스목록"]:
                    picking[bc]["쉽먼트박스목록"].append(ship_label)
                # 송장별 필요수량 추적
                picking[bc]["송장별수량"][ship_label] = picking[bc]["송장별수량"].get(ship_label, 0) + qty
            else:
                inv_key = (symbol, bc)
                inv_qty = st.session_state.pick_inventory_state.get(inv_key, None)
                picking[bc] = {
                    "상품명": row["상품명"], "필요수량": qty, "스캔수량": 0,
                    "회차기호": symbol if symbol else "N/A",
                    "박스번호": row.get("박스번호", ""), "박스넘버": row.get("박스넘버", ""),
                    "박스내수량": row.get("박스내수량", None), "배대지잔여": inv_qty,
                    "SKU_ID": row.get("SKU ID", ""), "물류센터": row.get("물류센터(FC)", ""),
                    "이미지URL": str(row.get("이미지URL", "") or "").strip(),
                    "쉽먼트박스목록": [ship_label],
                    "송장별수량": {ship_label: qty},
                }

    st.session_state.pick_picking_state = picking
    st.session_state.pick_shortage_items = shortage_items
    st.session_state.pick_selected_shipment = " + ".join(shipment_ids) if len(shipment_ids) > 1 else shipment_ids[0]
    st.session_state.pick_selected_shipments = shipment_ids
    st.session_state.pick_scan_log = []
    st.session_state.pick_last_scan_result = None
    st.session_state.pick_scan_counter = 0
    # 다량 모드 상태 초기화 (이전 쉽먼트의 모드가 끌려오지 않도록)
    st.session_state.pick_next_qty = 1
    st.session_state.pick_qty_input_mode = False

def pick_process_scan(barcode, qty=1):
    barcode = barcode.strip()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    state = st.session_state.pick_picking_state
    inventory = st.session_state.pick_inventory_state

    # #MULTI 트리거 → 수량 입력 모드 진입
    if barcode.upper() == "#MULTI":
        st.session_state.pick_qty_input_mode = True
        st.session_state.pick_next_qty = 1
        result = {"status": "multi_trigger", "message": "🔢 다량 입력 모드",
                  "detail": "수량을 입력한 후 상품 바코드를 스캔하세요",
                  "barcode": barcode, "상품명": "", "시간": now}
        st.session_state.pick_last_scan_result = result
        st.session_state.pick_scan_counter += 1
        return result

    if barcode not in state:
        df = st.session_state.pick_df_출고
        hint = ""
        if df is not None:
            match = df[df["바코드"] == barcode]
            if not match.empty:
                name = match["상품명"].iloc[0][:25]
                others = match["쉽먼트운송장번호"].unique()[:3]
                hint = f" → [{name}] 다른 쉽먼트에 있음: {', '.join(s[-6:] for s in others)}"
            else:
                hint = " → 출고지시서에 없는 바코드"
        result = {"status": "error", "message": "🚨 오피킹! 이 쉽먼트에 없는 바코드",
                  "detail": f"{barcode}{hint}", "barcode": barcode, "상품명": "", "시간": now}
        st.session_state.pick_scan_log.append(result)
        st.session_state.pick_last_scan_result = result
        st.session_state.pick_scan_counter += 1
        return result

    item = state[barcode]
    qty = max(1, int(qty))
    needed = item["필요수량"]
    prev_scanned = item["스캔수량"]

    # 이미 다 찬 상태에서 추가 스캔 → 전량 초과
    if prev_scanned >= needed:
        item["스캔수량"] += qty
        box_label = ""
        if "송장별수량" in item and item["송장별수량"]:
            box_label = list(item["송장별수량"].keys())[-1]
        result = {"status": "over", "message": f"⚠️ 수량 초과! {item['상품명'][:35]}",
                  "detail": f"필요 {needed}개 전부 스캔됨 (+{qty} 추가)",
                  "barcode": barcode, "상품명": item["상품명"], "시간": now, "박스": box_label}
        st.session_state.pick_scan_log.append(result)
        st.session_state.pick_last_scan_result = result
        st.session_state.pick_scan_counter += 1
        # 다량 모드 1회성 해제
        st.session_state.pick_next_qty = 1
        st.session_state.pick_qty_input_mode = False
        return result

    # 정상 처리: qty 만큼 증가 (초과분 포함)
    item["스캔수량"] += qty
    over_qty = max(0, item["스캔수량"] - needed)
    effective_inc = qty - over_qty  # 실제 유효 스캔 수

    # 박스 라벨: 마지막 유효 유닛 기준
    box_label = ""
    if "송장별수량" in item:
        cumulative = 0
        last_valid_scan = min(item["스캔수량"], needed)
        for ship_label, ship_qty in item["송장별수량"].items():
            cumulative += ship_qty
            if last_valid_scan <= cumulative:
                box_label = ship_label
                break

    # 재고 차감 (유효분만)
    symbol = item["회차기호"]
    inv_key = (symbol, barcode)
    shortage_warning = ""
    sheet_decrement = 0
    if inv_key in inventory and effective_inc > 0:
        available = inventory[inv_key]
        if available >= effective_inc:
            inventory[inv_key] -= effective_inc
            item["배대지잔여"] = inventory[inv_key]
            sheet_decrement = effective_inc
        else:
            sheet_decrement = available
            inventory[inv_key] = 0
            item["배대지잔여"] = 0
            shortage_warning = f" | ⚠ {symbol}회차 배대지 재고 소진!"
    elif inv_key in inventory:
        item["배대지잔여"] = inventory[inv_key]

    remaining = max(0, needed - item["스캔수량"])
    box_msg = f" → {box_label}" if box_label else ""
    qty_label = f" ×{qty}" if qty > 1 else ""
    if over_qty > 0:
        status = "over"
        message = f"⚠️ 수량 초과! {item['상품명'][:35]}{box_msg}"
        detail = f"스캔 {item['스캔수량']}/{needed} (+{over_qty} 초과){shortage_warning}"
    else:
        status = "shortage" if shortage_warning else "ok"
        message = f"✅ {item['상품명'][:35]}{box_msg}{qty_label}"
        detail = f"스캔 {item['스캔수량']}/{needed} (남은: {remaining}){shortage_warning}"

    result = {
        "status": status,
        "message": message,
        "detail": detail,
        "barcode": barcode, "상품명": item["상품명"], "시간": now, "박스": box_label,
        "처리수량": qty,
    }
    st.session_state.pick_scan_log.append(result)
    st.session_state.pick_last_scan_result = result
    st.session_state.pick_scan_counter += 1

    # 다량 모드는 1회성 → 1개 모드로 복귀
    st.session_state.pick_next_qty = 1
    st.session_state.pick_qty_input_mode = False

    if st.session_state.pick_use_gsheet and st.session_state.pick_gsheet_client:
        # 시트 쓰기는 메모리 큐에 넣고 즉시 반환한다 (네트워크 대기 없음).
        # 실제 전송은 플러셔 스레드가 2초마다 batch로 모아서 보낸다.
        # → 스캔 속도를 위해 스레드를 띄울 필요가 없어졌고, 백그라운드
        #   스레드에서 st.session_state를 읽던 위험도 사라졌다.
        _client = st.session_state.pick_gsheet_client
        log_row = [now, st.session_state.pick_selected_shipment or "", barcode,
                   item["상품명"][:40], result["status"], item["스캔수량"],
                   item["필요수량"], item.get("회차기호",""), item.get("박스번호","")]
        try:
            pick_append_log(_client, st.session_state.pick_sheet_url_출고, log_row)
            if (result["status"] in ("ok", "shortage") and sheet_decrement > 0
                    and st.session_state.pick_sheet_url_배대지
                    and st.session_state.pick_sheet_tab_배대지):
                pick_update_sheet_inventory(
                    _client,
                    st.session_state.pick_sheet_url_배대지,
                    st.session_state.pick_sheet_tab_배대지,
                    barcode, decrement=sheet_decrement,
                    # 배대지 시트 매칭용 박스번호 (배대지박스 식별자: M1/W3 등)
                    box_number=str(item.get("박스넘버") or '').strip().upper(),
                )
        except Exception:
            pass
    return result

def pick_get_progress():
    state = st.session_state.pick_picking_state
    if not state:
        return {"total":0,"scanned":0,"skus":0,"done_skus":0,"pct":0.0,"is_complete":False,"over":0,"shortage":0}
    total = sum(v["필요수량"] for v in state.values())
    scanned = sum(min(v["스캔수량"], v["필요수량"]) for v in state.values())
    over = sum(max(0, v["스캔수량"] - v["필요수량"]) for v in state.values())
    skus = len(state)
    done_skus = sum(1 for v in state.values() if v["스캔수량"] >= v["필요수량"])
    shortage = sum(1 for v in state.values()
                   if v.get("배대지잔여") is not None and v["배대지잔여"] == 0 and v["스캔수량"] < v["필요수량"])
    pct = scanned / total if total > 0 else 0
    return {"total":total,"scanned":scanned,"skus":skus,"done_skus":done_skus,
            "pct":pct,"is_complete":scanned>=total,"over":over,"shortage":shortage}

# ══════════════════════════════════════════════════════
# ── 입고 검수 (Inbound Inspection) ─────────────────────
#   입고예정 리스트(엑셀/CSV) 대비 실제 입고 바코드를 스캔해
#   과/부족/오입고를 자동 감지한다. (출고 피킹의 역방향)
# ══════════════════════════════════════════════════════


TAB_LABELS = {
    'small': '📦 소형 라벨',
    'large': '📋 대형 라벨 (90도 회전)',
    'reprint': '🔄 쉽먼트 재출력',
    'pick': '📦 피킹 & 분류',
    'stock': '📥 재고 확인',
    'notice': '📝 발주중단 공문',
    'pochg': '🚚 발주 변경 요청',
}
# 사용자가 설정에서 정한 순서대로 보여준다. 설정 · 관리자 탭은 항상 맨 뒤.
_tab_keys = kit_config.tab_order(list(TAB_LABELS))
_extra = ['settings'] + (['admin'] if kit_ui.is_admin() else [])
_extra_labels = {'settings': '⚙️ 설정', 'admin': '👑 계정 관리'}
_tabs = dict(zip(_tab_keys + _extra,
                 st.tabs([TAB_LABELS[k] for k in _tab_keys] + [_extra_labels[k] for k in _extra])))
tab1, tab2, tab7, tab8 = _tabs['small'], _tabs['large'], _tabs['reprint'], _tabs['pick']
tab_stock, tab5, tab_pochg = _tabs['stock'], _tabs['notice'], _tabs['pochg']

with _tabs['settings']:
    kit_ui.render_settings(TAB_LABELS, get_gsheet_client)
if 'admin' in _tabs:
    with _tabs['admin']:
        kit_ui.render_admin()

# ── 소형 탭 ────────────────────────────────────────────
with tab1:
    st.subheader('📋 열 번호 설정')
    st.caption('A=1, B=2, C=3, D=4 ... L=12, M=13')
    c1,c2 = st.columns(2)
    with c1:
        s_col_name     = st.number_input('상품명 열',     min_value=1, max_value=50, value=kit_config.label_col('s', 'name'),  key='s_name')
        s_col_barcode  = st.number_input('바코드 열',     min_value=1, max_value=50, value=kit_config.label_col('s', 'barcode'), key='s_barcode')
        s_col_material = st.number_input('재질 열',       min_value=1, max_value=50, value=kit_config.label_col('s', 'material'), key='s_material')
    with c2:
        s_col_insert   = st.number_input('이미지 삽입 열', min_value=1, max_value=50, value=kit_config.label_col('s', 'insert'), key='s_insert')
        s_start_row    = st.number_input('시작 행',       min_value=1, max_value=10,  value=kit_config.label_col('s', 'startrow'),  key='s_startrow')

    st.divider()
    st.subheader('✍️ 고정 문구')
    s_origin = st.text_input('제조국',   value=kit_config.label_text('s_origin'), key='s_origin')
    s_age    = st.text_input('사용연령', value=kit_config.label_text('s_age'),    key='s_age')
    if st.button('💾 이 문구 · 열 번호를 기본값으로 저장', key='s_save_defaults'):
        kit_config.save({
            'label_s_origin': s_origin, 'label_s_age': s_age,
            'label_s_name': s_col_name, 'label_s_barcode': s_col_barcode,
            'label_s_material': s_col_material, 'label_s_insert': s_col_insert,
            'label_s_startrow': s_start_row,
        })
        st.success('저장했습니다. 다음부터 이 값으로 나갑니다.')

    st.divider()
    s_file = st.file_uploader('📂 엑셀 파일 업로드', type=['xlsx'], key='s_file')

    if st.button('🚀 소형 라벨 생성 시작', type='primary', key='s_btn'):
        if not s_file:
            st.warning('⚠️ 엑셀 파일을 먼저 업로드해주세요!')
        else:
            with st.spinner('처리 중...'):
                settings={
                    'col_name':s_col_name,'col_barcode':s_col_barcode,
                    'col_material':s_col_material,'col_insert':s_col_insert,
                    'start_row':s_start_row,'origin':s_origin,'age':s_age,
                    'insert_w':244,'insert_h':157,'row_height':120,'col_width':38
                }
                output, ok, errors = process_excel(s_file, '소형', settings)
            if errors:
                for e in errors: st.error(e)
            fname = s_file.name.replace('.xlsx','_완성.xlsx')
            st.download_button('⬇️ 완성 파일 다운로드', output, file_name=fname,
                             mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # 배대지 부착용 폼텍 3102 라벨지 (B열 상품사진 · F/G 옵션 · J 수량)
    # 바코드 라벨은 위 '소형 라벨 생성'과 같은 create_small 을 그대로 쓴다 — 디자인 동일.
    def _small_label_image(item, px_w, px_h):
        return create_small(item.name, item.barcode, item.material, s_origin, s_age,
                            canvas_size=(px_w, px_h) if px_w and px_h else None)

    render_label_sheet_section(
        'ls_small', '3102', '소형', s_file,
        dict(col_name=s_col_name, col_barcode=s_col_barcode,
             col_material=s_col_material, start_row=s_start_row),
        [s_origin, s_age],
        dict(image=2, opt1=6, opt2=7, qty=10),
        _small_label_image,
    )

# ── 대형 탭 ────────────────────────────────────────────
with tab2:
    st.subheader('📋 열 번호 설정')
    st.caption('A=1, B=2, C=3, D=4 ... L=12, M=13')
    c1,c2 = st.columns(2)
    with c1:
        l_col_name     = st.number_input('상품명 열',     min_value=1, max_value=50, value=kit_config.label_col('l', 'name'),  key='l_name')
        l_col_barcode  = st.number_input('바코드 열',     min_value=1, max_value=50, value=kit_config.label_col('l', 'barcode'), key='l_barcode')
        l_col_material = st.number_input('재질 열',       min_value=1, max_value=50, value=kit_config.label_col('l', 'material'), key='l_material')
    with c2:
        l_col_insert   = st.number_input('이미지 삽입 열', min_value=1, max_value=50, value=kit_config.label_col('l', 'insert'), key='l_insert')
        l_start_row    = st.number_input('시작 행',       min_value=1, max_value=10,  value=kit_config.label_col('l', 'startrow'),  key='l_startrow')

    st.divider()
    st.subheader('✍️ 고정 문구')
    l_caution = st.text_area('취급주의', value=kit_config.label_text('l_caution'), height=80, key='l_caution')
    l_addr    = st.text_input('주소/전화', value=kit_config.label_text('l_addr'),   key='l_addr')
    l_origin  = st.text_input('제조국',   value=kit_config.label_text('l_origin'), key='l_origin')
    l_age     = st.text_input('사용연령', value=kit_config.label_text('l_age'),    key='l_age')
    if st.button('💾 이 문구 · 열 번호를 기본값으로 저장', key='l_save_defaults'):
        kit_config.save({
            'label_l_caution': l_caution, 'label_l_addr': l_addr,
            'label_l_origin': l_origin, 'label_l_age': l_age,
            'label_l_name': l_col_name, 'label_l_barcode': l_col_barcode,
            'label_l_material': l_col_material, 'label_l_insert': l_col_insert,
            'label_l_startrow': l_start_row,
        })
        st.success('저장했습니다. 다음부터 이 값으로 나갑니다.')

    st.divider()
    l_file = st.file_uploader('📂 엑셀 파일 업로드', type=['xlsx'], key='l_file')

    if st.button('🚀 대형 라벨 생성 시작', type='primary', key='l_btn'):
        if not l_file:
            st.warning('⚠️ 엑셀 파일을 먼저 업로드해주세요!')
        else:
            with st.spinner('처리 중...'):
                settings={
                    'col_name':l_col_name,'col_barcode':l_col_barcode,
                    'col_material':l_col_material,'col_insert':l_col_insert,
                    'start_row':l_start_row,
                    'fix_list':[l_caution,l_addr,l_origin,l_age],
                    'insert_w':298,'insert_h':208,'row_height':160,'col_width':58
                }
                output, ok, errors = process_excel(l_file, '대형', settings)
            if errors:
                for e in errors: st.error(e)
            fname = l_file.name.replace('.xlsx','_완성.xlsx')
            st.download_button('⬇️ 완성 파일 다운로드', output, file_name=fname,
                             mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # 배대지 부착용 폼텍 3218 라벨지 (B열 상품사진 · F/G 옵션 · J 수량)
    # create_large 는 마지막에 90도 회전하므로, 원하는 최종 크기의 가로/세로를 뒤집어 넘긴다.
    def _large_label_image(item, px_w, px_h):
        return create_large(item.name, item.barcode, item.material,
                            [l_caution, l_addr, l_origin, l_age],
                            canvas_size=(px_h, px_w) if px_w and px_h else None)

    render_label_sheet_section(
        'ls_large', '3218', '대형', l_file,
        dict(col_name=l_col_name, col_barcode=l_col_barcode,
             col_material=l_col_material, start_row=l_start_row),
        [l_caution, l_addr, l_origin, l_age],
        dict(image=2, opt1=6, opt2=7, qty=10),
        _large_label_image,
    )

# ── 재고 확인(반출 재고 채우기) 탭 ──────────────────────────
with tab_stock:
    st.header('📥 반출 재고 채우기')
    st.caption('위치를 먼저 지정하고 바코드를 스캔하면 재고 확인 시트의 X열(재고)에 누적, V열에 위치 기록 · 시트는 ⚙️ 설정 탭에서 지정')

    _STOCK_SHEET_URL = kit_config.stock_sheet_url()
    _STOCK_SHEET_TAB = kit_config.stock_tab()

    # 세션 초기화
    for _k, _v in [
        ('stock_location', ''),
        ('stock_scan_log', []),
        ('stock_next_qty', 1),
        ('stock_qty_input_mode', False),
        ('stock_qty_pending', False),   # 다량 모드에서 아직 수량을 안 받은 상태
        ('stock_scan_counter', 0),
        ('stock_last_result', None),
    ]:
        if _k not in st.session_state:
            st.session_state[_k] = _v

    # ── 시트 연결 ──
    _sclient = st.session_state.get('pick_gsheet_client') or get_gsheet_client()
    _stock_ready = _sclient is not None and bool(_STOCK_SHEET_URL)
    if not _stock_ready:
        if not _STOCK_SHEET_URL:
            st.error('❌ ⚙️ 설정 탭에서 재고 확인 시트 주소를 먼저 저장하세요')
        else:
            st.error('❌ Google Sheets 연결 실패 — ⚙️ 설정 탭의 [시트 연결 테스트]로 확인하세요')
    else:
        st.session_state['pick_gsheet_client'] = _sclient

        _sinfo_c1, _sinfo_c2, _sinfo_c3 = st.columns([3, 1, 1])
        with _sinfo_c1:
            st.text_input('구글 시트', value=_STOCK_SHEET_URL, disabled=True, key='stock_url_display')
        with _sinfo_c2:
            st.text_input('탭 이름', value=_STOCK_SHEET_TAB, disabled=True, key='stock_tab_display')
        with _sinfo_c3:
            st.markdown('<br>', unsafe_allow_html=True)
            if st.button('🔄 시트 새로고침', key='stock_reload', use_container_width=True):
                gs_flush_pending()   # 대기 중인 쓰기 먼저 반영 후 다시 읽기
                _gs_invalidate(_STOCK_SHEET_URL, _STOCK_SHEET_TAB)
                st.rerun()

        # ── 1단계: 위치 입력 ──
        st.divider()
        st.markdown('### 📍 1단계: 담을 위치 지정')
        _loc_c1, _loc_c2, _loc_c3 = st.columns([3, 1, 1])
        with _loc_c1:
            _loc_input = st.text_input(
                '위치/박스 (예: G박스, A-1구역)',
                value=st.session_state.stock_location,
                key='stock_location_input',
                placeholder='예: G박스, A-1구역',
            )
        with _loc_c2:
            if st.button('✅ 위치 설정', use_container_width=True, key='stock_set_loc'):
                _cleaned = _loc_input.strip()
                if _cleaned:
                    st.session_state.stock_location = _cleaned
                    st.success(f'위치: {_cleaned}')
                    st.rerun()
                else:
                    st.error('위치를 입력하세요')
        with _loc_c3:
            if st.button('🔄 위치 변경', use_container_width=True, key='stock_reset_loc'):
                st.session_state.stock_location = ''
                st.rerun()

        # ── 위치 일괄 변경 ──
        with st.expander('🔁 위치 일괄 변경 (예: G박스 → H박스)', expanded=False):
            st.caption('기존 위치에 있는 모든 상품의 V열 위치를 새 위치로 일괄 교체합니다. 같은 행에 다른 위치가 같이 있으면 그 위치는 유지됨.')
            _ch_c1, _ch_c2, _ch_c3 = st.columns([2, 2, 1])
            with _ch_c1:
                _old_loc = st.text_input(
                    '기존 위치',
                    key='stock_change_old',
                    placeholder='예: G박스',
                )
            with _ch_c2:
                _new_loc = st.text_input(
                    '변경할 위치 (비우면 위치 제거)',
                    key='stock_change_new',
                    placeholder='예: H박스',
                )
            with _ch_c3:
                st.markdown('<br>', unsafe_allow_html=True)
                _ch_preview = st.button('🔍 미리보기', use_container_width=True, key='stock_change_preview_btn')

            if _ch_preview:
                with st.spinner('대상 검색 중...'):
                    _prv = stock_change_location(
                        _sclient, _STOCK_SHEET_URL, _STOCK_SHEET_TAB,
                        _old_loc.strip(), _new_loc.strip(), preview_only=True,
                    )
                st.session_state['_stock_change_preview'] = _prv

            _prv = st.session_state.get('_stock_change_preview')
            if _prv:
                if not _prv.get('ok'):
                    st.error(_prv.get('error', '오류'))
                elif _prv['total_changed'] == 0:
                    st.warning(f"기존 위치 '{_old_loc.strip()}'에 해당하는 상품이 없습니다")
                else:
                    st.success(f'✅ 변경 대상 {_prv["total_changed"]}건 발견')
                    import pandas as _pds_ch
                    _ch_rows = [{
                        '행': c['row'], '바코드': c['barcode'],
                        '상품명': c['name'], '기존 위치': c['old'], '변경 후': c['new'],
                    } for c in _prv['changes']]
                    st.dataframe(_pds_ch.DataFrame(_ch_rows),
                                 use_container_width=True, hide_index=True)
                    if st.button(f'✅ 위 {_prv["total_changed"]}건 일괄 변경 실행',
                                 type='primary', key='stock_change_exec',
                                 use_container_width=True):
                        with st.spinner('시트 업데이트 중...'):
                            _res = stock_change_location(
                                _sclient, _STOCK_SHEET_URL, _STOCK_SHEET_TAB,
                                _old_loc.strip(), _new_loc.strip(), preview_only=False,
                            )
                        if _res.get('ok'):
                            st.success(f'✅ {_res["total_changed"]}건 변경 완료')
                            st.session_state.pop('_stock_change_preview', None)
                        else:
                            st.error(f'❌ {_res.get("error", "실패")}')

        if not st.session_state.stock_location:
            st.info('👆 먼저 위치를 입력하고 "✅ 위치 설정"을 눌러주세요')
            _stock_ready = False
        else:
            st.success(f'📍 현재 위치: **{st.session_state.stock_location}**')

    # ── 2단계 이하는 _stock_ready일 때만 실행 ──
    _stock_use_fragment = getattr(st, 'fragment', lambda f: f)
    if _stock_ready:
        st.divider()

    @_stock_use_fragment
    def _stock_scan_fragment():
        def _stock_rerun():
            try:
                st.rerun(scope='fragment')
            except TypeError:
                st.rerun()

        if st.session_state.stock_qty_input_mode:
            st.warning(f'🔢 다량 입력 모드 — 다음 스캔은 **{st.session_state.stock_next_qty}개**로 처리됩니다. `#MULTI` 다시 찍으면 해제.')

        if st.session_state.stock_next_qty > 1:
            st.markdown(
                f'<div style="background:#f59e0b;color:white;padding:0.4rem;border-radius:6px;text-align:center;font-weight:bold">'
                f'📦 다음 스캔: {st.session_state.stock_next_qty}개</div>',
                unsafe_allow_html=True)

        def _stock_process_scan(raw):
            _now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            barcode = str(raw or '').strip()
            if not barcode:
                return None
            if barcode.upper() == '#MULTI':
                st.session_state.stock_qty_input_mode = not st.session_state.stock_qty_input_mode
                # 모드가 켜지면 '수량 대기' — 다음에 오는 짧은 숫자를 수량으로 받는다
                st.session_state.stock_qty_pending = st.session_state.stock_qty_input_mode
                if not st.session_state.stock_qty_input_mode:
                    st.session_state.stock_next_qty = 1
                return {
                    'status': 'multi_trigger',
                    'message': '🔢 다량 입력 모드 ON' if st.session_state.stock_qty_input_mode else '1개 모드 복귀',
                    'barcode': barcode, 'name': '', 'qty': 0, 'stock': 0, 'time': _now,
                }
            # 수량은 '#MULTI 직후 한 번' 만 받는다. 예전에는 모드가 켜져 있는 동안 숫자면
            # 전부 수량으로 봐서, 수량을 넣은 뒤 찍은 상품 바코드가 EAN/UPC 처럼 숫자로만
            # 되어 있으면 그것까지 수량(예: 8801234567890개)으로 먹고 상품은 집계되지 않았다.
            # 길이 제한은 그 방어를 한 겹 더 — 상품 바코드는 8자리 이상이다.
            if (st.session_state.stock_qty_input_mode
                    and st.session_state.stock_qty_pending
                    and barcode.isdigit() and len(barcode) <= 5):
                n = int(barcode)
                if n > 0:
                    st.session_state.stock_next_qty = n
                    st.session_state.stock_qty_pending = False
                    return {
                        'status': 'qty_set',
                        'message': f'🔢 다음 스캔 수량 = {n}개',
                        'barcode': barcode, 'name': '', 'qty': n, 'stock': 0, 'time': _now,
                    }
            _qty = st.session_state.stock_next_qty if st.session_state.stock_qty_input_mode else 1
            _sclient = st.session_state.get('pick_gsheet_client')
            _res = stock_update_barcode(
                _sclient, _STOCK_SHEET_URL, _STOCK_SHEET_TAB,
                barcode, _qty, st.session_state.stock_location,
            )
            if st.session_state.stock_qty_input_mode:
                st.session_state.stock_next_qty = 1
                st.session_state.stock_qty_input_mode = False
                st.session_state.stock_qty_pending = False
            if _res['ok']:
                return {
                    'status': 'ok',
                    'message': f"✅ [{_res['name'][:30]}] +{_qty} → 재고 {_res['new_stock']}",
                    'barcode': barcode, 'name': _res['name'], 'qty': _qty,
                    'stock': _res['new_stock'], 'time': _now,
                }
            else:
                return {
                    'status': 'error',
                    'message': f"❌ {_res['error']}",
                    'barcode': barcode, 'name': '', 'qty': 0, 'stock': 0, 'time': _now,
                }

        # 콜백 안에서 위젯 자기 값을 비우는 것만이 Streamlit에서 안정적으로 동작.
        # (외부에서 pop은 fragment rerun과 함께 쓰면 값이 안 지워지는 경우가 있음)
        def _on_stock_scan():
            val = st.session_state.get('stock_scan_input', '')
            if not val:
                return
            _result = _stock_process_scan(val)
            if _result:
                st.session_state.stock_last_result = _result
                st.session_state.stock_scan_log.append(_result)
            st.session_state.stock_scan_counter += 1
            # 위젯 자기 값 초기화 — 콜백 안에서만 허용됨
            st.session_state['stock_scan_input'] = ''

        st.text_input(
            '바코드 입력',
            key='stock_scan_input',
            placeholder='바코드를 스캔하거나 #MULTI 입력...',
            label_visibility='collapsed',
            on_change=_on_stock_scan,
        )

        # 최근 결과 + 음성 안내
        _last = st.session_state.stock_last_result
        _tts_msg = None
        if _last:
            if _last['status'] == 'ok':
                st.success(_last['message'])
                _tts_msg = '재고완료'
            elif _last['status'] == 'error':
                st.error(_last['message'])
                _tts_msg = '다시 찍어주세요'
            else:
                st.info(_last['message'])

        # 최근 5건 간략 로그
        _recent = st.session_state.stock_scan_log[-5:]
        if _recent:
            for _e in reversed(_recent):
                _icon = {'ok': '✅', 'error': '❌', 'multi_trigger': '🔢', 'qty_set': '🔢'}.get(_e['status'], '?')
                st.caption(f"{_icon} {_e['time'][-8:]} | {_e['barcode']} | {_e['message']}")

        # 음성 안내(TTS) + 자동 포커스 (fragment rerun마다 실행)
        # scan_counter를 JS에 삽입해서 매번 고유 HTML 생성 → 캐시 방지
        from streamlit.components.v1 import html as _stock_html
        _tts_js = ''
        if _tts_msg:
            from html import escape as _html_esc
            _tts_js = f"""
            try{{
                window.speechSynthesis.cancel();
                setTimeout(function(){{
                    var u = new SpeechSynthesisUtterance('{_html_esc(_tts_msg)}');
                    u.lang = 'ko-KR'; u.rate = 1.2; u.volume = 1.0;
                    var voices = window.speechSynthesis.getVoices();
                    var koVoice = voices.find(v => v.lang && v.lang.startsWith('ko'));
                    if (koVoice) u.voice = koVoice;
                    window.speechSynthesis.speak(u);
                }}, 80);
            }}catch(e){{}}
            """
            st.session_state.stock_last_result = None

        _uid = st.session_state.stock_scan_counter
        _stock_html(f"""
        <script>
        /* scan_{_uid} */
        (function(){{
            const win = window.parent;
            const doc = win.document;
            {_tts_js}

            function findScan(){{
                const inputs = doc.querySelectorAll('input[type="text"]');
                for (const inp of inputs){{
                    if (inp.placeholder && inp.placeholder.includes('#MULTI')) return inp;
                }}
                return null;
            }}
            function focusScan(){{
                const inp = findScan();
                if (inp && doc.activeElement !== inp) {{ inp.focus(); inp.select(); }}
                return inp;
            }}

            // 즉시 + 여러 타이밍에서 포커스 시도 (fragment rerun 직후 새 input 빠르게 잡기)
            focusScan();
            [10,30,60,100,200,400,800].forEach(d => setTimeout(focusScan, d));

            // ★ Parent window에 영구 interval 설치 (iframe 사라져도 계속 작동)
            //   - 100ms마다 체크 → 스캐너 연사에도 빠르게 복귀
            //   - 사용자가 텍스트 선택 중이면 스킵
            if (!win.__stockFocusInstalled) {{
                win.__stockFocusInstalled = true;
                win.__stockFocusInterval = win.setInterval(function(){{
                    try {{
                        // 다른 input/textarea에 정상적으로 포커스 중이면 스킵
                        const ae = doc.activeElement;
                        const tag = (ae && ae.tagName || '').toLowerCase();
                        if (tag === 'input' || tag === 'textarea') {{
                            // 단, 스캔 input이면 OK
                            if (ae && ae.placeholder && ae.placeholder.includes('#MULTI')) return;
                            // 다른 입력창에서 작업 중이면 방해 X
                            return;
                        }}
                        // 텍스트 선택 중이면 스킵
                        const sel = doc.getSelection && doc.getSelection();
                        if (sel && sel.toString().length > 0) return;
                        // 스캔 input 찾아서 포커스
                        const inputs = doc.querySelectorAll('input[type="text"]');
                        for (const inp of inputs){{
                            if (inp.placeholder && inp.placeholder.includes('#MULTI')) {{
                                if (doc.activeElement !== inp) {{ inp.focus(); inp.select(); }}
                                return;
                            }}
                        }}
                    }} catch(e) {{}}
                }}, 100);  // 0.1초 — 스캐너 연사 따라가는 속도
            }}

            // ★ focusout 즉시 복귀 (parent에 영구 등록)
            if (!win.__stockBlurInstalled) {{
                win.__stockBlurInstalled = true;
                doc.addEventListener('focusout', function(ev){{
                    const t = ev.target;
                    if (t && t.placeholder && t.placeholder.includes('#MULTI')) {{
                        // 스캔 input에서 포커스 빠지면 0ms 뒤 즉시 복귀
                        setTimeout(function(){{
                            const inputs = doc.querySelectorAll('input[type="text"]');
                            for (const inp of inputs){{
                                if (inp.placeholder && inp.placeholder.includes('#MULTI')) {{
                                    if (doc.activeElement !== inp) {{ inp.focus(); inp.select(); }}
                                    return;
                                }}
                            }}
                        }}, 0);
                    }}
                }}, true);
            }}

            // ★ 문자 캡처 (parent에 영구 등록) — R 같은 첫 글자 누락 방지
            if (!win.__stockKeyInstalled) {{
                win.__stockKeyInstalled = true;
                const nativeSetter = Object.getOwnPropertyDescriptor(
                    win.HTMLInputElement.prototype, 'value'
                ).set;
                doc.addEventListener('keydown', function(ev){{
                    const inputs = doc.querySelectorAll('input[type="text"]');
                    let inp = null;
                    for (const x of inputs){{
                        if (x.placeholder && x.placeholder.includes('#MULTI')) {{ inp = x; break; }}
                    }}
                    if (!inp) return;
                    if (doc.activeElement === inp) return;
                    const ae = doc.activeElement;
                    const tag = (ae && ae.tagName || '').toLowerCase();
                    if (tag === 'input' || tag === 'textarea' || (ae && ae.isContentEditable)) return;
                    if (ev.ctrlKey || ev.metaKey || ev.altKey) return;
                    if (ev.key && ev.key.length === 1) {{
                        ev.preventDefault();
                        ev.stopPropagation();
                        nativeSetter.call(inp, (inp.value || '') + ev.key);
                        inp.dispatchEvent(new Event('input', {{bubbles: true}}));
                        inp.focus();
                    }} else if (ev.key === 'Enter') {{
                        ev.preventDefault();
                        ev.stopPropagation();
                        inp.focus();
                        setTimeout(function(){{
                            inp.dispatchEvent(new KeyboardEvent('keydown', {{
                                key: 'Enter', code: 'Enter', keyCode: 13, which: 13, bubbles: true
                            }}));
                        }}, 5);
                    }}
                }}, true);
            }}

            // visibilitychange 시 포커스
            if (!win.__stockVisInstalled) {{
                win.__stockVisInstalled = true;
                doc.addEventListener('visibilitychange', function(){{
                    if (doc.visibilityState === 'visible') {{
                        [50,200,500].forEach(d => setTimeout(focusScan, d));
                    }}
                }}, true);
            }}
        }})();
        </script>
        """, height=0)

    if _stock_ready:
        st.markdown('### 🔍 바코드 스캔')
        st.caption('💡 일반 스캔 = +1 / `#MULTI` → 숫자 → 바코드 = 그 수량만큼 +누적')
        _stock_scan_fragment()

        # ── 스캔 로그 ──
        if st.session_state.stock_scan_log:
            st.divider()
            with st.expander(f"📜 스캔 로그 ({len(st.session_state.stock_scan_log)}건)", expanded=True):
                import pandas as _pds
                _log_rows = []
                for e in reversed(st.session_state.stock_scan_log[-100:]):
                    _icon = {'ok': '✅', 'error': '❌', 'multi_trigger': '🔢', 'qty_set': '🔢'}.get(e['status'], '?')
                    _log_rows.append({
                        '시간': e['time'][-8:],
                        '결과': _icon,
                        '바코드': e['barcode'],
                        '수량': e['qty'] if e['qty'] else '',
                        '누적재고': e['stock'] if e['stock'] else '',
                        '상품명': e['name'][:40],
                    })
                st.dataframe(_pds.DataFrame(_log_rows), use_container_width=True, hide_index=True)

                if st.button('🗑️ 로그 초기화', key='stock_clear_log'):
                    st.session_state.stock_scan_log = []
                    st.session_state.stock_last_result = None
                    st.rerun()


# ── 출고 작업 지시서 탭 ───────────────────────────────

# ══════════════════════════════════════════════════════
# 탭4: PDF 병합
# ══════════════════════════════════════════════════════

# ══════════════════════════════════════════════════════
# 탭7: 로켓배송 발주 중단 공문 작성
# ══════════════════════════════════════════════════════
with tab5:
    st.header('📝 로켓배송 발주 중단 공문')
    st.caption('쿠팡 로켓배송 상품 영구적 발주 중단 요청 공문을 자동으로 생성합니다')

    BANNED_KEYWORDS = ['공급가 협의','발주량 협의','가격 인상','단가','시즌 종료','일시적','잠정적']
    REQUIRED_KEYWORDS = ['영구적 생산 중단','영구적 취급 중단','영구적 생산중단','영구적 취급중단']

    col_left, col_right = st.columns([2, 3])

    with col_left:
        st.subheader('📋 공문 정보 입력')

        # 기본 정보
        with st.expander('🏢 업체 정보', expanded=True):
            company_name = st.text_input('업체명 *', value=kit_config.company_name(), placeholder='예: (주)마켓피아', key='gm_company')
            representative = st.text_input('대표이사 성함 *', placeholder='예: 홍길동', key='gm_rep')
            manager_name = st.text_input('담당자명', placeholder='예: 김담당', key='gm_mgr')
            manager_contact = st.text_input('담당자 연락처', placeholder='예: 010-1234-5678', key='gm_contact')

        with st.expander('📄 문서 정보', expanded=True):
            doc_number = st.text_input(
                '문서번호',
                value=f'제 {datetime.now().year}-001호',
                key='gm_docnum'
            )
            doc_title = st.text_input(
                '제목',
                value='로켓배송 상품 영구적 발주 중단 요청의 건',
                key='gm_title',
                help='발주 중단 외 일반 공문도 가능 — 자유롭게 수정하세요'
            )
            doc_date = st.date_input('문서 날짜', value=datetime.now(), key='gm_date')

        with st.expander('✍️ 발주 중단 사유', expanded=True):
            reason_type = st.radio(
                '사유 유형',
                ['생산 중단', '취급 중단', '직접 입력'],
                horizontal=True,
                key='gm_reason_type'
            )

            if reason_type == '생산 중단':
                default_reason = '당사 제조사의 영구적 생산 중단으로 인하여 해당 상품의 지속적인 공급이 불가능하게 되었습니다.'
            elif reason_type == '취급 중단':
                default_reason = '당사의 영구적 취급 중단 결정으로 인하여 해당 상품의 지속적인 공급이 불가능하게 되었습니다.'
            else:
                default_reason = ''

            _direct_input = (reason_type == '직접 입력')
            reason_detail = st.text_area(
                '사유 상세 내용 *',
                value=default_reason,
                height=120,
                placeholder=(
                    '자유롭게 작성하세요 (일반 공문 등)' if _direct_input
                    else '반드시 "영구적 생산 중단" 또는 "영구적 취급 중단" 문구가 포함되어야 합니다.'
                ),
                key='gm_reason'
            )

            # 사유 유효성 검사 - 직접 입력 모드에서는 키워드 검사 스킵 (일반 공문 용도)
            if _direct_input:
                has_banned = False
                has_required = True  # 검증 통과 처리 (PDF 생성 가능)
                if reason_detail:
                    st.info('💬 직접 입력 모드 — 키워드 검증 없이 자유 작성 가능')
            else:
                has_banned = any(w in reason_detail for w in BANNED_KEYWORDS)
                has_required = any(w in reason_detail for w in REQUIRED_KEYWORDS)
                if reason_detail:
                    if has_banned:
                        banned_found = [w for w in BANNED_KEYWORDS if w in reason_detail]
                        st.error(f'⛔ 금지 키워드 포함: {", ".join(banned_found)}')
                    elif not has_required:
                        st.warning('⚠️ "영구적 생산 중단" 또는 "영구적 취급 중단" 문구가 필요합니다')
                    else:
                        st.success('✅ 사유 검증 통과')

        with st.expander('🖊️ 직인 이미지 (선택)', expanded=False):
            stamp_file = st.file_uploader('직인 이미지 업로드 (PNG 권장)', type=['png','jpg','jpeg'], key='gm_stamp')
            if stamp_file:
                st.image(stamp_file, width=100)
                stamp_size = st.slider('직인 크기', 50, 200, 80, key='gm_stamp_size')
                stamp_x = st.slider('직인 좌우 위치 (%)', 0, 100, 58, key='gm_stamp_x')
                stamp_y = st.slider('직인 상하 위치 (%)', 0, 100, 50, key='gm_stamp_y')
            else:
                stamp_size = 80
                stamp_x = 58
                stamp_y = 50

        st.subheader('📊 표(SKU 목록)')

        # ── 표 구성: 컬럼 수 / 헤더 이름 ──
        with st.expander('🛠️ 표 구성 — 열 추가 / 1행 제목 변경', expanded=False):
            num_cols = int(st.number_input('컬럼 수', min_value=1, max_value=6, value=3, key='gm_table_cols'))
            DEFAULT_HEADERS = ['SKU ID', 'SKU 명칭', '발주 중단 사유', '비고', '추가 1', '추가 2']
            col_headers = []
            _hc = st.columns(num_cols)
            for i in range(num_cols):
                with _hc[i]:
                    col_headers.append(
                        st.text_input(f'{i+1}열 제목', value=DEFAULT_HEADERS[i], key=f'gm_header_{i}')
                    )
            st.caption('💡 헤더에 "사유"가 들어가면 비워둔 칸은 위 사유 상세 내용으로 자동 채워집니다')

        sku_input_type = st.radio('입력 방식', ['엑셀 파일 업로드', '직접 입력'], horizontal=True, key='gm_sku_type')

        # table_rows: 표 본문 행들 (각 행은 길이 num_cols 의 문자열 리스트)
        # sku_list: 기존 호환용 (id/name) — 파일명 등에 사용
        sku_list = []
        table_rows = []

        def _autofill_reason(row_cells, headers, reason_text):
            """헤더에 '사유'가 들어간 컬럼에서 빈 칸을 reason_text로 자동 채움"""
            out = list(row_cells)
            for i, h in enumerate(headers):
                if '사유' in str(h) and (not out[i] or not str(out[i]).strip()):
                    out[i] = reason_text
            return out

        if sku_input_type == '엑셀 파일 업로드':
            st.caption(f'엑셀 시트의 처음 {num_cols}개 열을 순서대로 사용합니다 (1행은 헤더로 무시)')
            sku_excel = st.file_uploader('엑셀 파일 업로드', type=['xlsx','xls'], key='gm_excel')
            if sku_excel:
                try:
                    from openpyxl import load_workbook
                    wb = load_workbook(sku_excel)
                    ws = wb.active
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        cells = list(row) if row else []
                        # 첫 두 칸 비어있으면 스킵
                        if not any(str(c).strip() for c in cells[:max(2, num_cols)] if c is not None):
                            continue
                        row_vals = []
                        for c in range(num_cols):
                            v = cells[c] if c < len(cells) and cells[c] is not None else ''
                            row_vals.append(str(v).strip())
                        # 사유 자동 채움
                        row_vals = _autofill_reason(row_vals, col_headers, reason_detail)
                        table_rows.append(row_vals)
                        # 기존 호환: id/name
                        sku_list.append({
                            'id': row_vals[0] if num_cols >= 1 else '',
                            'name': row_vals[1] if num_cols >= 2 else '',
                        })
                    if table_rows:
                        st.success(f'✅ {len(table_rows)}개 행 로드됨')
                        import pandas as _pd_prev
                        st.dataframe(
                            _pd_prev.DataFrame(table_rows[:5], columns=col_headers),
                            hide_index=True,
                        )
                        if len(table_rows) > 5:
                            st.caption(f'... 외 {len(table_rows)-5}개')
                    else:
                        st.error('데이터를 읽지 못했습니다')
                except Exception as e:
                    st.error(f'파일 읽기 오류: {e}')
        else:
            st.caption('각 행에 컬럼별 값을 입력하세요. (행/열은 위 "표 구성"에서 조정)')
            num_rows = int(st.number_input('행 수', min_value=1, max_value=50, value=3, key='gm_num_rows'))
            for i in range(num_rows):
                _rc = st.columns(num_cols)
                row_vals = []
                for c in range(num_cols):
                    with _rc[c]:
                        ph = '(비우면 사유 자동입력)' if '사유' in col_headers[c] else ''
                        v = st.text_input(
                            f'{i+1}행 {col_headers[c]}',
                            placeholder=ph,
                            key=f'gm_cell_{i}_{c}',
                        )
                        row_vals.append(v)
                # 행에 아무것도 안 들어왔으면 스킵 (사유 자동 채움 후에도)
                if not any(str(v).strip() for v in row_vals):
                    continue
                row_vals = _autofill_reason(row_vals, col_headers, reason_detail)
                table_rows.append(row_vals)
                sku_list.append({
                    'id': row_vals[0] if num_cols >= 1 else '',
                    'name': row_vals[1] if num_cols >= 2 else '',
                })

    with col_right:
        st.subheader('👁️ 미리보기')

        # 본문 2번 문구 (직접 입력 모드면 사유를 그대로 사용, 아니면 기존 고정 문구)
        if reason_type == '직접 입력':
            body_p2 = reason_detail or ''
        else:
            body_p2 = ('2. 당사는 아래와 같은 불가피한 사유(영구적 생산 및 취급 중단)로 인해 해당 상품들의 '
                       '공급을 지속할 수 없게 되었습니다. 이에 따라 로켓배송 서비스의 안정적인 운영을 위해 '
                       '해당 제품들의 발주 중단을 정중히 요청드리는 바입니다.')

        # 공문 미리보기 HTML 생성
        def make_letter_html(company, rep, mgr, contact, docnum, date, title, body2,
                             headers, rows, stamp_data=None, stamp_sz=80, stamp_x=58, stamp_y=50):
            from html import escape as _esc
            date_str = date.strftime('%Y년 %m월 %d일') if hasattr(date, 'strftime') else str(date)
            ncols = len(headers)
            # 헤더 행 (균등 폭)
            col_w = round(100 / max(1, ncols), 2)
            header_cells = ''.join(
                f'<th style="border:1px solid black;padding:8px;width:{col_w}%">{_esc(h)}</th>'
                for h in headers
            )
            # 본문 행
            body_rows = ''
            for r in rows:
                cells_html = ''
                for c_i, cell in enumerate(r):
                    align = 'text-align:center;' if c_i == 0 else ''
                    cells_html += f'<td style="border:1px solid black;padding:6px 8px;{align}font-size:8pt">{_esc(str(cell))}</td>'
                body_rows += f'<tr>{cells_html}</tr>'
            if not body_rows:
                body_rows = (
                    f'<tr><td colspan="{ncols}" style="border:1px solid black;padding:20px;'
                    'text-align:center;color:#999">표 데이터를 입력해 주세요</td></tr>'
                )

            stamp_html = ''
            if stamp_data:
                import base64
                b64 = base64.b64encode(stamp_data).decode()
                stamp_html = f'<img src="data:image/png;base64,{b64}" style="position:absolute;left:{stamp_x}%;top:{stamp_y}%;transform:translate(-50%,-50%);width:{stamp_sz}px;height:{stamp_sz}px;object-fit:contain;mix-blend-mode:multiply;pointer-events:none"/>'

            return f"""
            <div style="background:white;padding:15px;width:100%;box-sizing:border-box;font-family:serif;font-size:9pt;line-height:1.6;color:black">
                <table style="width:100%;border-collapse:collapse;margin-bottom:8px">
                    <tr><td style="width:110px;font-weight:bold;padding:3px 0">문서번호 :</td><td style="padding:3px 0">{_esc(docnum)}</td></tr>
                    <tr><td style="font-weight:bold;padding:3px 0">수&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;신 :</td><td style="font-weight:bold;padding:3px 0">쿠팡 주식회사 귀중</td></tr>
                    <tr><td style="font-weight:bold;padding:3px 0">발&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;신 :</td><td style="padding:3px 0"><b>{_esc(company or '[업체명]')}</b>&nbsp;/&nbsp;{_esc(mgr or '[담당자명]')}&nbsp;{_esc(contact or '[연락처]')}</td></tr>
                    <tr><td style="font-weight:bold;padding:3px 0">제&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;목 :</td><td style="font-weight:bold;text-decoration:underline;padding:3px 0">{_esc(title)}</td></tr>
                </table>
                <hr style="border:1.5px solid black;margin:20px 0"/>
                <p>1. 귀사의 무궁한 발전을 기원합니다.</p>
                <p>{_esc(body2)}</p>
                <div style="text-align:center;font-weight:bold;font-size:13pt;margin:20px 0;letter-spacing:4px">- 아 래 -</div>
                <table style="width:100%;border-collapse:collapse;font-size:10pt;margin-bottom:30px">
                    <thead>
                        <tr style="background:#f3f4f6;text-align:center;font-weight:bold">
                            {header_cells}
                        </tr>
                    </thead>
                    <tbody>{body_rows}</tbody>
                </table>
                <div style="position:relative;text-align:center;margin-top:40px;padding-top:20px;border-top:1px solid #e5e7eb">
                    <div style="font-size:13pt;font-weight:bold;margin-bottom:30px;letter-spacing:2px">{date_str}</div>
                    <h1 style="font-size:20pt;font-weight:bold;letter-spacing:6px;margin-bottom:30px">{company or '[업체명]'}</h1>
                    <div style="display:inline-flex;align-items:center;gap:10px">
                        <span style="font-size:14pt;font-weight:bold;letter-spacing:4px">대표이사&nbsp;&nbsp;{rep or '[성함]'}</span>
                        <span style="font-size:13pt;font-weight:bold">(인)</span>
                    </div>
                    {stamp_html}
                </div>
            </div>"""

        stamp_data = stamp_file.read() if stamp_file else None
        if stamp_data:
            stamp_file.seek(0)

        html = make_letter_html(
            company_name, representative, manager_name, manager_contact,
            doc_number, doc_date, doc_title, body_p2,
            col_headers, table_rows,
            stamp_data, stamp_size, stamp_x, stamp_y
        )
        components_v1.html(html, height=1000, scrolling=True)

        st.divider()

        # PDF 생성 & 다운로드
        is_valid = (
            company_name and representative and table_rows
            and not has_banned and has_required
        ) if reason_detail else False

        if not is_valid:
            if not company_name:
                st.warning('업체명을 입력해주세요')
            elif not representative:
                st.warning('대표이사 성함을 입력해주세요')
            elif not table_rows:
                st.warning('표에 1개 이상 행을 입력해주세요')
            elif not reason_detail:
                st.warning('사유 내용을 입력해주세요')

        if st.button('📄 공문 PDF 생성', type='primary', key='gm_pdf_btn', disabled=not is_valid):
            try:
                buf = io.BytesIO()
                doc = SimpleDocTemplate(
                    buf, pagesize=A4,
                    leftMargin=25*mm, rightMargin=25*mm,
                    topMargin=20*mm, bottomMargin=20*mm
                )

                styles_normal = ParagraphStyle('normal', fontName='NanumReg', fontSize=10, leading=16)
                styles_bold   = ParagraphStyle('bold',   fontName='NanumBold', fontSize=10, leading=16)
                styles_title  = ParagraphStyle('title',  fontName='NanumBold', fontSize=14, leading=20, alignment=1)
                styles_center = ParagraphStyle('center', fontName='NanumBold', fontSize=10, leading=16, alignment=1)
                styles_big    = ParagraphStyle('big',    fontName='NanumBold', fontSize=18, leading=26, alignment=1, spaceAfter=10)
                styles_small  = ParagraphStyle('small',  fontName='NanumReg', fontSize=8, leading=11)

                date_str = doc_date.strftime('%Y년 %m월 %d일')
                story = []

                # 헤더 테이블 — 사용자가 입력한 값은 전부 이스케이프
                # (제목은 이미 하고 있었고 업체명/담당자/대표이사가 빠져 있었다. '<b' 같은
                #  태그 모양 조각이 들어오면 Paragraph 파싱이 깨진다)
                from html import escape as _pdf_esc
                header_data = [
                    [Paragraph('문서번호 :', styles_bold), Paragraph(_pdf_esc(doc_number), styles_normal)],
                    [Paragraph('수      신 :', styles_bold), Paragraph('쿠팡 주식회사 귀중', styles_bold)],
                    [Paragraph('발      신 :', styles_bold), Paragraph(f'{_pdf_esc(company_name)}  /  {_pdf_esc(manager_name)}  {_pdf_esc(manager_contact)}', styles_normal)],
                    [Paragraph('제      목 :', styles_bold), Paragraph(f'<u>{_pdf_esc(doc_title)}</u>', styles_bold)],
                ]
                header_table = Table(header_data, colWidths=[35*mm, None])
                header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
                story.append(header_table)
                story.append(HRFlowable(width='100%', thickness=1.5, color=colors.black, spaceAfter=12))

                story.append(Paragraph('1. 귀사의 무궁한 발전을 기원합니다.', styles_normal))
                story.append(Spacer(1, 8))
                story.append(Paragraph(_pdf_esc(body_p2), styles_normal))
                story.append(Spacer(1, 16))
                story.append(Paragraph('- 아 래 -', styles_center))
                story.append(Spacer(1, 12))

                # 동적 표
                ncols = len(col_headers)
                sku_table_data = [
                    [Paragraph(_pdf_esc(h), styles_bold) for h in col_headers]
                ]
                for row in table_rows:
                    sku_table_data.append([
                        Paragraph(_pdf_esc(str(cell)), styles_small)
                        for cell in row
                    ])

                # 사용 가능 폭(A4 - 좌우 마진) ≈ 160mm — 균등 분할
                total_w = 160 * mm
                col_widths = [total_w / ncols for _ in range(ncols)]
                sku_table = Table(sku_table_data, colWidths=col_widths)
                sku_table.setStyle(TableStyle([
                    ('GRID', (0,0), (-1,-1), 1, colors.black),
                    ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
                    ('ALIGN', (0,0), (0,-1), 'CENTER'),
                    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                    ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ]))
                story.append(sku_table)
                story.append(Spacer(1, 30))

                # 서명
                story.append(Paragraph(date_str, styles_center))
                story.append(Spacer(1, 20))
                story.append(Paragraph(_pdf_esc(company_name), styles_big))
                story.append(Spacer(1, 10))

                # 직인 + 대표이사 서명
                if stamp_data:
                    from reportlab.platypus import Flowable
                    sz = stamp_size * 0.8

                    class StampWithText(Flowable):
                        def __init__(self, stamp_bytes, stamp_sz, rep_name):
                            Flowable.__init__(self)
                            self.stamp_bytes = stamp_bytes
                            self.stamp_sz = stamp_sz
                            self.rep_name = rep_name
                            self.width = 160*mm
                            # 도장이 텍스트 위에 얹혀도 잘리지 않도록 충분한 높이
                            self.height = stamp_sz + 8

                        def draw(self):
                            canvas = self.canv
                            font_size = 10
                            text = f'대표이사   {self.rep_name}   (인)'
                            canvas.setFont('NanumBold', font_size)
                            text_w = canvas.stringWidth(text, 'NanumBold', font_size)
                            text_x = self.width / 2 - text_w / 2
                            # 텍스트를 flowable 세로 중앙에 배치
                            text_y = (self.height - font_size) / 2
                            canvas.drawString(text_x, text_y, text)

                            # "(인)" 위치 정확히 계산 → 그 위에 도장 중심 정렬
                            prefix = f'대표이사   {self.rep_name}   '
                            prefix_w = canvas.stringWidth(prefix, 'NanumBold', font_size)
                            in_w = canvas.stringWidth('(인)', 'NanumBold', font_size)
                            in_center_x = text_x + prefix_w + in_w / 2

                            # 도장 중심을 (인) 중심과 일치, 세로는 텍스트 라인 중앙
                            text_center_y = text_y + font_size / 2
                            stamp_x_pos = in_center_x - self.stamp_sz / 2
                            stamp_y_pos = text_center_y - self.stamp_sz / 2

                            from reportlab.lib.utils import ImageReader
                            stamp_reader = ImageReader(io.BytesIO(self.stamp_bytes))
                            canvas.drawImage(stamp_reader, stamp_x_pos, stamp_y_pos,
                                           width=self.stamp_sz, height=self.stamp_sz,
                                           preserveAspectRatio=True, mask='auto')

                    stamp_flow = StampWithText(stamp_data, sz, representative)
                    stamp_flow.hAlign = 'CENTER'
                    story.append(stamp_flow)
                else:
                    story.append(Paragraph(f'대표이사&nbsp;&nbsp;&nbsp;{_pdf_esc(representative)}&nbsp;&nbsp;&nbsp;(인)', styles_center))

                doc.build(story)
                buf.seek(0)

                today = datetime.now().strftime('%Y%m%d')
                st.success('✅ 공문 PDF 생성 완료!')
                st.download_button(
                    label='⬇️ 공문 PDF 다운로드',
                    data=buf,
                    file_name=f'공문_{(doc_title or "발주중단")[:30]}_{company_name}_{today}.pdf',
                    mime='application/pdf',
                    key='gm_dl'
                )
            except Exception as e:
                st.error(f'❌ PDF 생성 오류: {e}')
                import traceback
                st.code(traceback.format_exc())

# ══════════════════════════════════════════════════════
# 탭6: 쉽먼트 통합 관리
# ══════════════════════════════════════════════════════

# ── 쉽먼트 분석 헬퍼 함수들 ────────────────────────────

@lru_cache(maxsize=256)
def _pdf_page_texts(pdf_bytes):
    """PDF 페이지별 텍스트 목록.

    pdfplumber는 문자 단위 레이아웃 분석까지 해서 느리다(파일당 수백 ms).
    여기서는 정규식 매칭만 하므로 훨씬 빠른 pypdf로 먼저 뽑고,
    결과가 수상하면(텍스트가 비었거나 '박스' 표기가 하나도 없으면)
    해당 파일만 pdfplumber로 다시 읽는다. 느려도 정확한 쪽으로 떨어지는 구조.

    같은 파일을 한 번의 재출력에서 두 번 읽는다 (_extract_manifest_info 와
    _extract_manifest_products). bytes 는 해시 가능하므로 결과를 캐시해 두 번째
    호출은 공짜로 만든다. 반환값은 호출부가 수정하지 않으므로 공유해도 안전하다.
    """
    try:
        texts = [(p.extract_text() or '') for p in PdfReader(io.BytesIO(pdf_bytes)).pages]
        if texts and any('박스' in t for t in texts):
            return texts
    except Exception:
        pass
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        return [(p.extract_text() or '') for p in pdf.pages]


def _extract_manifest_info(pdf_bytes):
    """매니페스트 PDF 바이트에서 박스/송장/쉽먼트번호 정보 추출"""
    pages_info = []
    for i, text in enumerate(_pdf_page_texts(pdf_bytes)):
        box_match = re.search(r'박스\s*(\d+-\d+)', text)
        invoice_match = re.search(r'송장번호\s*\n?\s*(\d{12,})', text)
        if not invoice_match:
            invoice_match = re.search(r'(4\d{11})', text)
        # 쉽먼트번호 (7~10자리, 보통 8자리) — \b로 12자리 운송장번호와 혼동 방지
        shipment_id_match = re.search(r'쉽먼트\s*번호\s*\n\s*(\d{7,10})\b', text)
        if not shipment_id_match:
            shipment_id_match = re.search(r'쉽먼트\s*번호[\s:]+(\d{7,10})\b', text)
        if not shipment_id_match:
            shipment_id_match = re.search(r'쉽먼트\s*번호.{0,50}?\b(\d{7,10})\b', text, re.DOTALL)
        pages_info.append({
            'page_idx': i,
            'box_number': box_match.group(1) if box_match else None,
            'invoice_number': invoice_match.group(1) if invoice_match else None,
            'shipment_id': shipment_id_match.group(1) if shipment_id_match else None,
            'is_main_page': box_match is not None
        })
    return pages_info


def _extract_manifest_products(pdf_bytes):
    """매니페스트 PDF에서 박스별 상품 목록(상품번호/바코드) 추출.
    상품 테이블 행은 'R<12~13자리 바코드> <6~10자리 SKU ID>' 패턴.
    반환: [{'box_number', 'invoice_number', 'shipment_id',
            'products': [{'barcode', 'sku_id'}, ...]}]
    """
    boxes = []
    current = None
    for text in _pdf_page_texts(pdf_bytes):
        box_match = re.search(r'박스\s*(\d+-\d+)', text)
        if box_match:
            if current:
                boxes.append(current)
            invoice_match = re.search(r'송장번호\s*\n?\s*(\d{12,})', text)
            if not invoice_match:
                invoice_match = re.search(r'(4\d{11})', text)
            shipment_id_match = re.search(r'쉽먼트\s*번호\s*\n\s*(\d{7,10})\b', text)
            if not shipment_id_match:
                shipment_id_match = re.search(r'쉽먼트\s*번호[\s:]+(\d{7,10})\b', text)
            if not shipment_id_match:
                shipment_id_match = re.search(r'쉽먼트\s*번호.{0,50}?\b(\d{7,10})\b', text, re.DOTALL)
            current = {
                'box_number': box_match.group(1),
                'invoice_number': invoice_match.group(1) if invoice_match else None,
                'shipment_id': shipment_id_match.group(1) if shipment_id_match else None,
                'products': [],
            }
        if current is None:
            continue
        # 상품 테이블 행 (제품 바코드 + 상품번호)
        for m in re.finditer(r'(R\d{12,13})\s+(\d{6,10})\b', text):
            current['products'].append({
                'barcode': m.group(1),
                'sku_id': m.group(2),
            })
    if current:
        boxes.append(current)
    return boxes


def _extract_label_info(pdf_bytes):
    """라벨(동봉문서) PDF 바이트에서 박스/송장/쉽먼트번호 정보 추출"""
    pages_info = []
    for i, text in enumerate(_pdf_page_texts(pdf_bytes)):
        box_match = re.search(r'박스\s*(\d+-\d+)', text)
        invoice_match = re.search(r'(4\d{11})', text)
        # 동봉문서 '쉽먼트번호\n45232459' 추출 (레이아웃: 라벨 다음 줄에 값)
        shipment_id_match = re.search(r'쉽먼트\s*번호\s*\n\s*(\d{7,10})\b', text)
        if not shipment_id_match:
            shipment_id_match = re.search(r'쉽먼트\s*번호[\s:]+(\d{7,10})\b', text)
        if not shipment_id_match:
            # 유연 패턴: 쉽먼트번호 ~ 50자 안에 7~10자리 독립 숫자
            shipment_id_match = re.search(r'쉽먼트\s*번호.{0,50}?\b(\d{7,10})\b', text, re.DOTALL)
        pages_info.append({
            'page_idx': i,
            'box_number': box_match.group(1) if box_match else None,
            'invoice_number': invoice_match.group(1) if invoice_match else None,
            'shipment_id': shipment_id_match.group(1) if shipment_id_match else None,
        })
    return pages_info


def _group_manifest_pages(pages_info):
    """매니페스트 박스별 그룹핑"""
    groups, current = [], None
    for info in pages_info:
        if info['is_main_page']:
            if current:
                groups.append(current)
            current = {
                'box_number': info['box_number'],
                'invoice_number': info['invoice_number'],
                'shipment_id': info.get('shipment_id'),
                'page_indices': [info['page_idx']],
            }
        else:
            if current:
                current['page_indices'].append(info['page_idx'])
                # 서브페이지에만 쉽먼트번호가 있는 경우 대비 — 첫 번째 발견 값 유지
                if not current.get('shipment_id') and info.get('shipment_id'):
                    current['shipment_id'] = info['shipment_id']
    if current:
        groups.append(current)
    return groups


def _group_label_pages(pages_info):
    """라벨 박스별 그룹핑"""
    box_groups = OrderedDict()
    for info in pages_info:
        box = info['box_number']
        if box not in box_groups:
            box_groups[box] = {
                'box_number': box,
                'invoice_number': info['invoice_number'],
                'shipment_id': info.get('shipment_id'),
                'page_indices': [],
            }
        box_groups[box]['page_indices'].append(info['page_idx'])
        if not box_groups[box]['invoice_number'] and info['invoice_number']:
            box_groups[box]['invoice_number'] = info['invoice_number']
        if not box_groups[box].get('shipment_id') and info.get('shipment_id'):
            box_groups[box]['shipment_id'] = info['shipment_id']
    return list(box_groups.values())


def _render_labels_4up_combined(label_items, dpi=300):
    """여러 PDF의 여러 페이지를 순서대로 모아 4분할 A4로 렌더.
    송장별로 끊지 않고 전체를 연속해서 4개씩 배치 → 빈 슬롯 최소화.
    label_items: [(pdf_bytes, page_indices), ...]
    반환: PIL Image 리스트 (A4 페이지들)
    """
    images = []
    pdf_cache = {}
    for pdf_bytes, page_indices in label_items:
        key = id(pdf_bytes)
        if key not in pdf_cache:
            pdf_cache[key] = pdfium.PdfDocument(pdf_bytes)
        pdf_doc = pdf_cache[key]
        for idx in page_indices:
            page = pdf_doc[idx]
            bitmap = page.render(scale=dpi / 72)
            images.append(bitmap.to_pil())
    for pdf_doc in pdf_cache.values():
        try:
            pdf_doc.close()
        except Exception:
            pass

    a4_w, a4_h = int(8.27 * dpi), int(11.69 * dpi)
    slot_w, slot_h = a4_w // 2, a4_h // 2
    result = []
    for start in range(0, len(images), 4):
        chunk = images[start:start + 4]
        canvas = Image.new('RGB', (a4_w, a4_h), 'white')
        positions = [(0, 0), (slot_w, 0), (0, slot_h), (slot_w, slot_h)]
        for i, img in enumerate(chunk):
            ratio = img.width / img.height
            s_ratio = slot_w / slot_h
            if ratio > s_ratio:
                nw, nh = slot_w, int(slot_w / ratio)
            else:
                nh, nw = slot_h, int(slot_h * ratio)
            resized = img.resize((nw, nh), Image.LANCZOS)
            x = positions[i][0] + (slot_w - nw) // 2
            y = positions[i][1] + (slot_h - nh) // 2
            canvas.paste(resized, (x, y))
        result.append(canvas)
    return result


def _parse_csv_bytes(csv_bytes):
    """CSV 바이트 → 아이템 리스트"""
    text = csv_bytes.decode('utf-8-sig', errors='replace')
    reader = csv.reader(text.splitlines())
    rows = list(reader)
    if not rows:
        return []
    items = []
    for row in rows[1:]:
        if len(row) < 11 or not (row[1] if len(row) > 1 else '').strip():
            continue
        def safe(idx, default=''):
            return row[idx].strip() if idx < len(row) else default
        try:
            qty = int(safe(7, '0') or '0')
        except ValueError:
            qty = 0
        items.append({
            'orderNumber': safe(0),  # A열 발주번호
            'logisticsCenter': safe(1),
            'expectedDate': safe(3),
            'productBarcode': safe(5),
            'productName': safe(6),
            'quantity': qty,
            'shipmentNumber': safe(8),
            # 가공 전 원본 — 한진 엑셀에서 잘린 과학표기 감지에 사용
            'shipmentNumberRaw': safe(8),
            'orderDate': safe(9),
            'boxNumber': safe(10),
            'location': safe(12),
        })
    return items


def _parse_xlsx_bytes(xlsx_bytes):
    """엑셀 바이트 → 아이템 리스트 (헤더 기반 자동 매핑)"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    # 헤더 행 찾기
    headers = None
    for row in rows_iter:
        cells = [str(c).strip() if c is not None else '' for c in row]
        if '바코드' in cells and '상품명' in cells:
            headers = cells
            break
    if not headers:
        return []

    FIELD_MAP = {
        'orderNumber':     ['발주번호', '발주서번호'],
        'logisticsCenter': ['물류센터', 'FC'],
        'expectedDate':    ['입고예정일'],
        'productBarcode':  ['바코드', '상품바코드'],
        'productName':     ['상품명', '품명'],
        'quantity':        ['수량'],
        'shipmentNumber':  ['쉽먼트운송장', '송장번호', '운송장번호'],
        'orderDate':       ['발주일'],
        'boxNumber':       ['박스번호'],
        'location':        ['위치', '적재위치'],
    }

    col_map = {}
    for field, keywords in FIELD_MAP.items():
        for kw in keywords:
            for i, h in enumerate(headers):
                if h == kw or h.startswith(kw) or kw in h:
                    col_map[field] = i
                    break
            if field in col_map:
                break

    if 'productBarcode' not in col_map:
        return []

    items = []
    for row in rows_iter:
        cells = [str(c).strip() if c is not None else '' for c in row]
        def safe(field, default=''):
            idx = col_map.get(field)
            if idx is not None and idx < len(cells):
                return cells[idx]
            return default
        barcode = safe('productBarcode')
        if not barcode:
            continue
        try:
            qty = int(float(safe('quantity', '0') or '0'))
        except (ValueError, TypeError):
            qty = 0
        # 쉽먼트운송장번호가 과학표기법(4.62E+11)인 경우 정수로 변환
        shipment = safe('shipmentNumber')
        try:
            if 'E' in shipment or 'e' in shipment:
                shipment = str(int(float(shipment)))
        except (ValueError, TypeError):
            pass
        items.append({
            'orderNumber': safe('orderNumber'),
            'logisticsCenter': safe('logisticsCenter'),
            'expectedDate': safe('expectedDate'),
            'productBarcode': barcode,
            'productName': safe('productName'),
            'quantity': qty,
            'shipmentNumber': shipment,
            # 가공 전 원본 — 한진 엑셀에서 잘린 과학표기 감지에 사용
            'shipmentNumberRaw': safe('shipmentNumber'),
            'orderDate': safe('orderDate'),
            'boxNumber': safe('boxNumber'),
            'location': safe('location'),
        })
    wb.close()
    return items


# ── 입고 검수 탭 ──────────────────────────────────────


# ── 쉽먼트 재출력 탭 ──────────────────────────────────────
# 한진택배 엑셀 생성 시 참조하는 주소리스트 시트 (E열 키 = 물류센터명)
HANJIN_ADDR_SHEET_URL = kit_config.courier_addr_sheet_url()


def _norm_invoice_strict(v):
    """운송장번호 정규화 → (정규화값, 유효여부).

    엑셀이 긴 숫자를 과학표기로 저장하면서 자릿수를 버리는 경우가 있다(예: 4.62E+11).
    이때 int(float(...)) 로 변환하면 462000000000 처럼 그럴듯하지만 **틀린** 번호가
    만들어지고, 서로 다른 송장이 같은 값으로 뭉개져 여러 박스가 택배 1건으로 합쳐진다.
    유효숫자가 정수부 자릿수보다 적으면 원본 복구가 불가능하므로 ('', False) 를 돌린다.
    """
    from decimal import Decimal, InvalidOperation

    s = str(v if v is not None else '').strip()
    if not s or s.lower() in ('nan', 'none'):
        return '', False
    if s.isdigit():
        return s, True

    if 'e' not in s.lower():
        # 과학표기가 아니면 '123456789012.0' 형태만 정수화하고 나머지는 원본 유지
        # (하이픈/공백이 섞인 번호를 임의로 훼손하지 않기 위함)
        try:
            f = float(s)
        except (ValueError, TypeError):
            return s, True
        return (str(int(f)), True) if f.is_integer() else (s, True)

    try:
        d = Decimal(s)
    except (InvalidOperation, ValueError):
        return s, True
    if d <= 0:
        return '', False
    if len(d.as_tuple().digits) < len(str(int(d))):
        return '', False  # 잘린 과학표기 — 복구 불가
    return str(int(d)), True


def _hanjin_build_excel(pick_df, sclient, addr_url, addr_tab='주소리스트',
                        empty_msg='출고확인 데이터가 비어있습니다 (피킹&분류 탭에서 시트를 먼저 로드)',
                        box_num_map=None):
    """출고확인 데이터(pick_df_출고) + 주소리스트 시트 → 한진택배 양식 엑셀.

    그룹화: (물류센터, 쉽먼트운송장번호) — 그룹당 1행, 수량 합계
    주소 매칭: 물류센터 명을 주소리스트 E열(키)에서 검색 (정확/접두/부분)
    행 순서: 박스번호 순 (출고지시서/쉽먼트 출력 순서와 동일하게 맞춰 정리하기 쉽게).
             box_num_map({송장번호: 박스번호})을 주면 그대로 쓰고,
             안 주면 pick_df에서 직접 부여해 쓴다.

    반환: (BytesIO, 생성행수, 매칭실패목록, 송장번호불량목록)
    """
    import io
    import pandas as _pd
    from openpyxl import Workbook

    if pick_df is None or pick_df.empty:
        raise ValueError(empty_msg)
    if sclient is None:
        raise ValueError('Google Sheets 인증 실패 — 서비스 계정 키 확인 필요')

    # 1) 주소리스트 시트 읽기 (스냅샷 캐시 — 연속 생성 시 재조회 안 함)
    try:
        addr_rows = _gs_snapshot(sclient, addr_url, addr_tab)
    except Exception:
        raise ValueError(f'"{addr_tab}" 탭을 찾을 수 없습니다')
    if addr_rows is None:
        raise ValueError(f'"{addr_tab}" 탭을 찾을 수 없습니다')
    if len(addr_rows) < 2:
        raise ValueError('주소리스트 시트가 비어있습니다')

    # E열(인덱스 4)을 키로 하는 매핑
    addr_map = {}
    for r in addr_rows[1:]:
        if len(r) > 4 and str(r[4]).strip():
            addr_map[str(r[4]).strip()] = r

    # 2) (물류센터, 송장번호) 그룹화 + 수량 합계
    df = pick_df.copy()
    fc_col = '물류센터(FC)'
    ship_col = '쉽먼트운송장번호'
    qty_col = '수량'

    if fc_col not in df.columns or ship_col not in df.columns:
        raise ValueError(f'필수 컬럼 누락 (필요: {fc_col}, {ship_col})')

    df = df[df[fc_col].astype(str).str.strip() != '']

    # 송장번호 정규화 — 잘린 과학표기(4.62E+11)는 복구 불가이므로 제외하고 따로 보고.
    # 그냥 두면 서로 다른 송장이 한 값으로 뭉개져 N개 박스가 택배 1건으로 합쳐진다.
    _pairs = [_norm_invoice_strict(v) for v in df[ship_col]]
    invalid = sorted({
        f'{str(_fc).strip()} (원본 값: {str(_orig).strip()})'
        for _fc, _orig, (_nv, _ok) in zip(df[fc_col], df[ship_col], _pairs)
        if not _ok and str(_orig).strip() and str(_orig).strip().lower() not in ('nan', 'none')
    })
    df = df.assign(**{ship_col: [_nv for _nv, _ in _pairs]})
    df = df[df[ship_col].astype(str).str.strip() != '']

    if df.empty:
        _extra = f' (송장번호 판독 불가 {len(invalid)}건 제외됨)' if invalid else ''
        raise ValueError(f'유효한 행이 없습니다{_extra}')

    if qty_col in df.columns:
        df[qty_col] = _pd.to_numeric(df[qty_col], errors='coerce').fillna(0).astype(int)
    else:
        df[qty_col] = 1

    grouped = df.groupby([fc_col, ship_col], as_index=False)[qty_col].sum()

    # 2-1) 박스번호 순으로 행 정렬
    # groupby 기본 순서는 (물류센터, 송장번호)라 박스번호와 무관하다.
    # 출고지시서/쉽먼트 출력이 박스번호 순이므로 택배 엑셀도 같은 순서여야 대조가 쉽다.
    if box_num_map is None:
        try:
            box_num_map = assign_box_numbers(_pick_df_to_items(pick_df))
        except Exception:
            box_num_map = {}
    # 맵 키도 df와 같은 방식으로 정규화해야 매칭된다 (과학표기 등)
    _bn_map = {}
    for _k, _v in (box_num_map or {}).items():
        _nk, _ok = _norm_invoice_strict(_k)
        if _ok and str(_nk).strip():
            _bn_map[str(_nk).strip()] = _v
    if _bn_map:
        _bn = grouped[ship_col].astype(str).str.strip().map(_bn_map)
        grouped = (grouped.assign(_box_no=_bn)
                          .sort_values('_box_no', na_position='last', kind='stable')
                          .drop(columns='_box_no'))

    # 3) 택배 양식 행 만들기 — 열 구성은 ⚙️ 설정의 택배 양식을 따른다
    layout = kit_config.courier_layout()
    sender = kit_config.sender_defaults()
    output_rows = [[head for head, _ in layout]]
    unmatched = []

    for _, row in grouped.iterrows():
        fc = str(row[fc_col]).strip()
        ship = str(row[ship_col]).strip()
        qty = int(row[qty_col])

        # 매칭: 정확 → 접두(- 또는 _) → 부분
        key = None
        if fc in addr_map:
            key = fc
        else:
            prefix = fc.split('-')[0].split('_')[0].strip()
            if prefix and prefix in addr_map:
                key = prefix
            else:
                for k in addr_map:
                    if k and (k in fc or fc in k):
                        key = k
                        break

        if not key:
            unmatched.append(f'{fc} (송장 …{ship[-6:]})')
            continue

        info = addr_map[key]
        def _g(idx):
            return info[idx] if len(info) > idx else ''

        # 받는분(박스명) = 물류센터 + 송장끝6자리 + 수량
        box_label = f'{fc}-{ship[-6:]} ({qty}개)'

        values = {
            '보내는분': _g(0) or sender['보내는분'],
            '보내는전화1': _g(1) or sender['보내는전화1'],
            '보내는전화2': _g(2),
            '보내는주소': _g(3) or sender['보내는주소'],
            '받는분': box_label,
            '받는전화1': _g(5),
            '받는전화2': _g(6),
            '받는우편': _g(7),
            '받는주소': _g(8),
            '물류센터': fc,
            '송장번호': ship,
            '수량': qty,
            '빈칸': '',
        }
        output_rows.append([values[field] for _, field in layout])

    # 4) 엑셀 파일 생성
    wb = Workbook()
    ws = wb.active
    ws.title = kit_config.courier_name()[:31] or '택배출고'
    for r_i, r_data in enumerate(output_rows, start=1):
        for c_i, v in enumerate(r_data, start=1):
            ws.cell(row=r_i, column=c_i, value=v)
    # 헤더 폰트 굵게
    from openpyxl.styles import Font
    for cell in ws[1]:
        cell.font = Font(bold=True)
    # 컬럼 폭 적당히
    _widths = {'보내는주소': 40, '받는주소': 40, '받는분': 28}
    for c_i, (_, field) in enumerate(layout, start=1):
        ws.column_dimensions[get_column_letter(c_i)].width = _widths.get(field, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(output_rows) - 1, unmatched, invalid


def _items_to_hanjin_df(items):
    """재출력용 items 리스트(CSV/XLSX 파싱 결과) → 한진 엑셀 생성용 DataFrame.

    _hanjin_build_excel 이 요구하는 컬럼(물류센터(FC)/쉽먼트운송장번호/수량)만 만든다.
    송장번호는 **가공 전 원본**(shipmentNumberRaw)을 그대로 넘긴다 — 잘린 과학표기를
    _hanjin_build_excel 이 감지해서 걸러낼 수 있어야 하기 때문.
    """
    import pandas as _pd

    rows = []
    for it in (items or []):
        try:
            qty = int(it.get('quantity', 0) or 0)
        except (ValueError, TypeError):
            qty = 0
        _raw = it.get('shipmentNumberRaw')
        if _raw is None or not str(_raw).strip():
            _raw = it.get('shipmentNumber', '')
        rows.append({
            '물류센터(FC)': str(it.get('logisticsCenter', '') or '').strip(),
            '쉽먼트운송장번호': str(_raw if _raw is not None else '').strip(),
            '수량': qty,
        })
    return _pd.DataFrame(rows, columns=['물류센터(FC)', '쉽먼트운송장번호', '수량'])


def _pick_df_to_items(df):
    """피킹&분류 탭의 df_출고 → create_work_order_pdf 용 items 리스트 변환"""
    def _norm_ship(v):
        """운송장번호 정규화 (과학표기/float → 정수 문자열)"""
        s = str(v if v is not None else '').strip()
        if not s or s.lower() in ('nan', 'none'):
            return ''
        if 'E' in s or 'e' in s or '.' in s:
            try:
                return str(int(float(s)))
            except (ValueError, TypeError):
                pass
        return s

    # A열(발주번호) 헤더 탐색 — 이름 우선, 없으면 첫 번째 컬럼(A열 위치)
    _order_col = None
    for _c in df.columns:
        if str(_c).strip().replace(' ', '') in ('발주번호', '발주서번호', 'PO', 'PO번호'):
            _order_col = _c
            break
    if _order_col is None and len(df.columns) > 0:
        _order_col = df.columns[0]

    items = []
    for _, row in df.iterrows():
        try:
            qty = int(row.get('수량', 0) or 0)
        except (ValueError, TypeError):
            qty = 0
        items.append({
            'orderNumber': str(row.get(_order_col, '') or '').strip() if _order_col is not None else '',
            'logisticsCenter': str(row.get('물류센터(FC)', '') or '').strip(),
            'expectedDate': str(row.get('입고예정일(EDD)', '') or row.get('입고예정일', '') or '').strip(),
            'productBarcode': str(row.get('바코드', '') or '').strip(),
            'productName': str(row.get('상품명', '') or '').strip(),
            'quantity': qty,
            'shipmentNumber': _norm_ship(row.get('쉽먼트운송장번호', '')),
            # 가공 전 원본 — 한진 엑셀에서 잘린 과학표기 감지에 사용
            'shipmentNumberRaw': str(row.get('쉽먼트운송장번호', '') or '').strip(),
            'orderDate': str(row.get('발주일', '') or '').strip(),
            'boxNumber': str(row.get('박스번호', '') or '').strip(),
            'location': str(row.get('위치', '') or row.get('적재위치', '') or '').strip(),
        })
    return items


def _run_reprint_pipeline(rp_items, rp_manifest_files, rp_label_files,
                          existing_box_map=None, write_new_to_sheet=False,
                          sheet_client=None, sheet_url=None, sheet_tab=None,
                          on_progress=None, on_status=None):
    """재출력 공통 처리: 매니페스트/라벨 PDF 파싱 → 송장 매칭 → 박스번호 부여
    → 출고지시서 PDF 생성 → 통합 PDF 병합.

    existing_box_map: 기존 {송장: 박스번호}. None이면 fresh 부여(assign_box_numbers).
        dict면 assign_box_numbers_with_existing로 기존값 보존.
    write_new_to_sheet: True면 신규 부여분을 시트 M열에 기록 (기존값 보존).
    on_progress: 0~1 float 콜백 (진행률 UI용)
    on_status: str 콜백 (상태 메시지용)

    반환: 성공시 dict, 실패시 {'error': str}
    """
    def _p(v):
        if on_progress:
            try:
                on_progress(v)
            except Exception:
                pass

    def _s(msg):
        if on_status:
            try:
                on_status(msg)
            except Exception:
                pass

    _s('📄 데이터 분석 중...')
    if not rp_items:
        return {'error': '데이터에서 항목을 찾을 수 없습니다.'}

    csv_invoices = set()
    rp_grouped = OrderedDict()
    for item in rp_items:
        inv = item.get('shipmentNumber', '')
        if inv:
            csv_invoices.add(inv)
            rp_grouped.setdefault(inv, []).append(item)
    for key in rp_grouped:
        rp_grouped[key].sort(key=lambda x: (
            [int(c) if c.isdigit() else c.lower()
             for c in re.split(r'(\d+)', x.get('boxNumber', ''))],
            x.get('productName', '')
        ))
    _p(0.2)

    _s('📊 매니페스트/라벨 분석 중...')
    rp_manifest_data = []
    for fname, mf in rp_manifest_files:
        m_bytes = mf.read(); mf.seek(0)
        m_info = _extract_manifest_info(m_bytes)
        m_groups = _group_manifest_pages(m_info)
        sorted_m = sorted(m_groups, key=lambda g: g['invoice_number'] or '')
        rp_manifest_data.append((m_bytes, sorted_m))
    rp_label_data = []
    for fname, lf in rp_label_files:
        l_bytes = lf.read(); lf.seek(0)
        l_info = _extract_label_info(l_bytes)
        l_groups = _group_label_pages(l_info)
        sorted_l = sorted(l_groups, key=lambda g: g['invoice_number'] or '')
        rp_label_data.append((l_bytes, sorted_l))
    _p(0.4)

    all_manifest_invoices = set()
    invoice_to_shipment_id = {}
    for _, sorted_m in rp_manifest_data:
        for g in sorted_m:
            if g['invoice_number']:
                all_manifest_invoices.add(g['invoice_number'])
                if g.get('shipment_id'):
                    invoice_to_shipment_id[g['invoice_number']] = g['shipment_id']
    all_label_invoices = set()
    for _, sorted_l in rp_label_data:
        for g in sorted_l:
            if g['invoice_number']:
                all_label_invoices.add(g['invoice_number'])
                if g.get('shipment_id'):
                    invoice_to_shipment_id[g['invoice_number']] = g['shipment_id']

    all_pdf_invoices = all_manifest_invoices | all_label_invoices
    matched = csv_invoices & all_pdf_invoices
    not_in_manifest = csv_invoices - all_pdf_invoices
    not_in_csv = all_pdf_invoices - csv_invoices

    # ===== SKU 기반 fallback 매칭 + 매니페스트 순서 추출 =====
    # 매니페스트 페이지 순서대로 송장번호 + 송장별 바코드 집합 수집
    manifest_invoice_order = []
    manifest_barcodes_by_inv = {}  # invoice → set(barcodes)
    _seen_inv = set()
    for m_bytes, _ in rp_manifest_data:
        try:
            m_products = _extract_manifest_products(m_bytes)
        except Exception:
            m_products = []
        for box in m_products:
            inv = box.get('invoice_number')
            if not inv:
                continue
            if inv not in _seen_inv:
                manifest_invoice_order.append(inv)
                _seen_inv.add(inv)
                manifest_barcodes_by_inv[inv] = set()
            for p in box.get('products', []):
                bc = (p.get('barcode') or '').strip()
                if bc:
                    manifest_barcodes_by_inv[inv].add(bc)
    # 매니페스트에 송장번호가 한 번도 안 나오는 경우(레이블만 등) → label_invoices를 뒤에 추가
    for inv in sorted(all_pdf_invoices):
        if inv not in _seen_inv:
            manifest_invoice_order.append(inv)
            _seen_inv.add(inv)

    # 시트 송장별 바코드 집합
    sheet_barcodes_by_inv = {}
    for inv, items in rp_grouped.items():
        sheet_barcodes_by_inv[inv] = {
            (it.get('productBarcode') or '').strip()
            for it in items if (it.get('productBarcode') or '').strip()
        }

    # 송장번호 매칭 안 된 매니페스트 송장에 대해 SKU(바코드) 기반 fallback
    sku_matched = {}  # manifest_invoice → sheet_invoice
    unmatched_pdf = set(manifest_invoice_order) - matched
    unused_sheet = set(csv_invoices) - matched
    for m_inv in unmatched_pdf:
        m_bcs = manifest_barcodes_by_inv.get(m_inv, set())
        if not m_bcs:
            continue
        best_inv = None
        best_score = 0.0
        for s_inv in list(unused_sheet):
            s_bcs = sheet_barcodes_by_inv.get(s_inv, set())
            if not s_bcs:
                continue
            common = m_bcs & s_bcs
            denom = max(len(m_bcs), len(s_bcs))
            if denom == 0:
                continue
            score = len(common) / denom
            # 80% 이상 일치하면 같은 출고로 판단
            if score >= 0.8 and score > best_score:
                best_score = score
                best_inv = s_inv
        if best_inv:
            sku_matched[m_inv] = best_inv
            unused_sheet.discard(best_inv)

    # 출력 매핑: 매니페스트 송장번호 → 시트 송장번호
    inv_to_sheet_inv = {inv: inv for inv in matched}
    inv_to_sheet_inv.update(sku_matched)

    if not inv_to_sheet_inv:
        # 진단: 양쪽 샘플 값 비교로 포맷 차이 확인
        _csv_sample = sorted(list(csv_invoices))[:3] if csv_invoices else []
        _pdf_sample = sorted(list(all_pdf_invoices))[:3] if all_pdf_invoices else []
        _diag = (
            f'데이터와 매니페스트/라벨 간 매칭되는 송장번호가 없습니다 (송장 매칭 + SKU 매칭 모두 실패).\n\n'
            f'**진단**:\n'
            f'- 시트/CSV 송장번호 ({len(csv_invoices)}건) 샘플: `{_csv_sample}`\n'
            f'- PDF 송장번호 ({len(all_pdf_invoices)}건) 샘플: `{_pdf_sample}`\n\n'
            f'👉 두 값의 **포맷이 다르면** 매칭 실패 원인 (예: 과학표기 `4.62E+11` vs 정수 `462139010304`).\n'
            f'👉 SKU/바코드도 시트와 매니페스트가 서로 다르면 자동 매칭 불가.'
        )
        return {'error': _diag}

    # 매니페스트 페이지 순서로 정렬된 송장 (출력 순서)
    ordered_manifest_invs = [inv for inv in manifest_invoice_order if inv in inv_to_sheet_inv]
    # 매니페스트에 없지만 시트+라벨 매칭된 송장(드물지만) → 뒤에 추가
    for inv in sorted(inv_to_sheet_inv.keys()):
        if inv not in ordered_manifest_invs:
            ordered_manifest_invs.append(inv)

    # 박스번호 부여용 매칭(시트 송장번호 기준)
    matched_sheet_invs = set(inv_to_sheet_inv.values())

    _s('📄 출고지시서 생성 중...')
    # 박스번호는 시트 송장번호 기준으로 부여 (M열 키와 일치시킴)
    rp_all_items = []
    for s_inv in sorted(matched_sheet_invs):
        rp_all_items.extend(rp_grouped[s_inv])
    if existing_box_map is not None:
        rp_ship_to_box_num = assign_box_numbers_with_existing(rp_all_items, existing_box_map)
    else:
        rp_ship_to_box_num = assign_box_numbers(rp_all_items)

    # 출력 순서를 박스번호 순으로 재정렬한다.
    # 기본값은 매니페스트 업로드/페이지 순서라 박스번호와 무관해서,
    # 뽑은 출고지시서·쉽먼트·통합본을 박스 순으로 정리하려면 손으로 다시 골라야 했다.
    # 박스번호가 없는 송장(국내재고/부족 등)은 순서를 유지한 채 뒤로 보낸다.
    _box_order = {
        m_inv: rp_ship_to_box_num.get(inv_to_sheet_inv[m_inv])
        for m_inv in ordered_manifest_invs
    }
    ordered_manifest_invs.sort(
        key=lambda m_inv: (
            _box_order[m_inv] is None,
            _box_order[m_inv] if _box_order[m_inv] is not None else 0,
        )
    )

    # 신규 부여된 박스번호를 시트 M열에 기록
    new_box_only = {}
    sheet_write_result = None  # None: 시도 안 함 / >=0: 기록된 셀 수 / -1: API 실패
    if existing_box_map is not None:
        new_box_only = {s: n for s, n in rp_ship_to_box_num.items()
                        if s not in existing_box_map}
        if write_new_to_sheet and new_box_only and sheet_client and sheet_url and sheet_tab:
            sheet_write_result = pick_write_box_numbers(
                sheet_client, sheet_url, sheet_tab,
                new_box_only, only_empty=True,
            )

    # 출고지시서 PDF: 매니페스트 송장 키로 저장 (PDF 통합 시 같은 키 사용)
    rp_so_by_invoice = {}
    for m_inv in ordered_manifest_invs:
        s_inv = inv_to_sheet_inv[m_inv]
        inv_items = rp_grouped.get(s_inv, [])
        if not inv_items:
            continue
        auto_box_num = rp_ship_to_box_num.get(s_inv)
        center = inv_items[0].get('logisticsCenter', '')
        gk = f"{center}_{m_inv}" if center else m_inv
        # 출고지시서 헤더에 표시되는 쉽먼트번호: 매니페스트의 shipment_id 우선
        real_shipment_id = invoice_to_shipment_id.get(m_inv, m_inv)
        pdf_buf = create_work_order_pdf(gk, inv_items, real_shipment_id, auto_box_num)
        rp_so_by_invoice[m_inv] = pdf_buf
    _p(0.6)

    _s('📎 통합 PDF 생성 중...')
    rp_final_writer = PdfWriter()
    rp_shipment_only_writer = PdfWriter()
    rp_so_only_writer = PdfWriter()
    rp_total = rp_label_total = rp_shipment_total = rp_so_total = 0

    # 매니페스트 송장 단위로 페이지 인덱스 수집 (page-order 보존)
    manifest_pages_by_inv = {}
    for m_bytes, sorted_m in rp_manifest_data:
        for g in sorted_m:
            inv = g['invoice_number']
            if not inv or inv not in inv_to_sheet_inv:
                continue
            manifest_pages_by_inv.setdefault(inv, []).append((m_bytes, g['page_indices']))

    label_groups_by_inv = {}
    for l_bytes, sorted_l in rp_label_data:
        for g in sorted_l:
            inv = g['invoice_number']
            if not inv or inv not in inv_to_sheet_inv:
                continue
            label_groups_by_inv.setdefault(inv, []).append((l_bytes, g))

    # Pass 1: 매니페스트 페이지 순서대로 [출고지시서 → 매니페스트(동봉문서)] 배치
    # 한 매니페스트 파일에 송장이 여러 개 들어있어 같은 바이트를 송장마다 다시 파싱하지
    # 않도록 PdfReader 를 파일 단위로 재사용한다.
    _m_readers = {}
    for inv_num in ordered_manifest_invs:
        if inv_num in rp_so_by_invoice:
            so_buf = rp_so_by_invoice[inv_num]
            so_buf.seek(0)
            so_reader = PdfReader(so_buf)
            for page in so_reader.pages:
                rp_final_writer.add_page(page)
                rp_so_only_writer.add_page(page)
            rp_total += len(so_reader.pages)
            rp_so_total += len(so_reader.pages)
        for m_bytes, page_indices in manifest_pages_by_inv.get(inv_num, []):
            m_reader = _m_readers.get(id(m_bytes))
            if m_reader is None:
                m_reader = _m_readers[id(m_bytes)] = PdfReader(io.BytesIO(m_bytes))
            for pidx in page_indices:
                rp_final_writer.add_page(m_reader.pages[pidx])
                rp_shipment_only_writer.add_page(m_reader.pages[pidx])
                rp_total += 1
                rp_shipment_total += 1

    # Pass 2: 모든 송장의 라벨을 송장 순서대로 모아서 "한꺼번에" 4분할 배치
    # → 송장별로 끊지 않아 빈 슬롯 최소화, 4장씩 빽빽하게 배치됨
    _all_label_items = []
    for inv_num in ordered_manifest_invs:
        for l_bytes, grp in label_groups_by_inv.get(inv_num, []):
            _all_label_items.append((l_bytes, grp['page_indices']))

    if _all_label_items:
        four_up_pages = _render_labels_4up_combined(_all_label_items)
        for img in four_up_pages:
            img_buf = io.BytesIO()
            img.save(img_buf, format='PDF', resolution=300)
            img_buf.seek(0)
            lp = PdfReader(img_buf)
            rp_final_writer.add_page(lp.pages[0])
            rp_shipment_only_writer.add_page(lp.pages[0])
            rp_total += 1
            rp_label_total += 1
            rp_shipment_total += 1

    _p(0.9)
    rp_final_buf = io.BytesIO()
    rp_final_writer.write(rp_final_buf)
    rp_final_buf.seek(0)
    rp_shipment_only_buf = io.BytesIO()
    rp_shipment_only_writer.write(rp_shipment_only_buf)
    rp_shipment_only_buf.seek(0)
    rp_so_only_buf = io.BytesIO()
    rp_so_only_writer.write(rp_so_only_buf)
    rp_so_only_buf.seek(0)
    _p(1.0)
    _s('✅ 완료!')

    return {
        'final_bytes': rp_final_buf.getvalue(),
        'shipment_only_bytes': rp_shipment_only_buf.getvalue(),
        'so_only_bytes': rp_so_only_buf.getvalue(),
        'total': rp_total,
        'shipment_total': rp_shipment_total,
        'so_total': rp_so_total,
        'matched': len(inv_to_sheet_inv),
        'invoice_matched': len(matched),
        'sku_matched': sku_matched,  # {매니페스트 송장: 시트 송장} (SKU로만 매칭된 것)
        'not_in_manifest': sorted(not_in_manifest),
        # D열 입고예정일 함께 제공 — 송장당 하나라서 시트에서 행 찾기 쉬움
        # (발주번호는 한 송장에 여러 개라 표가 지저분해짐)
        'not_in_manifest_detail': [
            {
                '입고예정일(D열)': next(
                    (str(it.get('expectedDate', '') or '').strip()
                     for it in rp_grouped.get(_inv, [])
                     if str(it.get('expectedDate', '') or '').strip()),
                    '',
                ),
                '물류센터': next(
                    (str(it.get('logisticsCenter', '') or '').strip()
                     for it in rp_grouped.get(_inv, [])
                     if str(it.get('logisticsCenter', '') or '').strip()),
                    '',
                ),
                '송장번호': _inv,
                '행수': len(rp_grouped.get(_inv, [])),
            }
            for _inv in sorted(not_in_manifest)
        ],
        'not_in_csv': sorted(set(not_in_csv) - set(sku_matched.keys())),
        # 실제로 재출력된 시트 송장번호 — 한진 엑셀을 이 범위로 제한하는 데 사용
        'matched_sheet_invoices': sorted(matched_sheet_invs),
        'ship_to_box_num': rp_ship_to_box_num,
        'new_box_only': new_box_only,
        'sheet_write_result': sheet_write_result,
        'timestamp': datetime.now().strftime('%Y%m%d_%H%M'),
    }


with tab7:
    st.header('🔄 쉽먼트 재출력')
    st.caption('피킹&분류에 로드된 시트 데이터(또는 CSV)의 송장번호와 매니페스트/라벨을 매칭하여 재출력합니다')

    # ── 데이터 소스 모드 선택 ──
    _has_sheet_data = st.session_state.get('pick_df_출고') is not None
    _use_sheet_source = False
    if _has_sheet_data:
        _use_sheet_source = st.toggle(
            '📋 피킹&분류 시트 데이터 사용 (박스번호는 시트 M열에서 보존)',
            value=True,
            key='reprint_use_sheet',
            help='시트에 이미 박스번호가 저장되어 있어서 CSV 업로드 없이도 쉽먼트/라벨만 올리면 됩니다. 발주 취소돼도 박스번호 안 틀어짐.'
        )

    st.divider()

    # ── 🚚 한진택배 엑셀 만들기 (재출력된 송장 + 주소리스트 매칭) ──
    # 아래 업로드/실행 블록에는 st.stop() 경로가 있어서 그 뒤에 두면 렌더링되지 않는다.
    # 재출력 결과(session_state)만 읽으므로 위치는 여기가 안전하다.
    with st.expander(f'🚚 {kit_config.courier_name()} 엑셀 만들기 (물류센터 기준 자동 매칭)', expanded=False):
        st.caption('재출력된 송장을 (물류센터, 송장번호)로 그룹화 → 주소리스트 시트의 물류센터 키로 매칭 → 택배 송장 양식 엑셀 생성 (열 구성은 ⚙️ 설정 탭의 택배 양식)')
        st.caption(f'📋 주소리스트: 설정 탭에서 지정한 시트의 "{kit_config.courier_addr_tab()}" 탭 사용 (E열 키 = 물류센터명)')

        _hj_src = st.session_state.get('reprint_hanjin_src')
        if not _hj_src:
            st.info('먼저 아래에서 **쉽먼트 재출력**을 실행하세요. 실제로 재출력된 송장만 엑셀에 담습니다.')
        else:
            st.markdown(
                f'- 대상: 재출력 매칭 송장 **{_hj_src["n_invoices"]}건** '
                f'(`{_hj_src["label"]}`, {_hj_src["timestamp"]} 실행)'
            )

        if st.button(f'🚚 {kit_config.courier_name()} 엑셀 생성', key='reprint_hanjin_btn', type='primary',
                     use_container_width=True, disabled=not _hj_src):
            try:
                _sclient = st.session_state.get('pick_gsheet_client') or get_gsheet_client()
                _buf, _n_rows, _unmatched, _invalid = _hanjin_build_excel(
                    _hj_src['df'], _sclient, HANJIN_ADDR_SHEET_URL, kit_config.courier_addr_tab(),
                    empty_msg='재출력된 송장이 없습니다 (쉽먼트 재출력을 먼저 실행하세요)',
                    box_num_map=_hj_src.get('box_num_map'),
                )
                st.session_state['reprint_hanjin_result'] = {
                    'buf': _buf.getvalue(),
                    'n_rows': _n_rows,
                    'unmatched': _unmatched,
                    'invalid': _invalid,
                    'ts': datetime.now().strftime('%Y%m%d_%H%M'),
                }
            except Exception as _e:
                st.session_state.pop('reprint_hanjin_result', None)
                st.error(f'❌ {_e}')

        # 결과 표시 (rerun 후에도 유지)
        _rp_hr = st.session_state.get('reprint_hanjin_result')
        if _rp_hr:
            st.success(f'✅ {_rp_hr["n_rows"]}건 생성 완료')
            if _rp_hr.get('invalid'):
                st.error(
                    f'🚨 송장번호 판독 불가 {len(_rp_hr["invalid"])}건 — 엑셀에서 제외됨. '
                    '원본 파일에서 송장번호 열이 과학표기(4.62E+11)로 잘려 저장된 경우입니다. '
                    '해당 열을 텍스트 서식으로 바꿔 다시 받아주세요.'
                )
                with st.expander(f'제외된 송장 {len(_rp_hr["invalid"])}건 보기'):
                    for _iv in _rp_hr['invalid']:
                        st.caption(f'• {_iv}')
            if _rp_hr['unmatched']:
                with st.expander(f'⚠️ 매칭 실패 {len(_rp_hr["unmatched"])}건 (주소리스트에 없는 물류센터)'):
                    for _u in _rp_hr['unmatched']:
                        st.caption(f'• {_u}')
            st.download_button(
                label=f'⬇️ {kit_config.courier_name()} 엑셀 다운로드',
                data=_rp_hr['buf'],
                file_name=f'{kit_config.courier_name()}출고양식_{_rp_hr["ts"]}.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                key='reprint_hanjin_dl',
                use_container_width=True,
            )

    st.divider()

    st.subheader('📂 파일 업로드')
    if _use_sheet_source:
        st.caption('매니페스트/라벨 PDF만 업로드하세요 (CSV 불필요 — 시트 데이터 사용)')
    else:
        st.caption('CSV 1개(필수) + 매니페스트/라벨 PDF를 업로드하세요')

    reprint_files = st.file_uploader(
        '파일 선택 (CSV + PDF)' if not _use_sheet_source else '파일 선택 (PDF)',
        type=(['pdf'] if _use_sheet_source else ['csv', 'xlsx', 'pdf']),
        accept_multiple_files=True,
        key='reprint_files'
    )

    if reprint_files:
        rp_csv = None
        rp_manifest_files = []  # [(파일명, 파일)]
        rp_label_files = []     # [(파일명, 파일)]
        rp_unknown_files = []   # 이름 패턴 불일치로 무시된 파일

        for f in reprint_files:
            fname = f.name.lower()
            if fname.endswith('.csv') or fname.endswith('.xlsx'):
                rp_csv = f
            elif 'manifest' in fname:
                rp_manifest_files.append((f.name, f))
            elif 'label' in fname:
                rp_label_files.append((f.name, f))
            else:
                rp_unknown_files.append(f.name)

        if rp_unknown_files:
            st.warning(
                f'⚠️ 파일명에 `manifest`/`label`이 없어 무시된 파일 {len(rp_unknown_files)}개: '
                + ', '.join(f'`{n}`' for n in rp_unknown_files)
            )

        st.markdown('**분류 결과:**')
        if _use_sheet_source:
            st.markdown(f'- 📋 시트 데이터: `피킹&분류 로드됨 ({len(st.session_state.pick_df_출고)}행)`')
        elif rp_csv:
            st.markdown(f'- CSV: `{rp_csv.name}`')
        else:
            st.error('⚠️ CSV 파일은 필수입니다. 송장번호 매칭에 사용됩니다.')

        st.markdown(f'- 매니페스트 PDF: **{len(rp_manifest_files)}개**')
        st.markdown(f'- 라벨 PDF: **{len(rp_label_files)}개**')

        if not _use_sheet_source and not rp_csv:
            st.stop()

        if not rp_manifest_files and not rp_label_files:
            st.error('매니페스트 또는 라벨 PDF가 필요합니다.')
        else:
            st.divider()

            if st.button('🔄 쉽먼트 재출력 시작', type='primary', key='reprint_btn'):
                rp_progress = st.progress(0)
                rp_status = st.empty()

                try:
                    # 데이터 소스 준비
                    if _use_sheet_source:
                        rp_items = _pick_df_to_items(st.session_state.pick_df_출고)
                    else:
                        rp_csv.seek(0)
                        file_bytes = rp_csv.read()
                        rp_csv.seek(0)  # 같은 업로드로 재실행해도 b'' 가 읽히지 않도록
                        if rp_csv.name.lower().endswith('.xlsx'):
                            rp_items = _parse_xlsx_bytes(file_bytes)
                        else:
                            rp_items = _parse_csv_bytes(file_bytes)

                    # 시트 모드면 M열(기존값) 읽어 보존. 읽기 실패 시 경고 후 중단.
                    _existing_map = None
                    if _use_sheet_source:
                        _existing_map = dict(st.session_state.get('pick_ship_to_box') or {})
                        if not _existing_map and st.session_state.get('pick_gsheet_client'):
                            _read = pick_read_box_numbers(
                                st.session_state.pick_gsheet_client,
                                st.session_state.get('pick_sheet_url_출고', ''),
                                st.session_state.get('pick_sheet_tab_출고', ''),
                            )
                            if _read is None:
                                st.error('❌ 시트 M열 읽기 실패. 새로고침 후 다시 시도하세요.')
                                st.stop()
                            _existing_map = _read

                    result = _run_reprint_pipeline(
                        rp_items, rp_manifest_files, rp_label_files,
                        existing_box_map=_existing_map,
                        write_new_to_sheet=_use_sheet_source,
                        sheet_client=st.session_state.get('pick_gsheet_client'),
                        sheet_url=st.session_state.get('pick_sheet_url_출고', ''),
                        sheet_tab=st.session_state.get('pick_sheet_tab_출고', ''),
                        on_progress=rp_progress.progress,
                        on_status=rp_status.text,
                    )

                    if 'error' in result:
                        st.error(f'❌ {result["error"]}')
                        st.stop()

                    # 시트 세션 캐시 업데이트 — 매칭 안 된 기존 송장 번호를 지우지 않도록 MERGE
                    if _use_sheet_source and result.get('ship_to_box_num'):
                        _new_map = result['ship_to_box_num']
                        _merged_ship = dict(st.session_state.get('pick_ship_to_box') or {})
                        _merged_ship.update(_new_map)
                        st.session_state['pick_ship_to_box'] = _merged_ship
                        _sync_url = st.session_state.get('pick_sheet_url_출고', '')
                        _sync_tab = st.session_state.get('pick_sheet_tab_출고', '')
                        if _sync_url and _sync_tab:
                            _cache_key_sync = f"_pick_existing_box_{_sync_url}_{_sync_tab}"
                            _merged_cache = dict(st.session_state.get(_cache_key_sync) or {})
                            _merged_cache.update(_new_map)
                            st.session_state[_cache_key_sync] = _merged_cache
                    if result.get('sheet_write_result') == -1:
                        st.warning('⚠️ 시트 M열 쓰기 실패 — 박스번호가 시트에 저장되지 않았습니다. 다시 시도하세요.')

                    col_m1, col_m2, col_m3 = st.columns(3)
                    with col_m1:
                        st.metric('매칭됨', f'{result["matched"]}건')
                    with col_m2:
                        st.metric('데이터에만 존재', f'{len(result["not_in_manifest"])}건')
                    with col_m3:
                        st.metric('쉽먼트에만 존재', f'{len(result["not_in_csv"])}건')

                    if result['not_in_manifest']:
                        with st.expander(f'데이터에만 있는 송장 ({len(result["not_in_manifest"])}건) - 매니페스트 없음'):
                            if result.get('not_in_manifest_detail'):
                                import pandas as _pd_nim
                                st.dataframe(_pd_nim.DataFrame(result['not_in_manifest_detail']),
                                             use_container_width=True, hide_index=True)
                            st.caption('각 줄 오른쪽 📋 아이콘으로 개별 복사 (맨 위는 전체 복사용)')
                            st.code('\n'.join(result['not_in_manifest']))
                            for _inv in result['not_in_manifest']:
                                st.code(_inv, language=None)
                    if result['not_in_csv']:
                        with st.expander(f'쉽먼트에만 있는 송장 ({len(result["not_in_csv"])}건) - 이번에 미출력'):
                            st.caption('각 줄 오른쪽 📋 아이콘으로 개별 복사 (맨 위는 전체 복사용)')
                            st.code('\n'.join(result['not_in_csv']))
                            for _inv in result['not_in_csv']:
                                st.code(_inv, language=None)

                    # 한진 엑셀은 '실제로 재출력된 송장'만 대상으로 한다.
                    # (전체 데이터를 쓰면 매니페스트 없는 송장까지 택배 행으로 나감)
                    _hj_matched = set(result.get('matched_sheet_invoices') or [])
                    _hj_items = [
                        it for it in rp_items
                        if str(it.get('shipmentNumber', '') or '').strip() in _hj_matched
                    ]
                    st.session_state['reprint_hanjin_src'] = {
                        'df': _items_to_hanjin_df(_hj_items),
                        'label': ('피킹&분류 시트 데이터' if _use_sheet_source
                                  else f'업로드 파일 — {rp_csv.name}'),
                        'n_invoices': len(_hj_matched),
                        'timestamp': result['timestamp'],
                        # 택배 엑셀 행 순서를 지시서/쉽먼트 출력과 같은 박스번호 순으로
                        # 맞추기 위해 재출력이 실제로 부여한 맵을 그대로 넘긴다.
                        # (한진 df에는 박스번호 컬럼이 없어 여기서 주지 않으면 못 맞춘다)
                        'box_num_map': result.get('ship_to_box_num') or {},
                    }
                    # 재출력을 다시 돌렸으면 이전 엑셀은 무효
                    st.session_state.pop('reprint_hanjin_result', None)

                    st.session_state.reprint_result = {
                        'final_bytes': result['final_bytes'],
                        'shipment_only_bytes': result['shipment_only_bytes'],
                        'so_only_bytes': result['so_only_bytes'],
                        'total': result['total'],
                        'shipment_total': result['shipment_total'],
                        'so_total': result['so_total'],
                        'matched': result['matched'],
                        'timestamp': result['timestamp'],
                    }

                except Exception as e:
                    st.error(f'❌ 오류 발생: {e}')
                    import traceback
                    st.code(traceback.format_exc())

        # ===== 결과 표시 (session_state에서 읽어 버튼 유지) =====
        if 'reprint_result' in st.session_state:
            res = st.session_state.reprint_result
            st.divider()
            st.subheader('📋 재출력 결과')
            st.markdown(f"""
| 구분 | 수량 |
|------|------|
| 매칭된 송장 | {res['matched']}건 |
| 출고지시서 | {res['so_total']}p |
| 쉽먼트 (매니페스트+라벨) | {res['shipment_total']}p |
| **전체 합계** | **{res['total']}p** |
""")
            st.caption('순서: [출고지시서→매니페스트→라벨] 매칭 송장번호순 통합 배치')

            st.divider()

            rp_col_a, rp_col_b, rp_col_c = st.columns(3)
            with rp_col_a:
                st.download_button(
                    label=f"⬇️ 전체 통합 PDF ({res['total']}p)",
                    data=res['final_bytes'],
                    file_name=f"shipment_reprint_ALL_{res['timestamp']}.pdf",
                    mime='application/pdf',
                    key='reprint_dl',
                    type='primary',
                    use_container_width=True,
                )
            with rp_col_b:
                st.download_button(
                    label=f"⬇️ 쉽먼트만 ({res['shipment_total']}p)",
                    data=res['shipment_only_bytes'],
                    file_name=f"shipment_reprint_shipment_{res['timestamp']}.pdf",
                    mime='application/pdf',
                    key='reprint_dl_shipment',
                    use_container_width=True,
                )
            with rp_col_c:
                st.download_button(
                    label=f"⬇️ 출고지시서만 ({res['so_total']}p)",
                    data=res['so_only_bytes'],
                    file_name=f"shipment_reprint_so_{res['timestamp']}.pdf",
                    mime='application/pdf',
                    key='reprint_dl_so',
                    use_container_width=True,
                )


# ══════════════════════════════════════════════════════
# 탭8: 피킹 검증 시스템
# ══════════════════════════════════════════════════════
with tab8:
    import pandas as _pd

    st.header('📦 피킹 & 분류')
    st.caption('하나의 시트로 피킹검증 또는 입고분류를 모드 전환하며 사용')

    # ── 음성(TTS) 프로필 설정 ──
    # 브라우저 한국어 voice는 여자뿐 + pitch 무시 → 진짜 남자는 영어 voice 사용
    # 프로필: (display, lang, voice_gender, pitch, rate_offset, tone_freq, tone_pattern, hint)
    # tuple: (display, lang, gender, pitch, extra_rate, tone_freq, tone_pattern, description, voice_hint, edge_voice)
    # edge_voice가 있으면 서버에서 Microsoft 신경망 TTS(무료)로 mp3 생성 — 한국어 남자/다양한 여자 가능!
    _TTS_PROFILES = [
        ("👩 여자 한국어 (선희)",   'ko', 'female', 1.0, 0.0,   1500, 'ding-ding',    "Microsoft 신경망 한국어 여자 (SunHi) — 또렷한 톤",   '',  'ko-KR-SunHiNeural'),
        ("👩 여자 한국어 (지민)",   'ko', 'female', 1.0, 0.0,   880,  'do-re-mi',     "Microsoft 신경망 한국어 여자 (JiMin) — 부드러운 톤", '',  'ko-KR-JiMinNeural'),
        ("👨 남자 한국어 (인준)",   'ko', 'male',   1.0, 0.0,   220,  'bzz-bzz',      "진짜 한국어 남자 음성 (Microsoft 인준, 인터넷 필요)", '',  'ko-KR-InJoonNeural'),
        ("👨 남자 한국어 (현수)",   'ko', 'male',   1.0, 0.0,   330,  'low-low-high', "진짜 한국어 남자 음성 (Microsoft 현수, 인터넷 필요)", '',  'ko-KR-HyunsuMultilingualNeural'),
    ]
    _profile_labels = [p[0] for p in _TTS_PROFILES]
    _profile_map = {p[0]: p for p in _TTS_PROFILES}

    try:
        # 호환성: 이전 세션 잘못된 값 방어
        if st.session_state.get('pick_tts_profile_ui') not in _profile_labels:
            st.session_state.pop('pick_tts_profile_ui', None)
        _rate_val = st.session_state.get('pick_tts_rate', 1.0)
        if not (isinstance(_rate_val, (int, float)) and 0.8 <= float(_rate_val) <= 1.5):
            st.session_state.pop('pick_tts_rate', None)

        _tts_col1, _tts_col2, _tts_col3 = st.columns([2, 1, 1])
        with _tts_col1:
            _tts_profile = st.radio(
                "🔊 음성 프로필 (양쪽에서 다른 프로필 선택 → 명확히 구분)",
                _profile_labels,
                index=0, key="pick_tts_profile_ui",
                horizontal=False,
                help="두 노트북에서 다른 프로필을 선택하면 톤이 확실히 달라짐",
            )
            _sel = _profile_map[_tts_profile]
            st.session_state['pick_tts_lang'] = _sel[1]
            st.session_state['pick_tts_gender'] = _sel[2]
            st.session_state['pick_tts_pitch'] = _sel[3]
            st.session_state['pick_tts_extra_rate'] = _sel[4]
            st.session_state['pick_tts_tone_freq'] = _sel[5]
            st.session_state['pick_tts_tone_pattern'] = _sel[6]
            st.session_state['pick_tts_voice_hint'] = _sel[8]
            _prev_edge = st.session_state.get('pick_tts_edge_voice')
            st.session_state['pick_tts_edge_voice'] = _sel[9]
            # 남자 한국어 프로필 첫 선택 시 자주 쓰는 문구 미리 생성 (첫 스캔 지연 방지)
            if _sel[9] and _sel[9] != _prev_edge:
                try:
                    _edge_tts_warmup(_sel[9])
                except Exception:
                    pass
            st.caption(f'💡 {_sel[7]}')
        with _tts_col2:
            _tts_rate = st.slider(
                '🗣️ 속도',
                min_value=0.7, max_value=1.5, value=1.0, step=0.05,
                key='pick_tts_rate',
                help="느릴수록 또렷함",
            )
        with _tts_col3:
            st.markdown('<br>', unsafe_allow_html=True)
            if st.button('🔊 미리듣기', key='pick_tts_test', use_container_width=True):
                st.session_state['pick_tts_test_pending'] = True
            if st.button('🔄 기본값', key='pick_tts_reset', use_container_width=True,
                         help='프로필/속도 초기화'):
                for _k in ['pick_tts_profile_ui', 'pick_tts_rate']:
                    st.session_state.pop(_k, None)
                st.rerun()
    except Exception as _tts_ui_err:
        st.warning(f'⚠️ TTS 설정 UI 초기화 중 문제: {_tts_ui_err}. 페이지 새로고침하세요.')

    # 음성 이름 직접 선택기 (브라우저 voice 목록 표시)
    with st.expander('⚙️ 음성 고급 설정 — 내 브라우저에 설치된 voice 직접 선택', expanded=False):
        st.caption('"남자 선택해도 여자 목소리"로 들린다면 이 드롭다운에서 남자 이름의 voice를 직접 고르세요. 목록은 브라우저/OS에 따라 다릅니다.')
        from streamlit.components.v1 import html as _voice_pick_html
        _voice_pick_html("""
        <div style="padding:8px;background:#f3f4f6;border-radius:8px;font-family:system-ui">
          <div style="font-size:12px;color:#555;margin-bottom:6px">내 브라우저에 설치된 음성 목록:</div>
          <select id="__voice_select" style="width:100%;padding:6px;font-size:14px" onchange="__saveVoice()">
            <option value="">(자동 선택 — 성별 라디오 기준)</option>
          </select>
          <button onclick="__testVoice()" style="margin-top:6px;padding:4px 12px">🔊 이 음성 테스트</button>
          <span id="__voice_status" style="margin-left:10px;color:#1a56db;font-size:12px"></span>
        </div>
        <script>
        (function(){
            const win = window.parent;
            const doc = win.document;
            function listVoices() {
                const sel = doc.getElementById('__voice_select');
                if (!sel) return;
                const voices = win.speechSynthesis.getVoices();
                const saved = win.localStorage.getItem('__preferred_voice') || '';
                sel.innerHTML = '<option value="">(자동 선택 — 성별 라디오 기준)</option>';
                voices.forEach(function(v){
                    const opt = doc.createElement('option');
                    opt.value = v.name;
                    const genderHint = /injoon|minsu|male|man|david|mark|guy/i.test(v.name) ? ' 👨' :
                                       /heami|yuna|female|woman|zira|jenny/i.test(v.name) ? ' 👩' : '';
                    opt.textContent = v.name + ' [' + v.lang + ']' + genderHint;
                    if (v.name === saved) opt.selected = true;
                    sel.appendChild(opt);
                });
            }
            win.__saveVoice = function(){
                const sel = doc.getElementById('__voice_select');
                if (!sel) return;
                win.localStorage.setItem('__preferred_voice', sel.value || '');
                const status = doc.getElementById('__voice_status');
                if (status) status.textContent = sel.value ? '✅ 저장됨: ' + sel.value : '자동 선택 모드';
            };
            win.__testVoice = function(){
                const sel = doc.getElementById('__voice_select');
                const name = sel ? sel.value : '';
                const u = new SpeechSynthesisUtterance('안녕하세요. 이 음성으로 재고완료. 다시 찍어주세요.');
                u.lang = 'ko-KR';
                const voices = win.speechSynthesis.getVoices();
                if (name) {
                    const v = voices.find(function(x){ return x.name === name; });
                    if (v) u.voice = v;
                } else {
                    const v = voices.find(function(x){ return (x.lang||'').indexOf('ko') === 0; });
                    if (v) u.voice = v;
                }
                win.speechSynthesis.cancel();
                win.speechSynthesis.speak(u);
            };
            // Voices may load asynchronously
            if (win.speechSynthesis.getVoices().length > 0) {
                listVoices();
            }
            win.speechSynthesis.onvoiceschanged = listVoices;
            setTimeout(listVoices, 300);
            setTimeout(listVoices, 800);
        })();
        </script>
        """, height=150)

    # 음성 테스트 실행 (설정 바꾼 직후 들어보기 용)
    if st.session_state.get('pick_tts_test_pending'):
        from streamlit.components.v1 import html as _tts_test_html
        _gender_label = '남자입니다' if st.session_state.get('pick_tts_gender') == 'male' else '여자입니다'
        _tts_test_html(
            f"<script>{_tts_ko_script(f'안녕하세요. {_gender_label}. 재고완료. 다시 찍어주세요')}</script>",
            height=0,
        )
        st.session_state['pick_tts_test_pending'] = False

    # ── 데이터 소스 선택 ──
    pick_mode = st.radio(
        "데이터 소스",
        ["📊 구글 시트 (실시간)", "📂 CSV 파일 업로드"],
        index=0, key="pick_mode", horizontal=True,
    )

    if pick_mode == "📊 구글 시트 (실시간)":
        # 기본 고정 시트 (운영용)
        # 기본값은 ⚙️ 설정 탭의 피킹 시트 (각자 입력)
        _DEFAULT_SHEET_URL = kit_config.pick_ship_sheet_url()
        _DEFAULT_DAPAE_URL = kit_config.pick_dapae_sheet_url()
        # 페이지 새로고침 시에도 유지되도록 query_params에서 복원
        _qp = st.query_params
        if 'pick_url_출고' not in st.session_state:
            st.session_state['pick_url_출고'] = _DEFAULT_SHEET_URL or _qp.get('pu') or ''
        if 'pick_tab_출고' not in st.session_state:
            st.session_state['pick_tab_출고'] = kit_config.pick_ship_tab() or _qp.get('pt') or '출고확인'
        if 'pick_url_배대지' not in st.session_state:
            st.session_state['pick_url_배대지'] = _DEFAULT_DAPAE_URL or _qp.get('bu') or ''
        if 'pick_tab_배대지' not in st.session_state:
            st.session_state['pick_tab_배대지'] = kit_config.pick_dapae_tab() or _qp.get('bt') or '배대지입고리스트'

        st.markdown("##### 쉽먼트 시트 (출고지시서)")
        st.caption("기본 주소는 ⚙️ 설정 탭에서 저장합니다. 여기서 바꾸면 이번 화면에서만 적용됩니다.")
        gs_col1, gs_col2 = st.columns([3, 1])
        with gs_col1:
            url_출고 = st.text_input("구글 시트 URL", placeholder="https://docs.google.com/spreadsheets/d/...", key="pick_url_출고")
        with gs_col2:
            tab_출고 = st.text_input("탭 이름", key="pick_tab_출고")

        st.markdown("##### 배대지 입고 시트 (선택)")
        gs_col3, gs_col4 = st.columns([3, 1])
        with gs_col3:
            url_배대지 = st.text_input("구글 시트 URL", placeholder="비워두면 같은 시트 사용", key="pick_url_배대지")
        with gs_col4:
            tab_배대지 = st.text_input("탭 이름", key="pick_tab_배대지")

        # 입력값을 query_params에 저장 (새로고침 후 유지)
        if url_출고:
            st.query_params['pu'] = url_출고
            st.query_params['pt'] = tab_출고
        if url_배대지:
            st.query_params['bu'] = url_배대지
            st.query_params['bt'] = tab_배대지

        # 배대지 URL 비어있으면 출고지시서 URL 사용
        if not url_배대지.strip() and url_출고.strip():
            url_배대지 = url_출고

        if st.button("🔄 구글 시트 연결", use_container_width=True, key="pick_gsheet_btn"):
            if not url_출고.strip():
                st.error("쉽먼트 시트 URL을 입력해주세요")
            else:
                with st.spinner("구글 시트 연결 중..."):
                    success = pick_load_all_data(url_출고, tab_출고, url_배대지, tab_배대지)
                    if success:
                        pick_init_inventory()
                        st.success("✅ 구글 시트 연결 완료!")
                        st.rerun()
                    else:
                        st.error("연결 실패 — URL과 탭 이름을 확인하세요")
    else:
        pc1, pc2 = st.columns(2)
        with pc1:
            pick_csv_출고 = st.file_uploader("출고지시서 CSV", type=["csv"], key="pick_csv_출고")
        with pc2:
            pick_csv_배대지 = st.file_uploader("배대지 입고 CSV (선택)", type=["csv"], key="pick_csv_배대지")
        # 업로드된 파일이 바뀌었을 때만 다시 읽는다. 예전에는 rerun 마다 CSV 를 다시 파싱해
        # 새 DataFrame 을 만들었고, 그 때문에 버튼 하나만 눌러도 입고분류 진행 상태가
        # (데이터가 바뀐 줄 알고) 초기화됐다.
        def _upload_id(f):
            return getattr(f, 'file_id', None) or f'{f.name}:{getattr(f, "size", 0)}'

        if pick_csv_출고:
            _csv_id = _upload_id(pick_csv_출고)
            if (st.session_state.get('_pick_csv_출고_id') != _csv_id
                    or st.session_state.pick_df_출고 is None):
                pick_csv_출고.seek(0)
                df = _pd.read_csv(pick_csv_출고, encoding="utf-8-sig")
                st.session_state.pick_df_출고 = pick_clean_출고(df)
                st.session_state['_pick_csv_출고_id'] = _csv_id
                if st.session_state.pick_df_출고 is not None:
                    st.session_state.pick_data_loaded = True
                    st.session_state['pick_data_ver'] = st.session_state.get('pick_data_ver', 0) + 1
        if pick_csv_배대지:
            _dp_id = _upload_id(pick_csv_배대지)
            if (st.session_state.get('_pick_csv_배대지_id') != _dp_id
                    or st.session_state.pick_df_배대지 is None):
                pick_csv_배대지.seek(0)
                df = _pd.read_csv(pick_csv_배대지, encoding="utf-8-sig")
                st.session_state.pick_df_배대지 = pick_clean_배대지(df)
                st.session_state['_pick_csv_배대지_id'] = _dp_id
                if st.session_state.pick_data_loaded and st.session_state.pick_df_배대지 is not None:
                    pick_init_inventory()

    # ── 데이터 상태 표시 ──
    if st.session_state.pick_df_출고 is not None:
        n_rows = len(st.session_state.pick_df_출고)
        n_ship = st.session_state.pick_df_출고["쉽먼트운송장번호"].nunique()
        st.success(f"출고지시서: {n_rows}행 / {n_ship}개 쉽먼트")
    if st.session_state.pick_df_배대지 is not None:
        st.success(f"배대지 입고: {len(st.session_state.pick_df_배대지)}행 로드됨")

    # ── 시트 저장 상태 (쓰기는 2초마다 묶여서 전송된다) ──
    if st.session_state.get('pick_use_gsheet'):
        _pending = gs_pending_count()
        _sc1, _sc2 = st.columns([4, 1])
        with _sc1:
            if _pending:
                st.caption(f"⏳ 시트 저장 대기 {_pending}건 — 잠시 후 자동 반영됩니다")
            else:
                st.caption("✅ 시트에 모두 저장됨")
        with _sc2:
            if st.button("💾 지금 저장", key="pick_flush_now", use_container_width=True,
                         help="큐에 남은 시트 쓰기를 즉시 전송합니다"):
                gs_flush_pending()
                st.rerun()

    # ── 데이터 없으면 가이드 ──
    if not st.session_state.pick_data_loaded:
        st.info("위에서 데이터를 연결하세요 (구글 시트 또는 CSV)")
        st.markdown("""
**구글 시트 모드:**
1. 구글 시트 URL을 위 입력칸에 붙여넣기
2. 탭 이름을 정확히 입력 (예: 출고확인, 배대지입고리스트)
3. '구글 시트 연결' 클릭

**CSV 모드:**
1. 'CSV 파일 업로드' 선택
2. 출고지시서 CSV 업로드 (필수)
3. 배대지 입고 CSV 업로드 (선택)
        """)
        st.stop()

    # ── 현재 로드된 데이터 상태 (탭 변경 시 이전 데이터로 작업하는 실수 방지) ──
    if st.session_state.get('pick_use_gsheet') and st.session_state.pick_df_출고 is not None:
        _cur_tab = st.session_state.get('pick_sheet_tab_출고', '') or '(알 수 없음)'
        _cur_rows = len(st.session_state.pick_df_출고)
        _cur_ships = st.session_state.pick_df_출고['쉽먼트운송장번호'].nunique() if '쉽먼트운송장번호' in st.session_state.pick_df_출고.columns else 0
        st.info(f"📌 현재 로드된 데이터: **'{_cur_tab}'** 탭 — {_cur_rows}행 / 송장 {_cur_ships}건. 탭을 바꿨다면 반드시 '구글 시트 연결'을 다시 누르세요.")

    # ── 출고지시서 재출력 (피킹 시작 전에 박스번호 부여 + M열 기록) ──
    st.divider()
    with st.expander("📄 출고지시서 재출력 (쉽먼트/라벨 PDF 업로드 → 박스번호 부여)", expanded=False):
        st.caption("피킹 시작 전에 쉽먼트/라벨 PDF를 업로드하면, 현재 시트 송장과 매칭해 출고지시서 PDF를 만들고 박스번호를 시트 M열에 저장합니다. 기존에 M열에 값이 있으면 그대로 유지(발주 취소 내성). ⚠️ 동일 시트를 여러 사용자가 동시에 재출력하지 마세요 — 박스번호 충돌 가능.")

        pick_reprint_files = st.file_uploader(
            '매니페스트/라벨 PDF (파일명에 manifest/label 포함)',
            type=['pdf'],
            accept_multiple_files=True,
            key='pick_reprint_files'
        )

        if pick_reprint_files:
            pk_manifest_files = []
            pk_label_files = []
            pk_unknown_files = []
            for f in pick_reprint_files:
                fname = f.name.lower()
                if 'manifest' in fname:
                    pk_manifest_files.append((f.name, f))
                elif 'label' in fname:
                    pk_label_files.append((f.name, f))
                else:
                    pk_unknown_files.append(f.name)

            if pk_unknown_files:
                st.warning(
                    f'⚠️ 파일명에 `manifest`/`label`이 없어 무시된 파일 {len(pk_unknown_files)}개: '
                    + ', '.join(f'`{n}`' for n in pk_unknown_files)
                )

            st.markdown(f'- 매니페스트 PDF: **{len(pk_manifest_files)}개** / 라벨 PDF: **{len(pk_label_files)}개**')

            if not pk_manifest_files and not pk_label_files:
                st.warning('매니페스트 또는 라벨 PDF가 필요합니다. 파일명에 `manifest` 또는 `label`을 포함시켜 주세요.')
            elif st.button('🔄 출고지시서 재출력 시작', type='primary', key='pick_reprint_btn'):
                _pk_progress = st.progress(0)
                _pk_status = st.empty()
                try:
                    _pk_items = _pick_df_to_items(st.session_state.pick_df_출고)
                    _pk_existing = dict(st.session_state.get('pick_ship_to_box') or {})
                    if (not _pk_existing
                            and st.session_state.get('pick_gsheet_client')
                            and st.session_state.get('pick_sheet_url_출고')):
                        _pk_read = pick_read_box_numbers(
                            st.session_state.pick_gsheet_client,
                            st.session_state.pick_sheet_url_출고,
                            st.session_state.pick_sheet_tab_출고,
                        )
                        if _pk_read is None:
                            st.error('❌ 시트 M열 읽기 실패. 새로고침 후 다시 시도하세요.')
                            st.stop()
                        _pk_existing = _pk_read

                    _pk_result = _run_reprint_pipeline(
                        _pk_items, pk_manifest_files, pk_label_files,
                        existing_box_map=_pk_existing,
                        write_new_to_sheet=bool(st.session_state.get('pick_use_gsheet')),
                        sheet_client=st.session_state.get('pick_gsheet_client'),
                        sheet_url=st.session_state.get('pick_sheet_url_출고', ''),
                        sheet_tab=st.session_state.get('pick_sheet_tab_출고', ''),
                        on_progress=_pk_progress.progress,
                        on_status=_pk_status.text,
                    )

                    if 'error' in _pk_result:
                        st.error(f'❌ {_pk_result["error"]}')
                    else:
                        if _pk_result.get('ship_to_box_num'):
                            # 재출력 결과는 매칭된 송장만 포함하므로 기존 값에 MERGE (덮어쓰기 X)
                            _new_map = _pk_result['ship_to_box_num']
                            _merged_ship = dict(st.session_state.get('pick_ship_to_box') or {})
                            _merged_ship.update(_new_map)
                            st.session_state['pick_ship_to_box'] = _merged_ship
                            # 입고분류 모드 캐시도 merge 방식으로 동기화
                            _sync_url = st.session_state.get('pick_sheet_url_출고', '')
                            _sync_tab = st.session_state.get('pick_sheet_tab_출고', '')
                            if _sync_url and _sync_tab:
                                _cache_key_sync = f"_pick_existing_box_{_sync_url}_{_sync_tab}"
                                _merged_cache = dict(st.session_state.get(_cache_key_sync) or {})
                                _merged_cache.update(_new_map)
                                st.session_state[_cache_key_sync] = _merged_cache
                        if _pk_result.get('sheet_write_result') == -1:
                            st.warning('⚠️ 시트 M열 쓰기 실패 — 박스번호가 시트에 저장되지 않았습니다.')
                        _n_new = len(_pk_result.get('new_box_only') or {})
                        _n_reuse = _pk_result['matched'] - _n_new
                        _n_inv_match = _pk_result.get('invoice_matched', _pk_result['matched'])
                        _sku_map = _pk_result.get('sku_matched') or {}
                        _n_sku_match = len(_sku_map)
                        _msg = (
                            f'✅ 재출력 완료 — 매칭 {_pk_result["matched"]}건 '
                            f'(송장 {_n_inv_match}건 / SKU 매칭 {_n_sku_match}건). '
                            f'박스번호: 신규 {_n_new}건 / 기존 {_n_reuse}건 재사용'
                        )
                        st.success(_msg)

                        if _sku_map:
                            with st.expander(f'⚠️ SKU 기반 매칭 {_n_sku_match}건 — 송장번호 불일치 (PDF↔시트)', expanded=True):
                                st.caption('매니페스트 송장번호와 시트 송장번호가 다르지만, SKU(바코드) 구성이 같아서 자동 매칭됨. 시트의 송장번호를 매니페스트 값으로 갱신 권장.')
                                st.dataframe(
                                    [{'매니페스트 송장(=실제)': k, '시트 송장(현재)': v} for k, v in _sku_map.items()],
                                    use_container_width=True, hide_index=True,
                                )

                        _c1, _c2, _c3 = st.columns(3)
                        with _c1:
                            st.metric('매칭됨', f'{_pk_result["matched"]}건')
                        with _c2:
                            st.metric('데이터에만', f'{len(_pk_result["not_in_manifest"])}건')
                        with _c3:
                            st.metric('쉽먼트에만', f'{len(_pk_result["not_in_csv"])}건')

                        if _pk_result['not_in_manifest']:
                            with st.expander(f'데이터에만 있는 송장 ({len(_pk_result["not_in_manifest"])}건)'):
                                if _pk_result.get('not_in_manifest_detail'):
                                    import pandas as _pd_nim2
                                    st.dataframe(_pd_nim2.DataFrame(_pk_result['not_in_manifest_detail']),
                                                 use_container_width=True, hide_index=True)
                                st.caption('각 줄 오른쪽 📋 아이콘으로 개별 복사 (맨 위는 전체 복사용)')
                                st.code('\n'.join(_pk_result['not_in_manifest']))
                                for _inv in _pk_result['not_in_manifest']:
                                    st.code(_inv, language=None)
                        if _pk_result['not_in_csv']:
                            with st.expander(f'쉽먼트에만 있는 송장 ({len(_pk_result["not_in_csv"])}건)'):
                                st.caption('각 줄 오른쪽 📋 아이콘으로 개별 복사 (맨 위는 전체 복사용)')
                                st.code('\n'.join(_pk_result['not_in_csv']))
                                for _inv in _pk_result['not_in_csv']:
                                    st.code(_inv, language=None)

                        # 이번 재출력에 실제로 포함된 송장의 박스번호만 추림.
                        # (시트 M열 전체를 쓰면 이번에 안 찍는 박스 번호표까지 나온다)
                        _pk_boxmap = _pk_result.get('ship_to_box_num') or {}
                        _pk_scope = set(_pk_result.get('matched_sheet_invoices') or [])
                        _pk_box_nums = sorted({
                            int(v) for k, v in _pk_boxmap.items()
                            if (not _pk_scope or k in _pk_scope) and str(v).strip().isdigit()
                        })
                        _pk_new_nums = sorted({
                            int(v) for v in (_pk_result.get('new_box_only') or {}).values()
                            if str(v).strip().isdigit()
                        })

                        st.session_state['pick_reprint_result'] = {
                            'final_bytes': _pk_result['final_bytes'],
                            'shipment_only_bytes': _pk_result['shipment_only_bytes'],
                            'so_only_bytes': _pk_result['so_only_bytes'],
                            'total': _pk_result['total'],
                            'shipment_total': _pk_result['shipment_total'],
                            'so_total': _pk_result['so_total'],
                            'matched': _pk_result['matched'],
                            'timestamp': _pk_result['timestamp'],
                            'box_nums': _pk_box_nums,
                            'new_box_nums': _pk_new_nums,
                        }
                except Exception as e:
                    st.error(f'❌ 오류: {e}')
                    import traceback
                    st.code(traceback.format_exc())

        # 결과 다운로드 버튼 (rerun 후에도 유지)
        if 'pick_reprint_result' in st.session_state:
            _pres = st.session_state['pick_reprint_result']
            st.divider()
            st.markdown(
                f"**결과**: 출고지시서 {_pres['so_total']}p / "
                f"쉽먼트 {_pres['shipment_total']}p / 전체 {_pres['total']}p"
            )
            _dc1, _dc2, _dc3 = st.columns(3)
            with _dc1:
                st.download_button(
                    f"⬇️ 전체 통합 PDF ({_pres['total']}p)",
                    data=_pres['final_bytes'],
                    file_name=f"shipment_reprint_ALL_{_pres['timestamp']}.pdf",
                    mime='application/pdf',
                    key='pick_reprint_dl_all',
                    type='primary',
                    use_container_width=True,
                )
            with _dc2:
                st.download_button(
                    f"⬇️ 쉽먼트만 ({_pres['shipment_total']}p)",
                    data=_pres['shipment_only_bytes'],
                    file_name=f"shipment_reprint_shipment_{_pres['timestamp']}.pdf",
                    mime='application/pdf',
                    key='pick_reprint_dl_ship',
                    use_container_width=True,
                )
            with _dc3:
                st.download_button(
                    f"⬇️ 출고지시서만 ({_pres['so_total']}p)",
                    data=_pres['so_only_bytes'],
                    file_name=f"출고지시서_{_pres['timestamp']}.pdf",
                    mime='application/pdf',
                    key='pick_reprint_dl_so',
                    use_container_width=True,
                )

            # ── 박스번호 번호표 PDF (오려서 박스에 붙이는 용도) ──
            _all_nums = _pres.get('box_nums') or []
            _new_nums = _pres.get('new_box_nums') or []
            st.divider()
            st.markdown('#### 🔢 박스번호 번호표 PDF')
            if not _all_nums:
                st.caption('부여된 박스번호가 없습니다 (전부 국내재고/부족이면 번호가 안 붙습니다).')
            else:
                st.caption(
                    f'부여된 박스번호 **{len(_all_nums)}개** (신규 {len(_new_nums)}개). '
                    'A4에 격자로 찍히니 잘라서 박스에 붙이면 됩니다 (기본 5×7 = 한 장에 35개).'
                )
                _bc1, _bc2, _bc3 = st.columns(3)
                with _bc1:
                    _box_scope = st.radio(
                        '범위',
                        options=['전체', '신규만'],
                        horizontal=True,
                        key='pick_boxnum_scope',
                        help='신규만 = 이번에 새로 부여된 박스번호만 (이미 붙인 번호표는 다시 안 뽑음)',
                    )
                with _bc2:
                    _box_cols = st.number_input(
                        '가로 칸수', min_value=1, max_value=8, value=5, step=1,
                        key='pick_boxnum_cols',
                    )
                with _bc3:
                    _box_rows = st.number_input(
                        '세로 칸수', min_value=1, max_value=12, value=7, step=1,
                        key='pick_boxnum_rows',
                    )

                _target_nums = _new_nums if _box_scope == '신규만' else _all_nums
                if not _target_nums:
                    st.info('이번 재출력에서 새로 부여된 박스번호가 없습니다 (전부 기존 번호 재사용).')
                else:
                    _per_page = int(_box_cols) * int(_box_rows)
                    _pages = -(-len(_target_nums) // _per_page)
                    try:
                        _box_pdf = create_box_number_pdf(
                            _target_nums, cols=int(_box_cols), rows=int(_box_rows),
                        )
                    except Exception as _e:
                        _box_pdf = b''
                        st.error(f'❌ 번호표 PDF 생성 실패: {_e}')
                    if _box_pdf:
                        st.download_button(
                            f'⬇️ 박스번호 PDF ({len(_target_nums)}개 / {_pages}p)',
                            data=_box_pdf,
                            file_name=f"박스번호_{_target_nums[0]}-{_target_nums[-1]}_{_pres['timestamp']}.pdf",
                            mime='application/pdf',
                            key='pick_boxnum_dl',
                            type='primary',
                            use_container_width=True,
                        )

    # ── 시트 송장-상품 ↔ 동봉문서(매니페스트) 일치 검증 ──
    # 출고확인 시트는 사용자가 직접 송장을 분류해서 만든 것이라 송장에 엉뚱한 상품이
    # 들어갈 수 있음. 동봉문서(=쉽먼트의 정답)와 대조해서 잘못 분류된 상품을 찾음.
    st.divider()
    with st.expander("🔍 시트 송장-상품 ↔ 동봉문서 일치 검증", expanded=False):
        st.caption(
            "출고확인 시트의 송장-상품 매핑이 실제 매니페스트(동봉문서)와 일치하는지 검증합니다. "
            "**시트 상품 중 매니페스트에 없는 것은 🚨 오분류(엉뚱한 송장에 배정)**. "
            "매니페스트에 있는데 시트에 없는 것은 아직 도착 안 한 상품일 수 있어 참고용."
        )

        _df_pick = st.session_state.pick_df_출고
        _sku_col = None
        for _c in _df_pick.columns:
            _cn = str(_c).strip().upper().replace(' ', '')
            if _cn in ('SKUID', 'SKU_ID', 'SKU'):
                _sku_col = _c
                break
        if _sku_col is None:
            st.warning('⚠️ 시트에 `SKU ID` 컬럼(E열)이 없습니다. 시트 헤더를 확인하세요.')
        else:
            verify_files = st.file_uploader(
                '매니페스트 PDF (여러 개 가능)',
                type=['pdf'],
                accept_multiple_files=True,
                key='verify_manifest_files',
            )

            if verify_files and st.button('🔍 검증 시작', type='primary', key='verify_btn'):
                with st.spinner(f'매니페스트 {len(verify_files)}개 파싱 중...'):
                    _all_boxes = []
                    _parse_errors = []
                    for f in verify_files:
                        try:
                            _b = f.read(); f.seek(0)
                            _all_boxes.extend(_extract_manifest_products(_b))
                        except Exception as e:
                            _parse_errors.append((f.name, str(e)))

                if _parse_errors:
                    for _n, _e in _parse_errors:
                        st.error(f'❌ `{_n}` 파싱 실패: {_e}')

                # 송장별 SKU 집계 (매니페스트 = 정답)
                _manifest_skus = {}   # {invoice: {sku_id}}
                _manifest_box_count = {}
                for _box in _all_boxes:
                    _inv = _box.get('invoice_number')
                    if not _inv:
                        continue
                    _manifest_skus.setdefault(_inv, set()).update(
                        p['sku_id'] for p in _box['products']
                    )
                    _manifest_box_count[_inv] = _manifest_box_count.get(_inv, 0) + 1

                # 송장별 SKU 집계 (시트 = 검증 대상)
                def _norm_sku(v):
                    s = str(v or '').strip()
                    if not s:
                        return ''
                    try:
                        return str(int(float(s)))  # "71572440.0" → "71572440"
                    except (ValueError, TypeError):
                        return s

                _sheet_skus = {}
                for _, _row in _df_pick.iterrows():
                    _inv = str(_row.get('쉽먼트운송장번호', '') or '').strip()
                    _sku = _norm_sku(_row.get(_sku_col))
                    if not _inv or not _sku:
                        continue
                    _sheet_skus.setdefault(_inv, set()).add(_sku)

                # 비교 — 시트 기준으로 검증 (시트의 SKU가 매니페스트에 있는가?)
                _issues = []            # 오분류 있는 송장
                _sheet_only_inv = set(_sheet_skus.keys()) - set(_manifest_skus.keys())
                _clean_count = 0
                for _inv, _ssk in _sheet_skus.items():
                    _msk = _manifest_skus.get(_inv, set())
                    if not _msk:  # 매니페스트 자체가 없는 송장 → 별도 처리
                        continue
                    _wrong = _ssk - _msk   # 🚨 시트에만 있음 = 오분류
                    _missing = _msk - _ssk  # ℹ️ 매니페스트에만 있음 = 미도착 또는 참고
                    if not _wrong:
                        _clean_count += 1
                    else:
                        _issues.append({
                            'invoice': _inv,
                            'n_sheet': len(_ssk),
                            'n_manifest': len(_msk),
                            'wrong_in_sheet': sorted(_wrong),
                            'missing_vs_manifest': sorted(_missing),
                            'n_boxes': _manifest_box_count.get(_inv, 0),
                        })

                # 전체 결과
                _total_sheet_inv = len(_sheet_skus)
                _matched_inv = len(set(_sheet_skus.keys()) & set(_manifest_skus.keys()))

                st.markdown('### 검증 결과')
                _c1, _c2, _c3, _c4 = st.columns(4)
                with _c1:
                    st.metric('시트 송장', f'{_total_sheet_inv}건')
                with _c2:
                    st.metric('송장 매칭', f'{_matched_inv}건',
                              help='시트 송장 중 매니페스트에도 존재하는 송장 수')
                with _c3:
                    st.metric('🚨 오분류 송장', f'{len(_issues)}건',
                              help='시트에 매니페스트에 없는 상품이 배정된 송장')
                with _c4:
                    st.metric('매니페스트 없음', f'{len(_sheet_only_inv)}건',
                              help='시트엔 있지만 업로드한 매니페스트에 송장 자체가 없음')

                if len(_issues) == 0 and _sheet_only_inv == set():
                    st.success(f'✅ 완전 일치 — 시트 송장 {_total_sheet_inv}건 모두 매니페스트와 SKU 일치')
                else:
                    if _sheet_only_inv:
                        with st.expander(f'ℹ️ 매니페스트에 없는 송장 ({len(_sheet_only_inv)}건) — 매니페스트 PDF 추가 업로드 필요할 수도'):
                            st.code('\n'.join(sorted(_sheet_only_inv)))

                    if _issues:
                        st.error(f'🚨 오분류 발견 — {len(_issues)}개 송장에 매니페스트에 없는 상품이 배정되어 있음. 시트에서 해당 행의 송장번호를 올바르게 수정하세요.')
                        _rows = []
                        for _it in _issues:
                            _rows.append({
                                '송장번호': _it['invoice'],
                                '🚨 오분류 SKU': len(_it['wrong_in_sheet']),
                                '시트 SKU 총': _it['n_sheet'],
                                '매니페스트 SKU 총': _it['n_manifest'],
                                '미도착 가능': len(_it['missing_vs_manifest']),
                                '박스 수': _it['n_boxes'],
                            })
                        import pandas as _pdv
                        st.dataframe(_pdv.DataFrame(_rows),
                                     use_container_width=True, hide_index=True)

                        for _it in _issues:
                            with st.expander(f"🚨 송장 {_it['invoice']} — 오분류 {len(_it['wrong_in_sheet'])}개"):
                                st.error(
                                    f"**시트에 잘못 배정된 SKU ({len(_it['wrong_in_sheet'])}개)** — "
                                    f"이 송장의 매니페스트에 없는 상품임:\n\n"
                                    + ', '.join(f'`{s}`' for s in _it['wrong_in_sheet'])
                                )
                                if _it['missing_vs_manifest']:
                                    st.info(
                                        f"참고: 매니페스트엔 있는데 시트에 없는 SKU ({len(_it['missing_vs_manifest'])}개) — "
                                        f"아직 도착 안 했을 수 있음:\n\n"
                                        + ', '.join(f'`{s}`' for s in _it['missing_vs_manifest'])
                                    )

    # ── 🚚 한진택배 엑셀 만들기 (출고확인 + 주소리스트 매칭) ──
    st.divider()
    with st.expander(f'🚚 {kit_config.courier_name()} 엑셀 만들기 (물류센터 기준 자동 매칭)', expanded=False):
        st.caption('출고확인 데이터를 (물류센터, 송장번호)로 그룹화 → 주소리스트 시트의 물류센터 키로 매칭 → 택배 송장 양식 엑셀 생성 (열 구성은 ⚙️ 설정 탭의 택배 양식)')
        _hanjin_addr_url = HANJIN_ADDR_SHEET_URL
        st.caption(f'📋 주소리스트: 설정 탭에서 지정한 시트의 "{kit_config.courier_addr_tab()}" 탭 사용 (E열 키 = 물류센터명)')

        if st.button(f'🚚 {kit_config.courier_name()} 엑셀 생성', key='pick_hanjin_btn', type='primary', use_container_width=True):
            try:
                _sclient = st.session_state.get('pick_gsheet_client') or get_gsheet_client()
                _pick_df = st.session_state.get('pick_df_출고')
                _buf, _n_rows, _unmatched, _invalid = _hanjin_build_excel(
                    _pick_df, _sclient, _hanjin_addr_url, kit_config.courier_addr_tab(),
                )
                st.session_state['pick_hanjin_result'] = {
                    'buf': _buf.getvalue(),
                    'n_rows': _n_rows,
                    'unmatched': _unmatched,
                    'invalid': _invalid,
                    'ts': datetime.now().strftime('%Y%m%d_%H%M'),
                }
            except Exception as _e:
                st.session_state.pop('pick_hanjin_result', None)
                st.error(f'❌ {_e}')

        # 결과 표시 (rerun 후에도 유지)
        _hr = st.session_state.get('pick_hanjin_result')
        if _hr:
            st.success(f'✅ {_hr["n_rows"]}건 생성 완료')
            if _hr.get('invalid'):
                st.error(
                    f'🚨 송장번호 판독 불가 {len(_hr["invalid"])}건 — 엑셀에서 제외됨. '
                    '원본 파일에서 송장번호 열이 과학표기(4.62E+11)로 잘려 저장된 경우입니다. '
                    '해당 열을 텍스트 서식으로 바꿔 다시 받아주세요.'
                )
                with st.expander(f'제외된 송장 {len(_hr["invalid"])}건 보기'):
                    for _iv in _hr['invalid']:
                        st.caption(f'• {_iv}')
            if _hr['unmatched']:
                with st.expander(f'⚠️ 매칭 실패 {len(_hr["unmatched"])}건 (주소리스트에 없는 물류센터)'):
                    for _u in _hr['unmatched']:
                        st.caption(f'• {_u}')
            st.download_button(
                label=f'⬇️ {kit_config.courier_name()} 엑셀 다운로드',
                data=_hr['buf'],
                file_name=f'{kit_config.courier_name()}출고양식_{_hr["ts"]}.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                key='pick_hanjin_dl',
                use_container_width=True,
            )

    # ── 모드 토글 (피킹 검증 ↔ 입고 분류) ──
    st.divider()
    work_mode = st.radio(
        "🎯 작업 모드 선택",
        options=["📥 입고 분류", "📤 피킹 검증"],
        index=0,
        horizontal=True,
        key="pick_work_mode",
        help="입고 분류 = 배대지 박스 열고 박스별로 분류 / 피킹 검증 = 송장별 출고 박스 채우기",
    )
    st.divider()

    if work_mode == "📤 피킹 검증":
        if not st.session_state.pick_selected_shipment:
            # ── 송장번호 선택 ──
            st.markdown('<div class="shipment-input">', unsafe_allow_html=True)
            st.markdown("### 📋 쉽먼트 선택")
            st.caption("송장번호 입력(또는 바코드 스캔) 후 Enter → 자동으로 피킹 시작. 여러 개면 쉼표로 구분 후 🚀 피킹 시작 클릭.")

            pick_df = st.session_state.pick_df_출고

            # 스캔 후 입력창 초기화용 카운터
            if 'pick_ship_input_counter' not in st.session_state:
                st.session_state.pick_ship_input_counter = 0

            p_col1, p_col2 = st.columns([2, 1])
            with p_col1:
                _ship_input_key = f"pick_shipment_input_{st.session_state.pick_ship_input_counter}"
                input_shipment = st.text_input(
                    "송장번호 입력 (Enter로 자동 시작)",
                    placeholder="예: 461938764685 (스캐너로 스캔 후 Enter 또는 자동 입력)",
                    key=_ship_input_key,
                )
            with p_col2:
                centers = ["전체"] + sorted(pick_df["물류센터(FC)"].unique().tolist()) if "물류센터(FC)" in pick_df.columns else ["전체"]
                center = st.selectbox("물류센터", centers, key="pick_center_filter")

            # ── 자동 시작 로직 (단일 유효 송장 입력 시) ──
            if input_shipment and input_shipment.strip():
                # 쉼표/줄바꿈/공백으로 토큰 분리
                _toks = [t.strip() for t in re.split(r'[,\s\n]+', input_shipment.strip()) if t.strip()]
                if len(_toks) == 1:
                    _stgt = _toks[0]
                    _valid_ids = list(pick_df["쉽먼트운송장번호"].unique())
                    _resolved = None
                    if _stgt in _valid_ids:
                        _resolved = _stgt
                    else:
                        _mm = [s for s in _valid_ids if s.endswith(_stgt)]
                        if len(_mm) == 1:
                            _resolved = _mm[0]
                    if _resolved:
                        pick_init_picking([_resolved])
                        st.session_state.pick_start_audio_pending = True
                        st.session_state.pick_ship_input_counter += 1
                        st.rerun()

            pick_df = st.session_state.pick_df_출고
            if center != "전체" and "물류센터(FC)" in pick_df.columns:
                filtered = pick_df[pick_df["물류센터(FC)"] == center]
            else:
                filtered = pick_df

            summary = filtered.groupby("쉽먼트운송장번호").agg(
                SKU수=("바코드", "nunique"), 총수량=("수량", "sum"),
            ).reset_index().sort_values("총수량", ascending=False)

            selected_shipment = st.selectbox(
                "또는 목록에서 선택",
                options=summary["쉽먼트운송장번호"].tolist(),
                format_func=lambda x: (
                    f"{'✅ ' if x in st.session_state.pick_completed_shipments else ''}"
                    f"{x[-6:]} | "
                    f"{summary[summary['쉽먼트운송장번호']==x]['SKU수'].values[0]}종 "
                    f"{summary[summary['쉽먼트운송장번호']==x]['총수량'].values[0]}개"
                ),
                key="pick_shipment_select",
            )

            # 다중 송장 파싱: 쉼표/공백/줄바꿈으로 구분
            input_targets = []
            if input_shipment and input_shipment.strip():
                for token in re.split(r'[,\s\n]+', input_shipment.strip()):
                    token = token.strip()
                    if token:
                        input_targets.append(token)

            if st.button("🚀 피킹 시작", type="primary", use_container_width=True, key="pick_start_btn"):
                valid_ids = list(pick_df["쉽먼트운송장번호"].unique())
                resolved = []
                errors = []
                if input_targets:
                    for tgt in input_targets:
                        if tgt in valid_ids:
                            resolved.append(tgt)
                        else:
                            matches = [s for s in valid_ids if s.endswith(tgt)]
                            if len(matches) == 1:
                                resolved.append(matches[0])
                            elif len(matches) > 1:
                                errors.append(f"'{tgt}'에 매칭되는 쉽먼트가 {len(matches)}개입니다.")
                            else:
                                errors.append(f"'{tgt}'에 해당하는 쉽먼트를 찾을 수 없습니다.")
                elif selected_shipment:
                    resolved.append(selected_shipment)

                if errors:
                    for err in errors:
                        st.error(err)
                elif resolved:
                    pick_init_picking(resolved)
                    st.rerun()
                else:
                    st.warning("송장번호를 입력하거나 목록에서 선택해주세요.")

            # 송장 입력창 자동 포커스 — "다른 쉽먼트"로 돌아온 직후 바로 스캔 가능
            from streamlit.components.v1 import html as _ship_focus_html
            _ship_focus_html("""<script>
            (function(){
                const doc = window.parent.document;
                function findShipInput(){
                    const inputs = doc.querySelectorAll('input[type="text"]');
                    for (const inp of inputs){
                        if (inp.placeholder && inp.placeholder.includes('스캐너로 스캔')) return inp;
                    }
                    return null;
                }
                function isInteractingOther(){
                    const active = doc.activeElement;
                    if (!active) return false;
                    const tag = (active.tagName || '').toLowerCase();
                    if (tag === 'button' || tag === 'textarea') return true;
                    if (tag === 'input' && active.type !== 'text') return true;
                    // 다른 텍스트 입력창을 쓰는 중이면 포커스를 뺏지 않는다
                    if (tag === 'input' && active.type === 'text' &&
                        !(active.placeholder || '').includes('스캐너로 스캔')) return true;
                    // 셀렉트박스(목록에서 선택/물류센터) 조작 중이면 스킵
                    if (active.closest) {
                        if (active.closest('[data-baseweb="select"]')) return true;
                        if (active.closest('[data-baseweb="popover"]')) return true;
                        if (active.closest('[role="listbox"]')) return true;
                        if (active.closest('[role="combobox"]')) return true;
                    }
                    return false;
                }
                function focusShip(){
                    const inp = findShipInput();
                    if (!inp) return;
                    if (doc.activeElement === inp) return;
                    if (isInteractingOther()) return;
                    // 드롭다운이 열려있으면 포커스 안 함
                    if (doc.querySelector('[data-baseweb="popover"]')) return;
                    inp.focus({preventScroll: true});
                }
                focusShip();
                if (window._shipFocusInterval) clearInterval(window._shipFocusInterval);
                window._shipFocusInterval = setInterval(focusShip, 500);
            })();
            </script>""", height=0)
            st.markdown('</div>', unsafe_allow_html=True)
        else:
            # ── 피킹 진행 화면 ──
            shipment_id = st.session_state.pick_selected_shipment

            # 신규 진입 시 "확인을 시작하세요" 음성 안내 (1회)
            if st.session_state.get('pick_start_audio_pending'):
                from streamlit.components.v1 import html as _st_start_html
                _st_start_html(
                    f"<script>{_tts_ko_script('확인을 시작하세요')}</script>",
                    height=0,
                )
                st.session_state.pick_start_audio_pending = False

            # 이미지 URL 백그라운드 프리로드 (1회) — 스캔 시 즉시 표시되도록 브라우저 캐시에 미리 적재
            _preload_sig = ','.join(sorted([
                str(v.get('이미지URL', '') or '')
                for v in st.session_state.pick_picking_state.values()
                if v.get('이미지URL')
            ]))
            if _preload_sig and st.session_state.get('_pick_img_preload_sig') != _preload_sig:
                import json as _json_pl
                from streamlit.components.v1 import html as _st_pl_html
                _urls = list({
                    str(v.get('이미지URL', '') or '').strip()
                    for v in st.session_state.pick_picking_state.values()
                    if str(v.get('이미지URL', '') or '').strip()
                })
                _st_pl_html(
                    '<meta name="referrer" content="no-referrer">'
                    f"<script>(function(){{var u={_json_pl.dumps(_urls)};"
                    "u.forEach(function(s){try{var i=new Image();i.referrerPolicy='no-referrer';i.src=s;}catch(e){}});}})();</script>",
                    height=0,
                )
                st.session_state['_pick_img_preload_sig'] = _preload_sig

            hcol1, hcol2, hcol3 = st.columns([3, 1, 1])
            with hcol1:
                item0 = list(st.session_state.pick_picking_state.values())[0] if st.session_state.pick_picking_state else {}
                ships = st.session_state.get('pick_selected_shipments', [shipment_id])
                if len(ships) > 1:
                    ship_lines = " | ".join([f"**{i+1}번박스:** `{s[-6:]}`" for i, s in enumerate(ships)])
                    st.markdown(f"{ship_lines} | **센터:** {item0.get('물류센터','')}")
                else:
                    st.markdown(f"**쉽먼트:** `{shipment_id}` | **센터:** {item0.get('물류센터','')} | **회차:** {item0.get('회차기호','')}")
            with hcol2:
                if st.button("➕ 쉽먼트 추가", use_container_width=True, key="pick_add_btn"):
                    st.session_state.pick_show_add_input = True
                    st.rerun()
            with hcol3:
                if st.button("🔄 다른 쉽먼트", use_container_width=True, key="pick_change_btn"):
                    # 쉽먼트 관련 상태 완전 초기화 (로그/재고/완료목록은 유지)
                    st.session_state.pick_selected_shipment = None
                    st.session_state.pick_selected_shipments = []
                    st.session_state.pick_picking_state = {}
                    st.session_state.pick_shortage_items = []
                    st.session_state.pick_last_scan_result = None
                    st.session_state.pick_scan_counter = 0
                    st.session_state.pick_show_add_input = False
                    # 다량 모드 상태 초기화
                    st.session_state.pick_next_qty = 1
                    st.session_state.pick_qty_input_mode = False
                    st.rerun()

            # 쉽먼트 추가 입력 영역
            if st.session_state.get('pick_show_add_input'):
                with st.container():
                    ac1, ac2, ac3 = st.columns([3, 1, 1])
                    with ac1:
                        add_input = st.text_input("추가할 송장번호", key="pick_add_shipment_input",
                                                  placeholder="송장번호 입력 후 추가 클릭")
                    with ac2:
                        if st.button("✅ 추가", use_container_width=True, key="pick_add_confirm"):
                            target = (add_input or '').strip()
                            if target:
                                valid_ids = list(st.session_state.pick_df_출고["쉽먼트운송장번호"].unique())
                                resolved = None
                                if target in valid_ids:
                                    resolved = target
                                else:
                                    matches = [s for s in valid_ids if s.endswith(target)]
                                    if len(matches) == 1:
                                        resolved = matches[0]
                                if resolved and resolved not in ships:
                                    new_ships = list(ships) + [resolved]
                                    pick_init_picking(new_ships)
                                    st.session_state.pick_show_add_input = False
                                    if 'pick_add_shipment_input' in st.session_state:
                                        del st.session_state['pick_add_shipment_input']
                                    st.rerun()
                                elif resolved in ships:
                                    st.warning("이미 추가된 송장입니다")
                                else:
                                    st.error(f"'{target}' 송장을 찾을 수 없습니다")
                    with ac3:
                        if st.button("❌ 취소", use_container_width=True, key="pick_add_cancel"):
                            st.session_state.pick_show_add_input = False
                            st.rerun()

            st.markdown("---")

            # ── 바코드 스캔 (fragment으로 감싸서 전체 앱 리런 없이 조각만 재실행) ──
            _pick_use_fragment = getattr(st, 'fragment', lambda f: f)

            @_pick_use_fragment
            def _pick_scan_fragment():
                # ── 진행률 (매 스캔마다 갱신되도록 fragment 안에서 계산) ──
                _prog = pick_get_progress()
                _sid = st.session_state.pick_selected_shipment or ''
                pc1, pc2, pc3, pc4, pc5 = st.columns(5)
                pc1.metric("스캔", f"{_prog['scanned']}/{_prog['total']}")
                pc2.metric("SKU 완료", f"{_prog['done_skus']}/{_prog['skus']}")
                pc3.metric("진행률", f"{_prog['pct']:.0%}")
                pc4.metric("초과 스캔", f"{_prog['over']}건",
                           delta=f"+{_prog['over']}" if _prog['over'] > 0 else None, delta_color="inverse")
                pc5.metric("재고 부족", f"{_prog['shortage']}건",
                           delta=f"{_prog['shortage']}" if _prog['shortage'] > 0 else None, delta_color="inverse")
                st.progress(_prog["pct"])

                if _prog["is_complete"]:
                    st.markdown(
                        f'<div class="scan-complete">'
                        f'<strong style="font-size:1.4rem;">🎉 검증확인이 완료되었습니다. 출고하세요!</strong><br>'
                        f'<span style="font-size:1.05rem;">쉽먼트 {_sid[-6:]} — {_prog["total"]}개 전부 검증 완료</span>'
                        f'</div>',
                        unsafe_allow_html=True)
                    _newly_done = _sid not in st.session_state.pick_completed_shipments
                    st.session_state.pick_completed_shipments.add(_sid)
                    if _newly_done:
                        # 출고 직전이므로 큐에 남은 쓰기를 지금 시트에 반영.
                        # 백그라운드로 보낸다 — 여기서 직접 부르면 마지막 스캔 화면이
                        # 네트워크(429 백오프면 수십 초)를 기다리며 멈춘다.
                        _gs_threading.Thread(target=gs_flush_pending, daemon=True,
                                             name='gsheet-flush-done').start()
                        from streamlit.components.v1 import html as _st_html_done
                        _st_html_done(
                            f"<script>{_tts_ko_script('검증확인이 완료되었습니다. 출고하세요', extra_rate=-0.05)}</script>",
                            height=0,
                        )

                st.markdown("---")

                # 다량 모드 상태 초기화
                if 'pick_next_qty' not in st.session_state:
                    st.session_state.pick_next_qty = 1
                if 'pick_qty_input_mode' not in st.session_state:
                    st.session_state.pick_qty_input_mode = False

                # ── 입력 처리는 전부 on_change 콜백에서 ──
                # 예전에는 스캔마다 key 를 바꾼 새 입력창을 만들고 st.rerun(scope='fragment')
                # 을 한 번 더 불러서, 스캔 1건 = 조각 실행 2번 + 입력창 DOM 교체였다.
                # 그 사이에 스캐너가 다음 바코드를 쏘면 글자가 사라지거나 반응이 없었다.
                # 콜백은 조각이 다시 그려지기 전에 한 번만 돌고, 위젯 자기 값을 비우는
                # 것은 콜백 안에서만 안정적으로 동작한다 (재고 탭과 같은 방식).
                def _on_pick_qty_enter():
                    raw = str(st.session_state.get('pick_qty_input', '') or '').strip()
                    st.session_state['pick_qty_input'] = ''
                    if not raw:
                        return
                    try:
                        _qv = int(raw)
                    except ValueError:
                        st.session_state['pick_qty_error'] = '숫자만 입력 가능합니다'
                        return
                    if _qv < 1:
                        st.session_state['pick_qty_error'] = '1 이상을 입력하세요'
                        return
                    st.session_state.pop('pick_qty_error', None)
                    st.session_state.pick_next_qty = _qv
                    st.session_state.pick_qty_input_mode = False
                    # 수량 확정 안내 — 이전 결과(다량 입력 모드)가 다시
                    # 재생되지 않도록 결과를 교체한다
                    st.session_state.pick_last_scan_result = {
                        "status": "qty_set",
                        "message": f"📦 다음 스캔: {_qv}개",
                        "detail": "바코드를 스캔해 주세요",
                        "barcode": "", "상품명": "",
                        "시간": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    }
                    st.session_state.pick_scan_counter += 1

                def _on_pick_scan_enter():
                    raw = str(st.session_state.get('pick_scan_input', '') or '').strip()
                    st.session_state['pick_scan_input'] = ''
                    if not raw:
                        return
                    _pick_qty = int(st.session_state.get('pick_next_qty', 1) or 1)
                    pick_process_scan(raw, qty=_pick_qty)

                def _on_pick_qty_reset():
                    st.session_state.pick_next_qty = 1

                # 수량 입력 모드: 숫자 입력 후 Enter
                if st.session_state.pick_qty_input_mode:
                    st.warning('🔢 **수량을 입력하세요** — 숫자 입력 후 Enter')
                    st.text_input(
                        '다량 수량',
                        key='pick_qty_input',
                        placeholder='숫자 입력 후 Enter (예: 50)',
                        label_visibility='collapsed',
                        on_change=_on_pick_qty_enter,
                    )
                    if st.session_state.get('pick_qty_error'):
                        st.error(st.session_state['pick_qty_error'])

                # 수량 표시 + 1개 모드 리셋 버튼
                pqcol1, pqcol2 = st.columns([1, 1])
                with pqcol1:
                    if st.session_state.pick_next_qty > 1:
                        st.markdown(
                            f'<div style="background:#f59e0b;color:white;padding:0.5rem;border-radius:6px;text-align:center;font-weight:bold;font-size:1.1rem">'
                            f'📦 다음 스캔: {st.session_state.pick_next_qty}개'
                            f'</div>', unsafe_allow_html=True)
                    else:
                        st.markdown(
                            '<div style="background:#e5e7eb;padding:0.5rem;border-radius:6px;text-align:center">'
                            '1개 모드'
                            '</div>', unsafe_allow_html=True)
                with pqcol2:
                    if st.session_state.pick_next_qty > 1 and not st.session_state.pick_qty_input_mode:
                        # on_click 콜백: 버튼 클릭 → 조각 1회 실행으로 끝난다 (추가 rerun 없음)
                        st.button('🔄 1개 모드로 복귀', key='pick_qty_reset', use_container_width=True,
                                  on_click=_on_pick_qty_reset)

                # key 를 고정하면 스캔 후에도 같은 입력창 DOM 이 유지되어 포커스가 끊기지 않는다
                st.text_input("🔫 바코드 스캔 (스캐너 또는 직접 입력)", key='pick_scan_input',
                              placeholder="스캐너 대기 중... 바코드 (여러 개면 #MULTI 먼저)",
                              on_change=_on_pick_scan_enter)

                # 바코드 입력창에 자동 포커스 유지
                from streamlit.components.v1 import html as _st_html
                _st_html("""<script>
                (function(){
                    const doc = window.parent.document;
                    function findScanInput() {
                        const inputs = doc.querySelectorAll('input[type="text"]');
                        // 수량 입력창은 다량 모드일 때만 존재 → 있으면 그쪽에 포커스
                        for (const inp of inputs) {
                            if (inp.placeholder && inp.placeholder.includes('숫자 입력')) {
                                return inp;
                            }
                        }
                        // 기본: 바코드 스캐너 입력창
                        for (const inp of inputs) {
                            if (inp.placeholder && inp.placeholder.includes('스캐너')) {
                                return inp;
                            }
                        }
                        return null;
                    }
                    function isInteractingOther(e) {
                        // 드래그 중(스크롤바/리사이즈 등)이거나 dataframe 내부면 포커스 복귀 스킵
                        const t = e && e.target;
                        if (!t) return false;
                        if (t.closest && (
                            t.closest('[data-testid="stDataFrame"]') ||
                            t.closest('canvas') ||
                            t.closest('.glideDataEditor') ||
                            t.closest('[role="grid"]') ||
                            t.closest('[data-testid="stExpander"]')
                        )) return true;
                        return false;
                    }
                    function focusScan(e) {
                        if (e && isInteractingOther(e)) return;
                        const inp = findScanInput();
                        if (inp && doc.activeElement !== inp) {
                            // preventScroll: 포커스 이동 시 브라우저 자동 스크롤 방지
                            inp.focus({preventScroll: true});
                        }
                    }
                    // 즉시 포커스
                    focusScan();
                    // 짧은 간격으로 반복 (0.5초) — 스크롤 조작 방해 최소화
                    if (window._scanFocusInterval) clearInterval(window._scanFocusInterval);
                    window._scanFocusInterval = setInterval(function(){
                        // 사용자가 다른 곳과 상호작용 중(selection, 드래그)이면 스킵
                        const sel = doc.getSelection && doc.getSelection();
                        if (sel && sel.toString().length > 0) return;
                        focusScan();
                    }, 500);
                    // DOM 변경 감지는 제거 (dataframe 스크롤이 DOM 변경 유발 → 루프)
                    if (window._scanObserver) { window._scanObserver.disconnect(); window._scanObserver = null; }
                    // 다른 곳 클릭해도 입력창으로 복귀 (버튼/링크/테이블/expander 제외)
                    // parent 문서에 한 번만 등록 — 조각이 다시 그려질 때마다 쌓이지 않도록
                    if (!window.parent.__pickClickFocusInstalled) {
                        window.parent.__pickClickFocusInstalled = true;
                        doc.addEventListener('click', function(e){
                            const tag = (e.target.tagName||'').toLowerCase();
                            if (tag === 'button' || tag === 'a' || tag === 'input' || tag === 'select' || tag === 'textarea') return;
                            if (isInteractingOther(e)) return;
                            setTimeout(function(){ focusScan(e); }, 50);
                        }, true);
                    }
                })();
                </script>""", height=0)

                r = st.session_state.pick_last_scan_result
                if r:
                    _r_bc = str(r.get("barcode", "") or "")
                    _r_img = ''
                    if _r_bc and _r_bc in st.session_state.pick_picking_state:
                        _r_img = str(st.session_state.pick_picking_state[_r_bc].get("이미지URL", "") or "").strip()
                    # iframe 카드: no-referrer + onerror 동작 (st.markdown은 DOMPurify가 제거함)
                    render_scan_result_card(
                        r["status"], r["message"], r["detail"],
                        barcode=_r_bc, img_url=_r_img, img_size=120, height=210,
                    )
                    # 스캔 결과 소리 (비프음)
                    sound_js = {
                        "ok": "o.frequency.value=880;g.gain.value=0.3;o.start();setTimeout(()=>g.gain.value=0,150);setTimeout(()=>o.stop(),200);",
                        "error": "o.type='square';o.frequency.value=200;g.gain.value=0.5;o.start();setTimeout(()=>{o.frequency.value=150},150);setTimeout(()=>g.gain.value=0,500);setTimeout(()=>o.stop(),600);",
                        "over": "o.type='sawtooth';o.frequency.value=400;g.gain.value=0.4;o.start();setTimeout(()=>{o.frequency.value=300},100);setTimeout(()=>g.gain.value=0,300);setTimeout(()=>o.stop(),400);",
                        "shortage": "o.frequency.value=600;g.gain.value=0.3;o.start();setTimeout(()=>{o.frequency.value=400},100);setTimeout(()=>g.gain.value=0,250);setTimeout(()=>o.stop(),300);",
                    }
                    js_code = sound_js.get(r["status"], sound_js["ok"])

                    # 음성 안내 (한국어 TTS)
                    box_label = r.get("박스", "")
                    # 1번박스 → 일번박스 형태로 변환
                    _KOR_NUMS = {'1':'일','2':'이','3':'삼','4':'사','5':'오','6':'육','7':'칠','8':'팔','9':'구','10':'십'}
                    def _kor_box(label):
                        import re as _re
                        m = _re.match(r'(\d+)번박스', label or '')
                        if not m:
                            return label
                        n = m.group(1)
                        kor = _KOR_NUMS.get(n, n)
                        return f"{kor}번박스"
                    box_kor = _kor_box(box_label)

                    ships_count = len(st.session_state.get('pick_selected_shipments', []))
                    if r["status"] == "multi_trigger":
                        speak_text = "수량을 입력하세요"
                    elif r["status"] == "qty_set":
                        speak_text = "바코드를 스캔해 주세요"
                    elif r["status"] == "error":
                        speak_text = "없는 상품 입니다"
                    elif r["status"] == "over":
                        speak_text = "수량 초과"
                    elif r["status"] == "shortage":
                        if ships_count <= 1:
                            speak_text = "입고완료 재고 부족"
                        else:
                            speak_text = f"{box_kor} 재고 부족" if box_kor else "재고 부족"
                    elif ships_count <= 1:
                        speak_text = "입고완료"
                    elif box_kor:
                        speak_text = f"{box_kor}"
                    else:
                        speak_text = "확인"

                    # 매 스캔마다 새 컴포넌트로 강제 재실행 (같은 박스도 소리 나도록)
                    scan_id = st.session_state.pick_scan_counter
                    _tts_js_block = _tts_ko_script(speak_text)

                    from streamlit.components.v1 import html as st_html
                    st_html(f"""<script>
                    // scan_id={scan_id} (강제 재실행용)
                    try{{var a=new(window.AudioContext||window.webkitAudioContext)();var o=a.createOscillator();var g=a.createGain();o.connect(g);g.connect(a.destination);{js_code}}}catch(e){{}}
                    {_tts_js_block}
                    </script>""", height=0)

                # ── 피킹 현황 (매 스캔마다 갱신되도록 fragment 안에서 렌더링) ──
                st.markdown("---")
                st.subheader("📋 피킹 현황")
                rows = []
                for bc, info in st.session_state.pick_picking_state.items():
                    s, n = info["스캔수량"], info["필요수량"]
                    if s > n: status_txt = f"⚠️ 초과 ({s}/{n})"
                    elif s >= n: status_txt = "✅ 완료"
                    elif s > 0: status_txt = f"🔄 {s}/{n}"
                    else: status_txt = "⬜ 대기"
                    inv = info.get("배대지잔여")
                    ship_boxes = info.get("쉽먼트박스목록", [])
                    ship_box_str = ",".join(ship_boxes) if ship_boxes else ""
                    rows.append({
                        "상태": status_txt, "바코드": bc,
                        "상품명": info["상품명"],
                        "쉽먼트박스": ship_box_str,
                        "필요": n, "스캔": s, "남은": max(0, n - s),
                        "회차": info.get("회차기호",""), "박스": info.get("박스번호",""),
                        "배대지재고": f"{inv}" if inv is not None else "-",
                    })
                pick_order = {"🔄":0,"⬜":1,"✅":2,"⚠️":3}
                rows.sort(key=lambda x: pick_order.get(x["상태"][0], 9))
                st.dataframe(_pd.DataFrame(rows), use_container_width=True, hide_index=True,
                             height=min(500, len(rows) * 38 + 40))

                shortage = st.session_state.get("pick_shortage_items", [])
                if shortage:
                    with st.expander(f"⛔ 부족분 — 피킹 불가 ({len(shortage)}건)", expanded=False):
                        st.caption("출고지시서에 '부족'으로 표시된 항목입니다.")
                        st.dataframe(_pd.DataFrame(shortage), use_container_width=True, hide_index=True)

                if st.session_state.pick_scan_log:
                    with st.expander(f"📜 스캔 로그 ({len(st.session_state.pick_scan_log)}건)"):
                        log_display = []
                        for entry in reversed(st.session_state.pick_scan_log[-50:]):
                            icon = {"ok":"✅","over":"⚠️","error":"🚨","shortage":"📦"}.get(entry["status"],"?")
                            log_display.append({"시간":entry["시간"],"결과":icon,"바코드":entry["barcode"],"내용":entry["message"]})
                        st.dataframe(_pd.DataFrame(log_display), use_container_width=True, hide_index=True)

                    st.download_button(
                        "📥 스캔 로그 CSV",
                        data=_pd.DataFrame(st.session_state.pick_scan_log).to_csv(index=False, encoding="utf-8-sig"),
                        file_name=f"picking_log_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
                        mime="text/csv", use_container_width=True, key="pick_log_dl")

            _pick_scan_fragment()

    elif work_mode == "📥 입고 분류":
        import pandas as _pd2
        df_sort = st.session_state.pick_df_출고

        # ── 송장별 박스번호: 시트 M열(기존값) 우선 보존 + 신규만 부여 ──
        # 박스번호가 시트 M열에 영구 저장되므로 발주 취소돼도 재정렬 안 됨.
        sort_items_for_box = []
        for _, row in df_sort.iterrows():
            sort_items_for_box.append({
                'shipmentNumber': str(row.get('쉽먼트운송장번호', '') or '').strip(),
                'logisticsCenter': str(row.get('물류센터(FC)', '') or '').strip(),
                'expectedDate': str(row.get('입고예정일(EDD)', '') or row.get('입고예정일', '') or '').strip(),
                'boxNumber': str(row.get('박스번호', '') or '').strip(),
            })

        # 1) 시트 M열에서 기존 박스번호 읽기 (gsheet 모드 + 세션 캐시)
        #    실패(None) 시 캐시 저장 안 함 → 다음 rerun에서 재시도
        _pick_url = st.session_state.get('pick_sheet_url_출고', '')
        _pick_tab = st.session_state.get('pick_sheet_tab_출고', '')
        _existing_box_map = {}
        _read_failed = False
        _gsheet_ready = (st.session_state.get('pick_use_gsheet')
                         and st.session_state.get('pick_gsheet_client')
                         and _pick_url and _pick_tab)
        if _gsheet_ready:
            _cache_key = f"_pick_existing_box_{_pick_url}_{_pick_tab}"
            if _cache_key in st.session_state:
                _existing_box_map = st.session_state[_cache_key]
            else:
                _read = pick_read_box_numbers(
                    st.session_state.pick_gsheet_client, _pick_url, _pick_tab
                )
                if _read is None:
                    _read_failed = True
                    st.warning('⚠️ 시트 M열 읽기 실패 — 이번 세션에서 박스번호 쓰기는 건너뜁니다. 페이지를 새로고침하세요.')
                else:
                    _existing_box_map = _read
                    st.session_state[_cache_key] = _read

        # 2) 기존 유지 + 신규 송장만 순차 부여
        sort_ship_to_box = assign_box_numbers_with_existing(
            sort_items_for_box, _existing_box_map
        )

        # 3) 신규 부여된 것만 시트 M열에 동기 기록 (기존값 보존, 실패 시 캐시 갱신 안 함)
        if _gsheet_ready and not _read_failed:
            _new_only = {s: n for s, n in sort_ship_to_box.items()
                         if s not in _existing_box_map}
            _written_key = f"_pick_box_written_{_pick_url}_{_pick_tab}"
            _last_written = st.session_state.get(_written_key, {})
            _to_write = {s: n for s, n in _new_only.items()
                         if _last_written.get(s) != n}
            if _to_write:
                with st.spinner(f'📝 시트 M열에 박스번호 {len(_to_write)}건 기록 중...'):
                    _result = pick_write_box_numbers(
                        st.session_state.pick_gsheet_client,
                        _pick_url, _pick_tab, _to_write, only_empty=True,
                    )
                if _result >= 0:
                    # 성공 시에만 캐시 갱신
                    _merged = dict(_existing_box_map)
                    _merged.update(_new_only)
                    st.session_state[_cache_key] = _merged
                    _last_written.update(_to_write)
                    st.session_state[_written_key] = _last_written
                else:
                    st.warning('⚠️ 시트 M열 쓰기 실패 — 박스번호가 시트에 저장되지 않았습니다. 다음 rerun에서 재시도됩니다.')

        # 쉽먼트 재출력 탭에서 재사용하도록 세션에 저장
        st.session_state['pick_ship_to_box'] = sort_ship_to_box

        # ── 바코드 → (배대지박스, 송장박스, 수량, 상품명, 송장) 매핑 ──
        def _build_sort_state():
            state = {}  # barcode → {상품명, items: [{dapae_box, out_box, ship, needed, scanned}]}
            for _, row in df_sort.iterrows():
                bc = str(row.get('바코드', '')).strip()
                ship = str(row.get('쉽먼트운송장번호', '')).strip()
                h_qty = int(row.get('수량', 0) or 0)
                name = str(row.get('상품명', '')).strip()
                if not bc or not ship or h_qty <= 0:
                    continue
                # 부족 전용 행은 박스 채울 수가 없으므로 분류 대상에서 제외.
                # (박스번호 = '부족(-1)' 같이 부족만 있는 행)
                # '▲M7(1),부족(-1)' 처럼 박스+부족 혼합인 행은 피킹가능으로 파싱되어 포함됨.
                pick_status = str(row.get('피킹상태', '피킹가능')).strip()
                if pick_status == '부족':
                    continue
                # 출고 박스번호 (송장별 자동 부여, 국내재고 전용 송장은 None)
                out_box = sort_ship_to_box.get(ship)
                if out_box is None:
                    continue  # 국내재고/부족 전용 송장은 박스 분류 제외
                # 배대지 박스번호 (K열에서 파싱한 값: M1, W3 등)
                dapae_raw = str(row.get('박스넘버') or '').strip().upper()
                if not dapae_raw or dapae_raw == 'NAN':
                    dapae_raw = ''
                # 박스 분류에 필요한 실제 수량 = 박스내수량(파싱값)
                # H열 수량은 '주문 수량'이라 부족분 포함될 수 있음 → 박스에 들어갈 양과 다름
                # 예: '▲M7(1),부족(-1)' → 박스내수량=1, H열=2 → needed는 1이어야 박스 완료 가능
                try:
                    parsed_box_qty = int(row.get('박스내수량')) if row.get('박스내수량') is not None else None
                except (ValueError, TypeError):
                    parsed_box_qty = None
                needed_qty = parsed_box_qty if (parsed_box_qty and parsed_box_qty > 0) else h_qty
                # 시트 L열(확인수량)에서 이전 진행 상태 복원
                try:
                    prev_scanned = int(row.get('확인수량', 0) or 0)
                except (ValueError, TypeError):
                    prev_scanned = 0
                prev_scanned = max(0, min(prev_scanned, needed_qty))
                if bc not in state:
                    state[bc] = {
                        '상품명': name,
                        # 입고분류는 pick_picking_state(피킹검증 전용)가 비어있으므로
                        # 시트에서 파싱한 이미지URL을 여기 직접 보관
                        '이미지URL': str(row.get('이미지URL', '') or '').strip(),
                        'items': [],
                    }
                state[bc]['items'].append({
                    'out_box': str(out_box),      # 출고 박스번호 (1~58번)
                    'dapae_box': dapae_raw,       # 배대지 박스번호 (M1, W3 등)
                    'box_num': str(out_box),      # 기존 코드 호환: 음성/화면 안내용
                    'box_key': str(out_box),      # 기존 코드 호환
                    'sym': '',
                    'ship': ship,
                    'needed': needed_qty,
                    'scanned': prev_scanned,      # L열에서 복원한 값
                })
            return state

        # 초기화 — 데이터 세대 번호(pick_data_ver)가 바뀌었을 때만 다시 만든다
        _cur_data_ver = st.session_state.get('pick_data_ver', 0)
        if 'sort_state' not in st.session_state or st.session_state.get('sort_data_ver') != _cur_data_ver:
            st.session_state.sort_state = _build_sort_state()
            st.session_state.sort_data_ver = _cur_data_ver
            st.session_state.sort_scan_counter = 0
            st.session_state.sort_last_result = None
            # 유령 TTS 방지: 새 데이터 로드 시 played_id를 0으로 동기화
            st.session_state['_sort_tts_played_id'] = 0

        sort_state = st.session_state.sort_state

        # 이미지 URL 백그라운드 프리로드 (1회) — 스캔 시 즉시 표시
        _sort_preload_urls = sorted({
            str((v or {}).get('이미지URL', '') or '').strip()
            for v in sort_state.values()
            if str((v or {}).get('이미지URL', '') or '').strip().lower().startswith(('http://', 'https://'))
        })
        _sort_preload_sig = ','.join(_sort_preload_urls)
        if _sort_preload_sig and st.session_state.get('_sort_img_preload_sig') != _sort_preload_sig:
            import json as _json_spl
            from streamlit.components.v1 import html as _st_spl_html
            _st_spl_html(
                '<meta name="referrer" content="no-referrer">'
                f"<script>(function(){{var u={_json_spl.dumps(_sort_preload_urls)};"
                "u.forEach(function(s){try{var i=new Image();i.referrerPolicy='no-referrer';i.src=s;}catch(e){}});}})();</script>",
                height=0,
            )
            st.session_state['_sort_img_preload_sig'] = _sort_preload_sig

        # ── 배대지 박스별 집계 (작업 순서 추천용) ──
        # 배대지 박스 = K열에서 파싱된 M1, W3 등 (작업자가 물리적으로 열 박스)
        import re as _re_bq
        box_qty_map = {}  # dapae_box → {box_num, total_qty, ships, out_boxes}
        for v in sort_state.values():
            for it in v['items']:
                dp = str(it.get('dapae_box', '')).strip().upper()
                if not dp or dp == 'NAN':
                    continue
                if not _re_bq.match(r'^[A-Z]*\d+$', dp):
                    continue
                ent = box_qty_map.setdefault(dp, {
                    'box_num': dp, 'sym': '',
                    'total_qty': 0, 'ships': set(), 'out_boxes': set(),
                })
                ent['total_qty'] += it['needed']
                if it.get('ship'):
                    ent['ships'].add(it['ship'])
                if it.get('out_box'):
                    ent['out_boxes'].add(it['out_box'])

        # ── 출고(송장) 박스별 집계 (라벨 PDF / 크기 분류용) ──
        # 송장박스 = 1~58번 (우리가 준비하는 박스, 사용자 라벨 부착용)
        out_box_map = {}  # out_box(str) → {box_num, total_qty, ship}
        for v in sort_state.values():
            for it in v['items']:
                ob = str(it.get('out_box', '')).strip()
                if not ob or not ob.isdigit():
                    continue
                ent = out_box_map.setdefault(ob, {
                    'box_num': ob,
                    'total_qty': 0,
                    'ship': it.get('ship', ''),
                })
                ent['total_qty'] += it['needed']

        def _box_size(qty):
            """수량 기준으로 박스 크기 분류"""
            if qty >= 50:
                return ('대', '🟢')
            elif qty >= 30:
                return ('중', '🟡')
            else:
                return ('소', '🔵')

        def _box_sort_key(box_str):
            """박스 번호 정렬 키: (알파벳 부분, 숫자 부분) 튜플
            'W1' → ('W', 1), 'M3' → ('M', 3), '1' → ('', 1)
            """
            import re as _re_sk
            m = _re_sk.match(r'^([A-Z]*)(\d+)$', str(box_str).upper())
            if m:
                return (m.group(1), int(m.group(2)))
            return (str(box_str), 0)

        # 송장(출고) 박스 크기별 그룹 (우리가 준비할 박스 기준)
        boxes_large = []
        boxes_med = []
        boxes_small = []
        for key, info in out_box_map.items():
            size_label, _ = _box_size(info['total_qty'])
            entry = (_box_sort_key(key), info['box_num'], info['total_qty'])
            if size_label == '대':
                boxes_large.append(entry)
            elif size_label == '중':
                boxes_med.append(entry)
            else:
                boxes_small.append(entry)
        boxes_large.sort()
        boxes_med.sort()
        boxes_small.sort()

        # 송장 박스 크기 매핑 (스캔 결과 표시용)
        box_size_lookup = {}  # out_box → (size_label, emoji)
        for key, info in out_box_map.items():
            box_size_lookup[key] = _box_size(info['total_qty'])

        # ── 상단 요약 ──
        total_qty = sum(it['needed'] for v in sort_state.values() for it in v['items'])
        total_scanned = sum(it['scanned'] for v in sort_state.values() for it in v['items'])

        hh1, hh2, hh3, hh4 = st.columns(4)
        hh1.metric('출고박스', f'{len(out_box_map)}개')
        hh2.metric('배대지박스', f'{len(box_qty_map)}개')
        hh3.metric('스캔 진행', f'{total_scanned}/{total_qty}')
        hh4.metric('진행률', f'{(total_scanned/total_qty*100 if total_qty else 0):.0f}%')

        # ── 박스 크기별 준비 안내 ──
        with st.expander('📦 박스 준비 안내 (크기별)', expanded=True):
            size_rows = [
                {
                    '크기': '🟢 대형',
                    '기준': '50개 이상',
                    '개수': f'{len(boxes_large)}개',
                    '박스 번호': ', '.join(f'{b[1]}번({b[2]})' for b in boxes_large) if boxes_large else '-',
                },
                {
                    '크기': '🟡 중형',
                    '기준': '30~49개',
                    '개수': f'{len(boxes_med)}개',
                    '박스 번호': ', '.join(f'{b[1]}번({b[2]})' for b in boxes_med) if boxes_med else '-',
                },
                {
                    '크기': '🔵 소형',
                    '기준': '30개 미만',
                    '개수': f'{len(boxes_small)}개',
                    '박스 번호': ', '.join(f'{b[1]}번({b[2]})' for b in boxes_small) if boxes_small else '-',
                },
            ]
            st.dataframe(_pd2.DataFrame(size_rows), use_container_width=True, hide_index=True)
            st.caption(f'💡 총 {len(out_box_map)}개 박스 준비 ・ 박스 옆 괄호는 들어갈 총 수량')

            # ── 폼텍 3100 라벨 PDF 다운로드 (송장박스 1~N번) ──
            label_entries = []
            for key, info in sorted(out_box_map.items(),
                                    key=lambda x: _box_sort_key(x[1]['box_num'])):
                size_lbl, size_emoji = _box_size(info['total_qty'])
                label_entries.append((info['box_num'], info['total_qty'], f'{size_emoji}{size_lbl}'))
            try:
                label_pdf_buf = _box_labels_pdf_bytes_cached(label_entries)
                st.download_button(
                    label=f'🏷️ 박스 라벨 PDF 다운로드 ({len(label_entries)}장 / 폼텍 3100)',
                    data=label_pdf_buf,
                    file_name=f'box_labels_{datetime.now().strftime("%Y%m%d_%H%M")}.pdf',
                    mime='application/pdf',
                    key='sort_label_dl',
                    use_container_width=True,
                )
                st.caption('📄 폼텍 3100 (38.1×21.2mm, A4 65칸) 라벨지에 출력하세요')
            except Exception as _e:
                st.caption(f'라벨 생성 오류: {_e}')

            # #MULTI 트리거 라벨 PDF
            try:
                multi_pdf_buf = _multi_trigger_label_pdf_bytes()
                st.download_button(
                    label='🔢 다량 입력 트리거 바코드 PDF (1회만 출력)',
                    data=multi_pdf_buf,
                    file_name='multi_trigger_label.pdf',
                    mime='application/pdf',
                    key='sort_multi_trigger_dl',
                    use_container_width=True,
                )
                st.caption('📄 A4 한 장에 큰 바코드. 인쇄해서 잘 보이는 곳에 부착하세요')
            except Exception as _e:
                st.caption(f'트리거 라벨 생성 오류: {_e}')

        # ── 송장별 필요 배대지 박스 집계 (집합 커버 계산용) ──
        # 각 송장이 어느 배대지 박스에 있는지: {송장: {배대지박스 set}}
        import re as _re_ship
        ship_need_boxes = {}  # ship_id → set of dapae_box
        for v in sort_state.values():
            for it in v['items']:
                ship = it.get('ship')
                if not ship:
                    continue
                dp = str(it.get('dapae_box', '')).strip().upper()
                if not dp or dp == 'NAN':
                    continue
                if not _re_ship.match(r'^[A-Z]*\d+$', dp):
                    continue
                ship_need_boxes.setdefault(ship, set()).add(dp)

        # ── 배대지 박스별 완료 상태 계산 (sort_state 기준) ──
        # box_done: dapae_box → True(모두 스캔됨) / False
        box_done = {}
        box_progress = {}  # dapae_box → (scanned, needed)
        for _bc_bd, _v_bd in sort_state.items():
            for _it_bd in _v_bd['items']:
                _dp_bd = str(_it_bd.get('dapae_box', '')).strip().upper()
                if not _dp_bd or _dp_bd == 'NAN':
                    continue
                s, n = box_progress.get(_dp_bd, (0, 0))
                box_progress[_dp_bd] = (s + _it_bd['scanned'], n + _it_bd['needed'])
        for _dp_bd, (s, n) in box_progress.items():
            box_done[_dp_bd] = (n > 0 and s >= n)

        # ── 시작 박스 지정 ──
        st.markdown('### 🎯 작업할 배대지 박스')

        # 1순위 추천: 미완료 박스 중 수량이 가장 많은 배대지 박스
        top_box = None
        _incomplete_boxes = {k: v for k, v in box_qty_map.items() if not box_done.get(k, False)}
        if _incomplete_boxes:
            top_box_key = max(_incomplete_boxes.keys(), key=lambda k: _incomplete_boxes[k]['total_qty'])
            top_info = _incomplete_boxes[top_box_key]
            top_out_boxes = sorted(top_info['out_boxes'], key=_box_sort_key)
            out_box_str = ', '.join(f'{b}번' for b in top_out_boxes)
            st.success(
                f'🏆 **우선 작업 추천**: **{top_box_key} 박스부터 열어주세요**\n\n'
                f'→ 수량 {top_info["total_qty"]}개\n\n'
                f'📦 **준비할 출고박스 ({len(top_out_boxes)}개)**: {out_box_str}'
            )
            top_box = top_box_key
        elif box_qty_map:
            st.success('🎉 **모든 배대지 박스 완료!**')

        # 드롭다운 정렬: 미완료 우선(수량 많은 순) → 완료는 맨 아래(알파벳+번호 순)
        def _dropdown_sort_key(item):
            k, v = item
            done = box_done.get(k, False)
            # (완료 여부, -수량, box_sort_key) — 미완료(False=0)가 완료(True=1)보다 앞
            return (1 if done else 0, -v['total_qty'], _box_sort_key(k))

        recommended_boxes = sorted(box_qty_map.items(), key=_dropdown_sort_key)
        all_box_nums_sorted = [k for k, _ in recommended_boxes]

        def _fmt_box_num(x):
            info = box_qty_map[x]
            size_lbl, size_emo = _box_size(info['total_qty'])
            ship_cnt = len(info['ships'])
            done = box_done.get(x, False)
            prefix = '✅ ' if done else ''
            suffix = ' — 완료' if done else ''
            return f"{prefix}{x}번 ({info['total_qty']}개 / {size_emo}{size_lbl} / 송장 {ship_cnt}개){suffix}"

        # 창고 수용력 + 자동 추천 박스 세트
        sac0a, sac0b = st.columns([1, 3])
        with sac0a:
            capacity = st.number_input(
                '창고 수용력',
                min_value=1, max_value=10, value=1,
                key='sort_capacity',
                help='창고에 동시에 펼쳐놓을 수 있는 배대지 박스 개수 (1 = 한 박스씩 차례대로)',
            )

        # 집합 커버 추천: 수량 절대 우선 (수량 많은 박스부터 → 빠르게 출고박스 채움)
        def _recommend_box_set(target_count, already_committed=None):
            committed = set(already_committed or [])
            picks = []
            # 완료된 박스는 추천 대상에서 제외
            available = set(k for k in box_qty_map.keys() if not box_done.get(k, False)) - committed
            while len(picks) < target_count and available:
                best = None
                best_score = (-1, -1)
                for box in available:
                    qty = box_qty_map[box]['total_qty']
                    ship_cnt = len(box_qty_map[box]['ships'])
                    # 점수: (수량, -송장수) — 수량 우선 + 동률이면 송장 적은 것(단순한 박스)
                    # 송장 적을수록 다른 배대지 박스 의존성↓ → 다른 박스 동시에 안 깔아도 됨
                    score = (qty, -ship_cnt)
                    if score > best_score:
                        best_score = score
                        best = box
                if best is None:
                    break
                picks.append(best)
                available.discard(best)
            return picks

        rec_set = _recommend_box_set(int(capacity))
        if rec_set and int(capacity) > 1:
            # 수용력이 2 이상일 때만 세트 안내 (1이면 위의 1순위 추천과 동일하므로 중복 표시 안 함)
            rec_set_sorted = sorted(rec_set, key=_box_sort_key)
            completed_with_set = sum(
                1 for need in ship_need_boxes.values() if need.issubset(set(rec_set))
            )
            with sac0b:
                st.success(
                    f'💡 **지금 열 박스 {int(capacity)}개**: '
                    f'**{", ".join(str(b) + "번" for b in rec_set_sorted)}**  '
                    f'→ 이것만 열면 **{completed_with_set}개 송장** 완성'
                )
        elif int(capacity) == 1:
            with sac0b:
                st.caption('💡 **1박스씩 처리** — 수량 많은 박스부터 차례로 끝내면 출고박스가 빠르게 채워져요. 동시 작업 원하면 수용력↑')

        # ── 활성 박스 (멀티 선택) ──
        if 'sort_active_boxes' not in st.session_state:
            st.session_state.sort_active_boxes = []

        sac1, sac2, sac3 = st.columns([4, 1, 1])
        with sac1:
            active_boxes = st.multiselect(
                '🎯 지금 열어놓은 배대지 박스 (바코드 #1, #2... 찍으면 자동 추가)',
                options=all_box_nums_sorted,
                format_func=_fmt_box_num,
                default=st.session_state.sort_active_boxes,
                key='sort_active_boxes_ms',
                help='여러 박스를 동시에 선택 가능. 라벨의 #N 바코드를 찍으면 자동 토글',
            )
            # multiselect 변경 반영
            st.session_state.sort_active_boxes = active_boxes

        with sac2:
            if st.button('🔄 현황 새로고침', key='sort_refresh_view',
                         use_container_width=True,
                         help='아래 테이블(이 박스 상품, 박스별 진행 현황)을 최신 스캔 반영해서 다시 그리기'):
                st.rerun()

        with sac3:
            if st.button('🗑️ 초기화', key='sort_reset', use_container_width=True):
                st.session_state.sort_state = _build_sort_state()
                st.session_state.sort_scan_counter = 0
                st.session_state.sort_last_result = None
                st.session_state.sort_active_boxes = []
                st.rerun()

        if active_boxes:
            # 활성 박스별로 필요한 출고박스 번호 상세 표시 (출고박스 크기별 그룹핑)
            active_info_lines = []
            all_needed_out_boxes = set()
            for b in sorted(active_boxes, key=_box_sort_key):
                info = box_qty_map[b]
                size_lbl, size_emo = _box_size(info['total_qty'])
                out_boxes = sorted(info['out_boxes'], key=_box_sort_key)
                all_needed_out_boxes.update(info['out_boxes'])
                # 출고박스를 크기별로 그룹화 (box_size_lookup 사용)
                _out_L, _out_M, _out_S = [], [], []
                for _o in out_boxes:
                    _osl, _ = box_size_lookup.get(_o, ('소', '🔵'))
                    if _osl == '대':
                        _out_L.append(_o)
                    elif _osl == '중':
                        _out_M.append(_o)
                    else:
                        _out_S.append(_o)
                parts = []
                if _out_L:
                    parts.append(f"🟢대형 {len(_out_L)}개: " + ', '.join(f'{o}번' for o in _out_L))
                if _out_M:
                    parts.append(f"🟡중형 {len(_out_M)}개: " + ', '.join(f'{o}번' for o in _out_M))
                if _out_S:
                    parts.append(f"🔵소형 {len(_out_S)}개: " + ', '.join(f'{o}번' for o in _out_S))
                active_info_lines.append(
                    f"**{b} 박스** ({size_emo}{size_lbl}, {info['total_qty']}개)\n"
                    + '\n'.join(f'　→ {p}' for p in parts)
                )
            total_out_count = len(all_needed_out_boxes)
            header = f'📦 **활성 배대지 박스 {len(active_boxes)}개 — 준비할 출고박스 총 {total_out_count}개**'
            st.info(header + '\n\n' + '\n\n'.join(active_info_lines))

            # ── 활성 박스별 빠른 완료 처리 버튼 (expander 밖으로 노출) ──
            def _push_sheet_updates(updates_for_sheet):
                if not updates_for_sheet:
                    return False
                if not (st.session_state.get('pick_use_gsheet')
                        and st.session_state.get('pick_gsheet_client')
                        and st.session_state.get('pick_sheet_url_출고')
                        and st.session_state.get('pick_sheet_tab_출고')):
                    return False
                _client = st.session_state.pick_gsheet_client
                _url = st.session_state.pick_sheet_url_출고
                _tab = st.session_state.pick_sheet_tab_출고
                # 큐에 넣기만 하므로 즉시 끝난다 (플러셔가 batch로 전송)
                for _ub, _us, _uq in updates_for_sheet:
                    try:
                        pick_update_check_qty(_client, _url, _tab, _ub, _us, _uq)
                    except Exception:
                        pass
                return True

            def _collect_ship_updates(target_box):
                updates = []
                for _bc2, _v2 in sort_state.items():
                    _ship_cum = {}
                    touched_in_box = False
                    for _it2 in _v2['items']:
                        _dp2 = str(_it2.get('dapae_box', '')).strip().upper()
                        _ship2 = _it2.get('ship', '')
                        if _dp2 == target_box:
                            touched_in_box = True
                        if _ship2:
                            _ship_cum[_ship2] = _ship_cum.get(_ship2, 0) + _it2['scanned']
                    if touched_in_box:
                        for _s2, _c2 in _ship_cum.items():
                            updates.append((_bc2, _s2, _c2))
                return updates

            def _mark_box_done_bulk(target_box):
                """배대지 박스 안의 모든 항목을 완료 처리 + 시트 L열 일괄 업데이트.
                실수로 완료한 경우 되돌릴 수 있도록 원래 scanned 값을 스냅샷에 저장."""
                snapshot = []
                for _bc2, _v2 in sort_state.items():
                    for _idx2, _it2 in enumerate(_v2['items']):
                        if str(_it2.get('dapae_box', '')).strip().upper() == target_box:
                            snapshot.append((_bc2, _idx2, _it2['scanned']))
                            _it2['scanned'] = _it2['needed']
                st.session_state.pick_bulk_complete_snapshot[target_box] = snapshot
                updates_for_sheet = _collect_ship_updates(target_box)
                if _push_sheet_updates(updates_for_sheet):
                    return True, len(updates_for_sheet)
                return False, 0

            def _undo_box_done_bulk(target_box):
                """완료 처리한 박스의 원래 scanned 값을 복원 + 시트 L열 되돌림."""
                snapshot = st.session_state.pick_bulk_complete_snapshot.get(target_box)
                if not snapshot:
                    return False, 0
                for _bc2, _idx2, _orig in snapshot:
                    try:
                        sort_state[_bc2]['items'][_idx2]['scanned'] = _orig
                    except (KeyError, IndexError, TypeError):
                        pass
                updates_for_sheet = _collect_ship_updates(target_box)
                pushed = _push_sheet_updates(updates_for_sheet)
                st.session_state.pick_bulk_complete_snapshot.pop(target_box, None)
                return pushed, len(updates_for_sheet)

            st.markdown('##### ✅ 박스 완료 처리 (이미 끝낸 배대지 박스)')
            st.caption('💡 실수로 완료 처리했을 때는 옆에 나타나는 **↩️ 완료 취소** 버튼으로 되돌릴 수 있어요. '
                       '취소 시 원래 스캔 수량이 복원되고 누락 상품을 이어서 스캔할 수 있습니다.')
            _done_cols = st.columns(min(4, len(active_boxes)) or 1)
            _snapshot_map = st.session_state.get('pick_bulk_complete_snapshot', {})
            for _i, _b in enumerate(sorted(active_boxes, key=_box_sort_key)):
                _info = box_qty_map[_b]
                _cur_s = 0
                _cur_n = 0
                for _bc_x, _v_x in sort_state.items():
                    for _it_x in _v_x['items']:
                        if str(_it_x.get('dapae_box', '')).strip().upper() == _b:
                            _cur_s += _it_x['scanned']
                            _cur_n += _it_x['needed']
                _is_done = _cur_s >= _cur_n and _cur_n > 0
                _has_snapshot = _b in _snapshot_map
                with _done_cols[_i % len(_done_cols)]:
                    if _has_snapshot:
                        # 강제 완료 처리한 박스 — 취소 가능
                        if st.button(
                            f'↩️ {_b} 완료 취소',
                            key=f'sort_quick_undo_{_b}',
                            use_container_width=True,
                            type='secondary',
                            help='원래 스캔 수량을 복원하고 누락 상품을 이어서 스캔할 수 있게 합니다',
                        ):
                            ok, n = _undo_box_done_bulk(_b)
                            if ok:
                                st.success(f'↩️ {_b} 완료 취소 — 시트 L열 복원 중 ({n}건)')
                            else:
                                st.success(f'↩️ {_b} 완료 취소 (세션만, 시트 미연결)')
                            st.rerun()
                    elif _is_done:
                        st.button(
                            f'✔ {_b} 이미 완료',
                            key=f'sort_quick_done_{_b}',
                            use_container_width=True,
                            type='secondary',
                            disabled=True,
                        )
                    else:
                        if st.button(
                            f'✅ {_b} 완료 처리',
                            key=f'sort_quick_done_{_b}',
                            use_container_width=True,
                            type='primary',
                        ):
                            ok, n = _mark_box_done_bulk(_b)
                            if ok:
                                st.success(f'✅ {_b} 완료 처리 — 시트 L열 업데이트 중 ({n}건)')
                            else:
                                st.success(f'✅ {_b} 완료 처리 (세션만, 시트 미연결)')
                            st.rerun()

            # ── 활성 배대지 박스별 → 각 출고박스에 들어갈 내용물 리스트 ──
            st.markdown('#### 📋 출고박스에 담을 내용물 (배대지 박스 → 출고박스별)')
            active_set_upper = set(str(b).strip().upper() for b in active_boxes)

            # 각 출고박스별 배대지 구성 (out_box → {dapae_box → needed_qty}) — 한 번만 계산
            ob_composition = {}
            for _bc_c, _v_c in sort_state.items():
                for _it_c in _v_c['items']:
                    _ob_c = str(_it_c.get('out_box', '')).strip()
                    if not _ob_c.isdigit():
                        continue
                    _dp_c = str(_it_c.get('dapae_box', '')).strip().upper()
                    if not _dp_c or _dp_c == 'NAN':
                        continue
                    ob_composition.setdefault(_ob_c, {})
                    ob_composition[_ob_c][_dp_c] = ob_composition[_ob_c].get(_dp_c, 0) + _it_c['needed']

            for b in sorted(active_boxes, key=_box_sort_key):
                info = box_qty_map[b]
                size_lbl, size_emo = _box_size(info['total_qty'])
                # 이 배대지 박스에서 나가는 출고박스별 항목 집계
                ob_items = {}  # out_box → list of {바코드, 상품명, 필요, 스캔, 남음}
                for _bc, _v in sort_state.items():
                    for _it in _v['items']:
                        _dp = str(_it.get('dapae_box', '')).strip().upper()
                        if _dp != b:
                            continue
                        _ob = str(_it.get('out_box', '')).strip()
                        if not _ob.isdigit():
                            continue
                        ob_items.setdefault(_ob, []).append({
                            '바코드': _bc,
                            '상품명': _v['상품명'][:35],
                            '필요': _it['needed'],
                            '스캔': _it['scanned'],
                            '남음': max(0, _it['needed'] - _it['scanned']),
                        })
                # 합계 표시
                total_ob_n = sum(it['필요'] for items in ob_items.values() for it in items)
                total_ob_s = sum(it['스캔'] for items in ob_items.values() for it in items)
                head_pct = (total_ob_s / total_ob_n * 100) if total_ob_n else 0
                head_status = '✅' if total_ob_s >= total_ob_n and total_ob_n > 0 else ('🔄' if total_ob_s > 0 else '⬜')
                exp_label = (
                    f'{head_status} {b} 배대지박스 ({size_emo}{size_lbl}, {total_ob_s}/{total_ob_n}개, '
                    f'{head_pct:.0f}%) → 출고박스 {len(ob_items)}개'
                )
                with st.expander(exp_label, expanded=False):
                    # ── 박스 일괄 완료 처리 (이미 물리적으로 끝난 박스용) ──
                    btn_col1, btn_col2 = st.columns([2, 1])
                    _has_snap_b = b in st.session_state.get('pick_bulk_complete_snapshot', {})
                    with btn_col1:
                        if _has_snap_b:
                            st.caption('↩️ 강제 완료 처리된 박스 — 우측 **완료 취소** 버튼으로 원래 스캔 수량 복원 + 누락 상품 이어서 스캔 가능')
                        else:
                            st.caption('💡 이 배대지 박스가 이미 물리적으로 완료된 경우 → 우측 버튼으로 시트 L열 일괄 채움')
                    with btn_col2:
                        if _has_snap_b:
                            if st.button(
                                f'↩️ {b} 완료 취소',
                                key=f'sort_box_undo_{b}',
                                use_container_width=True,
                                type='secondary',
                                help='원래 스캔 수량 복원 + 누락 상품 이어서 스캔 가능',
                            ):
                                ok, n = _undo_box_done_bulk(b)
                                if ok:
                                    st.success(f'↩️ {b} 배대지박스 완료 취소 — 시트 L열 복원 중 ({n}건)')
                                else:
                                    st.success(f'↩️ {b} 배대지박스 완료 취소 (세션만, 시트 미연결)')
                                st.rerun()
                        else:
                            if st.button(
                                f'✅ {b} 박스 완료 처리',
                                key=f'sort_box_done_{b}',
                                use_container_width=True,
                                type='primary' if total_ob_s < total_ob_n else 'secondary',
                            ):
                                ok, n = _mark_box_done_bulk(b)
                                if ok:
                                    st.success(f'✅ {b} 배대지박스 완료 처리 — 시트 L열 업데이트 중 ({n}건)')
                                else:
                                    st.success(f'✅ {b} 배대지박스 완료 처리 (세션만, 시트 미연결)')
                                st.rerun()

                    # ── 🎯 상품 단위 리스트 (수량 많은 순) — 제일 많은 것부터 준비 ──
                    sku_rows = []
                    for _bc_s, _v_s in sort_state.items():
                        total_in_box = 0
                        total_scanned_in_box = 0
                        out_box_break = {}  # out_box → qty
                        for _it_s in _v_s['items']:
                            _dp_s = str(_it_s.get('dapae_box', '')).strip().upper()
                            if _dp_s != b:
                                continue
                            total_in_box += _it_s['needed']
                            total_scanned_in_box += _it_s['scanned']
                            _ob_s = str(_it_s.get('out_box', '')).strip()
                            if _ob_s:
                                out_box_break[_ob_s] = out_box_break.get(_ob_s, 0) + _it_s['needed']
                        if total_in_box <= 0:
                            continue
                        # 출고박스 브레이크다운 (박스 번호 순)
                        ob_parts = sorted(out_box_break.items(), key=lambda x: _box_sort_key(x[0]))
                        ob_str = ', '.join(f'{k}번({v})' for k, v in ob_parts)
                        _sku_status = '✅' if total_scanned_in_box >= total_in_box else (
                            '🔄' if total_scanned_in_box > 0 else '⬜'
                        )
                        sku_rows.append({
                            '상태': _sku_status,
                            '이 박스 수량': total_in_box,
                            '스캔': total_scanned_in_box,
                            '남음': max(0, total_in_box - total_scanned_in_box),
                            '바코드': _bc_s,
                            '상품명': _v_s['상품명'],
                            '출고박스 배분': ob_str,
                        })
                    # 수량 많은 순 정렬 (수량 동률이면 남은 수량 많은 순)
                    sku_rows.sort(key=lambda r: (-r['이 박스 수량'], -r['남음']))
                    if sku_rows:
                        st.markdown('##### 🎯 이 박스 상품 (수량 많은 순) — 제일 많이 담긴 것부터 준비')
                        st.dataframe(
                            _pd2.DataFrame(sku_rows),
                            use_container_width=True, hide_index=True,
                            height=min(500, len(sku_rows) * 38 + 40),
                        )

                    st.markdown('---')
                    st.markdown('##### 📦 출고박스별 상세 (출고박스마다 담을 상품)')
                    for _ob in sorted(ob_items.keys(), key=_box_sort_key):
                        _items = ob_items[_ob]
                        _n = sum(x['필요'] for x in _items)
                        _s = sum(x['스캔'] for x in _items)
                        _ob_status = '✅' if _s >= _n and _n > 0 else ('🔄' if _s > 0 else '⬜')
                        _size_lbl2, _size_emo2 = box_size_lookup.get(_ob, ('', ''))
                        # 이 출고박스에 들어가는 모든 배대지 박스 구성 (현재 b는 굵게 표시)
                        _comp = ob_composition.get(_ob, {})
                        _all_parts = sorted(_comp.items(), key=lambda x: _box_sort_key(x[0]))
                        _total_qty = sum(v for _, v in _all_parts)
                        if _all_parts:
                            _parts_str = ', '.join(
                                (f'**{k}({v})**' if k == b else f'{k}({v})')
                                for k, v in _all_parts
                            )
                            _tail = f"  ·  📦 **구성 {_total_qty}개**: {_parts_str}"
                        else:
                            _tail = ''
                        st.markdown(
                            f"**{_ob_status} {_ob}번 출고박스** "
                            f"{_size_emo2}{_size_lbl2} — {_s}/{_n}개 ({len(_items)} SKU)"
                            f"{_tail}"
                        )
                        st.dataframe(
                            _pd2.DataFrame(_items),
                            use_container_width=True, hide_index=True,
                            height=min(250, len(_items) * 38 + 40),
                        )

        st.markdown('---')

        # ── 바코드 스캔 (fragment으로 감싸서 전체 앱 리런 없이 조각만 재실행) ──
        # 외부 테이블(이 박스 상품 리스트 등)은 '현황 새로고침' 버튼으로 수동 갱신.
        _use_fragment = getattr(st, 'fragment', lambda f: f)

        @_use_fragment
        def _scan_fragment():
            # ── 바코드 스캔 ──
            # 다량 모드 상태
            if 'sort_next_qty' not in st.session_state:
                st.session_state.sort_next_qty = 1
            if 'sort_qty_input_mode' not in st.session_state:
                st.session_state.sort_qty_input_mode = False

            # ── 입력 처리는 전부 on_change 콜백에서 (피킹검증/재고 탭과 같은 방식) ──
            # 예전: 스캔마다 key 가 바뀐 새 입력창 + st.rerun(scope='fragment') 1회 추가
            #       = 스캔 1건에 조각 실행 2번, 입력창 DOM 교체로 포커스/글자 유실.
            # 지금: 콜백이 조각 실행 전에 한 번 돌고, 조각은 결과만 1번 그린다.
            def _on_sort_qty_enter():
                raw = str(st.session_state.get('sort_qty_input', '') or '').strip()
                st.session_state['sort_qty_input'] = ''
                if not raw:
                    return
                try:
                    qty_val = int(raw)
                except ValueError:
                    st.session_state['sort_qty_error'] = '숫자만 입력 가능합니다'
                    return
                if qty_val < 1:
                    st.session_state['sort_qty_error'] = '1 이상을 입력하세요'
                    return
                st.session_state.pop('sort_qty_error', None)
                st.session_state.sort_next_qty = qty_val
                st.session_state.sort_qty_input_mode = False
                # 수량 확정 안내 — 이전 결과(다량 입력 모드)가 다시
                # 재생되지 않도록 결과를 교체한다
                st.session_state.sort_last_result = {
                    'status': 'qty_set',
                    'barcode': '',
                    'message': f'📦 다음 스캔: {qty_val}개',
                    'detail': '바코드를 스캔해 주세요',
                }
                st.session_state.sort_scan_counter += 1

            def _on_sort_qty_reset():
                st.session_state.sort_next_qty = 1

            # 수량 입력 모드: 큰 알림 + 전체 너비 입력창
            if st.session_state.sort_qty_input_mode:
                st.warning('🔢 **수량을 입력하세요** — 숫자 입력 후 Enter')
                st.text_input(
                    '다량 수량',
                    key='sort_qty_input',
                    placeholder='숫자 입력 후 Enter (예: 50)',
                    label_visibility='collapsed',
                    on_change=_on_sort_qty_enter,
                )
                if st.session_state.get('sort_qty_error'):
                    st.error(st.session_state['sort_qty_error'])

            # 수량 표시 + 1개 모드 리셋 버튼
            qcol1, qcol2 = st.columns([1, 1])
            with qcol1:
                if st.session_state.sort_next_qty > 1:
                    st.markdown(
                        f'<div style="background:#f59e0b;color:white;padding:0.5rem;border-radius:6px;text-align:center;font-weight:bold;font-size:1.1rem">'
                        f'📦 다음 스캔: {st.session_state.sort_next_qty}개'
                        f'</div>', unsafe_allow_html=True)
                else:
                    st.markdown(
                        '<div style="background:#e5e7eb;padding:0.5rem;border-radius:6px;text-align:center">'
                        '1개 모드'
                        '</div>', unsafe_allow_html=True)
            with qcol2:
                if st.session_state.sort_next_qty > 1 and not st.session_state.sort_qty_input_mode:
                    st.button('🔄 1개 모드로 복귀', key='sort_qty_reset', use_container_width=True,
                              on_click=_on_sort_qty_reset)

            def _process_sort_scan(bc):
                bc = bc.strip()

                # #MULTI 트리거 → 수량 입력 모드 진입
                if bc.upper() == '#MULTI':
                    st.session_state.sort_qty_input_mode = True
                    st.session_state.sort_next_qty = 1
                    return {
                        'status': 'multi_trigger',
                        'barcode': bc,
                        'message': '🔢 다량 입력 모드',
                        'detail': '수량을 입력한 후 상품 바코드를 스캔하세요',
                    }

                # #W1, #M2, #1 바코드 → 박스 토글
                import re as _re_bc
                box_label_match = _re_bc.match(r'^#([A-Za-z]*\d+)$', bc)
                if box_label_match:
                    bn = box_label_match.group(1).upper()
                    if bn in box_qty_map:
                        cur = list(st.session_state.get('sort_active_boxes', []))
                        if bn in cur:
                            cur.remove(bn)
                            msg_suffix = '제외'
                        else:
                            cur.append(bn)
                            msg_suffix = '선택'
                        st.session_state.sort_active_boxes = cur
                        # 화면의 multiselect 위젯 값도 같이 맞춘다. 이걸 안 하면 다음 전체
                        # rerun 때 위젯에 남아있던 옛 값이 default 를 이기고 그 값으로
                        # sort_active_boxes 를 다시 덮어써서, 바코드로 고른 박스가 풀린다.
                        try:
                            st.session_state['sort_active_boxes_ms'] = list(cur)
                        except Exception:
                            pass
                        info = box_qty_map[bn]
                        size_lbl, size_emo = _box_size(info['total_qty'])
                        return {
                            'status': 'box_toggle',
                            'barcode': bc,
                            'box_num': bn,
                            'message': f'📦 {bn}번 박스 {msg_suffix}',
                            'detail': f'{size_emo}{size_lbl}형 · {info["total_qty"]}개 · 송장 {len(info["ships"])}개',
                        }
                    else:
                        return {'status': 'error', 'barcode': bc,
                                'message': f'🚨 {bn}번 박스 없음', 'detail': bc}

                if bc not in sort_state:
                    return {'status': 'error', 'barcode': bc,
                            'message': '🚨 출고지시서에 없는 바코드',
                            'detail': bc}
                item_data = sort_state[bc]
                # 아직 채워야 할 박스 중 후보 선택
                candidates = [it for it in item_data['items'] if it['scanned'] < it['needed']]
                if not candidates:
                    return {'status': 'over', 'barcode': bc,
                            '상품명': item_data['상품명'],
                            'message': '⚠️ 이 상품은 모두 분류 완료',
                            'detail': item_data['상품명'][:35]}

                # 활성 배대지 박스 집합 필터: 선택된 배대지 박스들의 상품만 유효
                _active_boxes_set = set(
                    str(b).strip().upper() for b in st.session_state.get('sort_active_boxes', [])
                )
                if _active_boxes_set:
                    box_candidates = []
                    for it in candidates:
                        it_dp = str(it.get('dapae_box', '')).strip().upper()
                        if it_dp and it_dp in _active_boxes_set:
                            box_candidates.append(it)
                    if not box_candidates:
                        # 이 상품의 배대지 박스가 활성 목록에 없음
                        other_boxes = sorted(set(
                            str(it.get('dapae_box', ''))
                            for it in item_data['items']
                            if it.get('dapae_box')
                        ))
                        return {'status': 'wrong_box', 'barcode': bc,
                                '상품명': item_data['상품명'],
                                'message': '🚨 활성 배대지 박스에 없는 상품',
                                'detail': f'이 상품은 {", ".join(other_boxes)} 배대지 박스에 있음'}
                    candidates = box_candidates

                # 다량 모드: next_qty 만큼 차감 (여러 박스에 걸쳐 자동 분배)
                requested_qty = int(st.session_state.get('sort_next_qty', 1))
                requested_qty = max(1, requested_qty)
                processed = 0
                last_target = candidates[0]
                touched_ships = set()  # 이번 스캔으로 영향받은 송장(시트 L열 업데이트용)
                # 후보들을 순회하며 각 박스 채워가기
                remaining_to_scan = requested_qty
                idx = 0
                while remaining_to_scan > 0 and idx < len(candidates):
                    it = candidates[idx]
                    space = it['needed'] - it['scanned']
                    if space <= 0:
                        idx += 1
                        continue
                    take = min(space, remaining_to_scan)
                    it['scanned'] += take
                    processed += take
                    remaining_to_scan -= take
                    last_target = it
                    if it.get('ship'):
                        touched_ships.add(it['ship'])
                    if it['scanned'] >= it['needed']:
                        idx += 1

                # 영향받은 송장별 누적 스캔수량 (L열 덮어쓰기용)
                touched_updates = []
                if touched_ships:
                    ship_cum = {}
                    for _it in item_data['items']:
                        _s = _it.get('ship')
                        if _s in touched_ships:
                            ship_cum[_s] = ship_cum.get(_s, 0) + _it['scanned']
                    for _s, _c in ship_cum.items():
                        touched_updates.append((bc, _s, _c))

                # 모두 차감 후 남은 수량 (수량 초과)
                over_qty = requested_qty - processed

                # 다량 모드는 1회용: 원래대로 복귀
                st.session_state.sort_next_qty = 1
                st.session_state.sort_qty_input_mode = False

                # 이 박스(box_key)가 다 채워졌는지 확인
                target_box_key = last_target['box_key']
                box_complete = True
                for _bc, _v in sort_state.items():
                    for _it in _v['items']:
                        if _it['box_key'] == target_box_key and _it['scanned'] < _it['needed']:
                            box_complete = False
                            break
                    if not box_complete:
                        break

                return {
                    'status': 'ok',
                    'barcode': bc,
                    '상품명': item_data['상품명'],
                    'box_key': last_target['box_key'],
                    'box_num': last_target['box_num'],
                    'sym': last_target['sym'],
                    'ship': last_target['ship'],
                    'remaining': last_target['needed'] - last_target['scanned'],
                    'box_complete': box_complete,
                    'processed_qty': processed,
                    'over_qty': over_qty,
                    'touched_updates': touched_updates,  # [(bc, ship, cum_scanned), ...]
                }

            def _on_sort_scan_enter():
                sort_scanned = str(st.session_state.get('sort_scan_input', '') or '').strip()
                # 위젯 자기 값 비우기 — 콜백 안에서만 안정적으로 동작
                st.session_state['sort_scan_input'] = ''
                if not sort_scanned:
                    return
                scan_result = _process_sort_scan(sort_scanned)
                st.session_state.sort_last_result = scan_result
                st.session_state.sort_scan_counter += 1

                # 구글 시트 L열(확인 수량) 업데이트 (백그라운드)
                # touched_updates: 이번 스캔으로 영향받은 (bc, ship, 누적스캔수량) 목록
                # 누적값으로 덮어쓰므로 시트를 다시 로드해도 진행 상태가 보존됨
                if (scan_result.get('status') == 'ok'
                        and st.session_state.get('pick_use_gsheet')
                        and st.session_state.get('pick_gsheet_client')
                        and st.session_state.get('pick_sheet_url_출고')
                        and st.session_state.get('pick_sheet_tab_출고')):
                    _updates = scan_result.get('touched_updates') or []
                    if not _updates:
                        # fallback: 단일 스캔 시 last_target 정보로
                        _bc = scan_result.get('barcode', '')
                        _ship = scan_result.get('ship', '')
                        _qty = scan_result.get('processed_qty', 1)
                        _updates = [(_bc, _ship, _qty)]
                    _client = st.session_state.pick_gsheet_client
                    _url = st.session_state.pick_sheet_url_출고
                    _tab = st.session_state.pick_sheet_tab_출고
                    # 큐에 넣기만 하므로 즉시 끝난다 (플러셔가 batch로 전송)
                    for _ub, _us, _uq in _updates:
                        try:
                            pick_update_check_qty(_client, _url, _tab, _ub, _us, _uq)
                        except Exception:
                            pass

            # key 를 고정하면 스캔 후에도 같은 입력창 DOM 이 유지되어 포커스가 끊기지 않는다
            st.text_input(
                '🔫 바코드 스캔',
                key='sort_scan_input',
                placeholder='박스에서 꺼낸 상품의 바코드를 스캔하세요 (여러 개면 #MULTI 바코드 먼저)',
                on_change=_on_sort_scan_enter,
            )

            # 자동 포커스 (multiselect/number input 상호작용 중에는 포커스 안 가로챔)
            from streamlit.components.v1 import html as _sort_html
            _sort_html("""<script>
            (function(){
                const doc = window.parent.document;
                function findScan(){
                    const inputs = doc.querySelectorAll('input[type="text"]');
                    for (const inp of inputs){
                        if (inp.placeholder && inp.placeholder.includes('박스에서 꺼낸')) return inp;
                    }
                    return null;
                }
                function findQtyInput(){
                    const inputs = doc.querySelectorAll('input[type="text"]');
                    for (const inp of inputs){
                        if (inp.placeholder && inp.placeholder.includes('숫자 입력')) return inp;
                    }
                    return null;
                }
                function isInteractingOther(){
                    const active = doc.activeElement;
                    if (!active) return false;
                    const tag = (active.tagName || '').toLowerCase();
                    // number input (수량), textarea, button
                    if (tag === 'button' || tag === 'textarea') return true;
                    if (tag === 'input' && active.type !== 'text') return true;
                    // Streamlit BaseWeb select (multiselect) 내부
                    if (active.closest) {
                        if (active.closest('[data-baseweb="select"]')) return true;
                        if (active.closest('[data-baseweb="popover"]')) return true;
                        if (active.closest('[role="listbox"]')) return true;
                        if (active.closest('[role="combobox"]')) return true;
                    }
                    return false;
                }
                function focusScan(){
                    // 수량 입력창이 표시되어 있으면 우선 그 창으로 포커스
                    const qty = findQtyInput();
                    if (qty) {
                        if (doc.activeElement !== qty) qty.focus({preventScroll: true});
                        return;
                    }
                    const inp = findScan();
                    if (!inp) return;
                    if (doc.activeElement === inp) return;
                    if (isInteractingOther()) return;
                    // 팝오버/드롭다운이 열려있으면 포커스 안 함
                    if (doc.querySelector('[data-baseweb="popover"]')) return;
                    // preventScroll: 포커스 이동 시 브라우저 자동 스크롤 방지
                    inp.focus({preventScroll: true});
                }
                focusScan();
                if (window._sortFocusInterval) clearInterval(window._sortFocusInterval);
                window._sortFocusInterval = setInterval(focusScan, 500);
            })();
            </script>""", height=0)

            # ── 결과 표시 + 음성 ──
            r = st.session_state.get('sort_last_result')
            # 같은 결과를 rerun 시 반복 재생/표시하지 않도록 scan_id로 가드
            _last_played_id = st.session_state.get('_sort_tts_played_id', -1)
            _current_scan_id = st.session_state.sort_scan_counter
            _is_fresh_scan = (r is not None and _current_scan_id != _last_played_id)

            # ── 박스 완료 음성 발생 전에 정확성 재검증 (거짓 완료 방지) ──
            try:
                _box_complete_validated = False
                _miss_items = []
                if r and _is_fresh_scan and isinstance(r, dict) and r.get('box_complete'):
                    _bk = r.get('box_key')
                    for _bc_v, _v_v in sort_state.items():
                        for _it_v in _v_v.get('items', []):
                            if _it_v.get('box_key') == _bk and _it_v.get('scanned', 0) < _it_v.get('needed', 0):
                                _miss_items.append({
                                    '바코드': _bc_v,
                                    '상품명': _v_v.get('상품명', '')[:30],
                                    '필요': _it_v.get('needed', 0),
                                    '스캔': _it_v.get('scanned', 0),
                                    '남음': _it_v.get('needed', 0) - _it_v.get('scanned', 0),
                                })
                    _box_complete_validated = (len(_miss_items) == 0)
                    if not _box_complete_validated:
                        # 완료 검출 거짓이었으므로 결과 강제 수정 → 음성/UI 다른 분기로
                        r = dict(r)  # 사본
                        r['box_complete'] = False
                        r['_box_validation_failed'] = True
                        r['_miss_items'] = _miss_items
            except Exception:
                # 검증 실패하더라도 음성/UI는 진행 (이전 동작 유지)
                pass

            if r and _is_fresh_scan:
                _KOR_ALPHA = {
                    'W': '더블유', 'M': '엠', 'L': '엘', 'S': '에스',
                    'A': '에이', 'B': '비', 'C': '씨', 'D': '디', 'E': '이', 'F': '에프',
                    'G': '지', 'H': '에이치', 'I': '아이', 'J': '제이', 'K': '케이',
                    'N': '엔', 'O': '오', 'P': '피', 'Q': '큐', 'R': '알', 'T': '티',
                    'U': '유', 'V': '브이', 'X': '엑스', 'Y': '와이', 'Z': '지',
                }
                def _num_to_kor(n):
                    # 워밍업(_edge_tts_warmup_phrases)과 같은 함수를 써야 캐시 키가 일치한다
                    return _tts_num_to_kor(n)
                def _box_to_kor(n_str):
                    """W1 → '더블유 일', M3 → '엠 삼', 1 → '일'"""
                    import re as _re_k
                    s = str(n_str).strip().upper()
                    m = _re_k.match(r'^([A-Z]*)(\d+)$', s)
                    if not m:
                        return s
                    alpha_part = m.group(1)
                    num_part = int(m.group(2))
                    alpha_kor = ' '.join(_KOR_ALPHA.get(ch, ch) for ch in alpha_part)
                    num_kor = _num_to_kor(num_part)
                    if alpha_kor:
                        return f'{alpha_kor} {num_kor}'
                    return num_kor

                if r['status'] == 'multi_trigger':
                    st.markdown(
                        '<div class="scan-complete" style="background:#f59e0b;color:white;padding:2rem;border-radius:12px;text-align:center;border-left:8px solid #d97706">'
                        '<div style="font-size:2.2rem;font-weight:bold;">🔢 수량을 입력하세요</div>'
                        '<div style="font-size:1.1rem;margin-top:0.8rem;">수량 입력 후 "수량 확정" 또는 바로 상품 바코드 스캔</div>'
                        '</div>',
                        unsafe_allow_html=True)
                    speak = '수량을 입력하세요'
                elif r['status'] == 'qty_set':
                    st.markdown(
                        f'<div class="scan-complete" style="background:#f59e0b;color:white;padding:2rem;border-radius:12px;text-align:center;border-left:8px solid #d97706">'
                        f'<div style="font-size:2.2rem;font-weight:bold;">{r["message"]}</div>'
                        f'<div style="font-size:1.1rem;margin-top:0.8rem;">바코드를 스캔해 주세요</div>'
                        f'</div>',
                        unsafe_allow_html=True)
                    speak = '바코드를 스캔해 주세요'
                elif r['status'] == 'box_toggle':
                    st.markdown(
                        f'<div class="scan-complete" style="background:#3b82f6;color:white;padding:1.5rem;border-radius:10px;text-align:center;border-left:8px solid #1e40af">'
                        f'<div style="font-size:1.8rem;font-weight:bold;">{r["message"]}</div>'
                        f'<div style="font-size:1rem;margin-top:0.5rem;opacity:0.95;">{r["detail"]}</div>'
                        f'</div>',
                        unsafe_allow_html=True)
                    _kor_bt = _box_to_kor(str(r['box_num']))
                    speak = f'{_kor_bt}번'
                elif r['status'] == 'error':
                    st.markdown(
                        f'<div class="scan-error"><strong style="font-size:1.4rem;">📥 보류</strong><br>'
                        f'쉽먼트 정보 없음 - 따로 보관<br>{r["detail"]}</div>',
                        unsafe_allow_html=True)
                    speak = '보류'
                elif r['status'] == 'wrong_box':
                    st.markdown(
                        f'<div class="scan-error"><strong style="font-size:1.3rem;">{r["message"]}</strong><br>{r["detail"]}</div>',
                        unsafe_allow_html=True)
                    speak = '다른 박스 상품'
                elif r['status'] == 'over':
                    st.markdown(f'<div class="scan-warning"><strong style="font-size:1.2rem;">{r["message"]}</strong><br>{r["detail"]}</div>', unsafe_allow_html=True)
                    speak = '분류 완료'
                else:
                    box_num_str = str(r['box_num']).strip().upper()
                    kor_n = _box_to_kor(box_num_str)
                    size_label, size_emoji = box_size_lookup.get(box_num_str, ('', ''))
                    size_str = f' ({size_emoji}{size_label}형)' if size_label else ''
                    if r.get('_box_validation_failed'):
                        # 완료 검출됐으나 검증 결과 미스캔 발견 → 경고 모드
                        _miss_n = len(r.get('_miss_items') or [])
                        st.markdown(
                            f'<div class="scan-warning" style="background:#f59e0b;color:white;padding:1.5rem;border-radius:10px;text-align:center;border-left:8px solid #d97706">'
                            f'<div style="font-size:1.8rem;font-weight:bold;">⚠️ {box_num_str}번 — 완료 같지만 {_miss_n}건 미스캔!</div>'
                            f'<div style="font-size:1rem;margin-top:0.5rem;">아래 미스캔 항목을 다시 확인해주세요.</div>'
                            f'</div>',
                            unsafe_allow_html=True)
                        speak = f'{kor_n}번 미스캔 있음'
                        import pandas as _pd_miss2
                        st.dataframe(_pd_miss2.DataFrame(r['_miss_items']),
                                     use_container_width=True, hide_index=True)
                    elif r.get('box_complete'):
                        # 박스 완료! 큰 알림 + 포장 안내
                        st.markdown(
                            f'<div class="scan-complete" style="background:#10b981;color:white;padding:2rem;border-radius:12px;text-align:center;border-left:8px solid #059669">'
                            f'<div style="font-size:2.5rem;font-weight:bold;">🎉 {box_num_str}번 {size_str} 완료!</div>'
                            f'<div style="font-size:1.3rem;margin-top:0.8rem;">📦 포장하고 출고지시서 종이를 끼워주세요</div>'
                            f'<div style="font-size:1rem;margin-top:0.5rem;opacity:0.9;word-break:break-all;">마지막 상품: {r["상품명"]}</div>'
                            f'</div>',
                            unsafe_allow_html=True)
                        speak = f'{kor_n}번 완료. 포장하세요'
                    else:
                        processed_qty = r.get('processed_qty', 1)
                        over_qty = r.get('over_qty', 0)
                        qty_str = f' × {processed_qty}개' if processed_qty > 1 else ''
                        over_str = f' ⚠️ {over_qty}개 초과' if over_qty > 0 else ''
                        _bc_sort = str(r.get('barcode', '') or '')
                        # 이미지: 입고분류는 sort_state에서 (pick_picking_state는 피킹검증 전용이라 비어있음)
                        _img_sort = ''
                        if _bc_sort:
                            _img_sort = str((sort_state.get(_bc_sort) or {}).get('이미지URL', '') or '').strip()
                            if not _img_sort and _bc_sort in st.session_state.get('pick_picking_state', {}):
                                _img_sort = str(st.session_state.pick_picking_state[_bc_sort].get('이미지URL', '') or '').strip()
                        render_scan_result_card(
                            'ok',
                            f'✅ {box_num_str}번{size_str}{qty_str} → {r["상품명"]}',
                            f'송장 {r["ship"][-6:]} | 남은 수량: {r["remaining"]}개{over_str}',
                            barcode=_bc_sort, img_url=_img_sort, img_size=110, height=215,
                        )
                        if processed_qty > 1:
                            speak = f'{kor_n}번 {processed_qty}개'
                        else:
                            speak = f'{kor_n}번'

                # 소리 + 음성
                scan_id_s = st.session_state.sort_scan_counter
                beep_js = "o.frequency.value=880;g.gain.value=0.3;o.start();setTimeout(()=>g.gain.value=0,150);setTimeout(()=>o.stop(),200);"
                if r['status'] in ('error', 'wrong_box'):
                    beep_js = "o.type='square';o.frequency.value=200;g.gain.value=0.5;o.start();setTimeout(()=>{o.frequency.value=150},150);setTimeout(()=>g.gain.value=0,500);setTimeout(()=>o.stop(),600);"
                elif r['status'] == 'over':
                    beep_js = "o.type='sawtooth';o.frequency.value=400;g.gain.value=0.4;o.start();setTimeout(()=>g.gain.value=0,300);setTimeout(()=>o.stop(),400);"
                elif r.get('box_complete'):
                    # 박스 완료 - 축하 멜로디 (3음)
                    beep_js = ("o.frequency.value=523;g.gain.value=0.4;o.start();"
                               "setTimeout(()=>{o.frequency.value=659},120);"
                               "setTimeout(()=>{o.frequency.value=784},240);"
                               "setTimeout(()=>g.gain.value=0,400);"
                               "setTimeout(()=>o.stop(),500);")
                _tts_js_s = _tts_ko_script(speak)
                _sort_html(f"""<script>
                // sort_id={scan_id_s}
                try{{var a=new(window.AudioContext||window.webkitAudioContext)();var o=a.createOscillator();var g=a.createGain();o.connect(g);g.connect(a.destination);{beep_js}}}catch(e){{}}
                {_tts_js_s}
                </script>""", height=0)
                # 이 scan_id에 대해서는 TTS 재생 완료로 마크 → 같은 rerun 반복돼도 다시 안 울림
                st.session_state['_sort_tts_played_id'] = _current_scan_id

            # ── 박스별 진행 현황 (fragment 안: 매 스캔마다 갱신됨) ──
            # 성능 메모: 이 블록은 스캔마다 다시 그려진다. 예전에는 박스마다
            # st.dataframe(Arrow 직렬화 + 브라우저 그리드) 을 하나씩 만들어서
            # 박스가 40~60개면 스캔 한 번에 그리드 60개를 다시 만들었다.
            # 지금은 '진행 중' 박스만 그리드로, 완료/대기 박스는 가벼운 마크다운 표로 그린다.
            # 상세 행도 박스마다 sort_state 전체를 다시 돌지 않고 한 번에 모은다.
            st.markdown('---')
            st.subheader('📋 박스별 진행 현황')
            box_summary = {}
            box_detail_rows = {}   # box_key → [상세 행]
            for bc, v in sort_state.items():
                for it in v['items']:
                    key = it['box_key']
                    ent = box_summary.setdefault(key, {
                        'box_num': it['box_num'], 'sym': it['sym'],
                        'needed': 0, 'scanned': 0, 'sku_total': 0, 'sku_done': 0,
                        'ships': set(),
                    })
                    ent['needed'] += it['needed']
                    ent['scanned'] += it['scanned']
                    ent['sku_total'] += 1
                    if it.get('ship'):
                        ent['ships'].add(it['ship'])
                    if it['scanned'] >= it['needed']:
                        ent['sku_done'] += 1
                    _d_s = it['scanned']
                    _d_n = it['needed']
                    box_detail_rows.setdefault(key, []).append({
                        '상태': '✅' if _d_s >= _d_n and _d_n > 0 else ('🔄' if _d_s > 0 else '⬜'),
                        '바코드': bc,
                        '상품명': v['상품명'],
                        '배대지박스': it.get('dapae_box', '') or '-',
                        '송장': it.get('ship', ''),
                        '필요': _d_n,
                        '스캔': _d_s,
                        '남음': max(0, _d_n - _d_s),
                    })

            def _rows_to_md_table(rows):
                """가벼운 마크다운 표 (그리드 없이) — 완료/대기 박스 상세용."""
                cols = ['상태', '바코드', '상품명', '배대지박스', '송장', '필요', '스캔', '남음']
                out = ['| ' + ' | '.join(cols) + ' |', '|' + '---|' * len(cols)]
                for r_ in rows:
                    out.append('| ' + ' | '.join(
                        str(r_[c]).replace('|', '｜') for c in cols) + ' |')
                return '\n'.join(out)

            def _box_prog_sort_key(kv):
                key, ent = kv
                done = (ent['needed'] > 0 and ent['scanned'] >= ent['needed'])
                return (1 if done else 0, ent['sym'], _box_sort_key(ent['box_num']))

            st.caption('💡 각 박스 헤더를 클릭하면 그 박스에 들어가는 모든 상품이 펼쳐집니다')
            for key, ent in sorted(box_summary.items(), key=_box_prog_sort_key):
                pct = (ent['scanned'] / ent['needed'] * 100) if ent['needed'] else 0
                is_done = (ent['needed'] > 0 and ent['scanned'] >= ent['needed'])
                if is_done:
                    status = '✅ 완료'
                elif ent['scanned'] > 0:
                    status = f'🔄 {pct:.0f}%'
                else:
                    status = '⬜ 대기'
                _bn_key_p = str(ent['box_num']).strip().upper()
                _size_label, _size_emoji = box_size_lookup.get(_bn_key_p, ('', ''))
                _size_str = f'{_size_emoji}{_size_label}' if _size_label else ''
                exp_title = (
                    f"{status} · **{ent['box_num']}번 박스** {_size_str} · "
                    f"{ent['scanned']}/{ent['needed']}개 · "
                    f"SKU {ent['sku_done']}/{ent['sku_total']} · "
                    f"송장 {len(ent['ships'])}개"
                )
                in_progress = (not is_done and ent['scanned'] > 0)
                with st.expander(exp_title, expanded=in_progress):
                    detail_rows = box_detail_rows.get(key, [])
                    detail_rows.sort(key=lambda r: (
                        _box_sort_key(r['배대지박스']) if r['배대지박스'] != '-' else ('ZZZ', 99999),
                        r['송장'], r['바코드'],
                    ))
                    if not detail_rows:
                        st.caption('이 박스에 해당하는 항목이 없습니다')
                    elif in_progress:
                        # 지금 채우는 박스만 그리드 — 스캔마다 갱신되는 것이 중요한 곳
                        st.dataframe(
                            _pd2.DataFrame(detail_rows),
                            use_container_width=True, hide_index=True,
                            height=min(500, len(detail_rows) * 38 + 40),
                        )
                    else:
                        st.markdown(_rows_to_md_table(detail_rows))

            # 미스캔 항목
            incomplete = []
            for bc, v in sort_state.items():
                for it in v['items']:
                    if it['scanned'] < it['needed']:
                        incomplete.append({
                            '박스': f"{it['box_num']}번",
                            '바코드': bc,
                            '상품명': v['상품명'][:35],
                            '필요': it['needed'],
                            '스캔': it['scanned'],
                            '남음': it['needed'] - it['scanned'],
                        })
            if incomplete:
                with st.expander(f'⏳ 미스캔 ({len(incomplete)}건)', expanded=False):
                    st.dataframe(_pd2.DataFrame(incomplete), use_container_width=True, hide_index=True)

        _scan_fragment()



# ══════════════════════════════════════════════════════
# ── 쿠팡 발주서 변경 요청 (입고예정일 & 납품센터) ─────
#   Supplier Hub → 물류 → 상품 공급상태 관리 → + 요청 등록
#   → '발주서 변경 > 입고예정일 & FC(물류센터) 동시 변경' → [대량 업로드]
#   에 올릴 엑셀을 만든다.
# ══════════════════════════════════════════════════════
with tab_pochg:
    import pandas as _pd_chg      # 이 앱은 pandas 를 쓰는 자리에서 그때그때 불러온다

    st.header('🚚 쿠팡 발주서 변경 요청서 만들기')
    st.caption(f'🔖 기능 버전 **{coupang_po_change.VERSION}** — '
               '입고예정일 · 납품센터(FC) 동시 변경 대량 업로드용')

    _cpc = coupang_po_change

    st.info(
        '쿠팡이 2026-09-19 이후 발행 발주서부터 Supplier Hub에서 셀프 변경을 열어줍니다. '
        '**[엑셀 다운로드]로 받은 실제 양식이 생기면 아래 4단계에 올려주세요.** '
        '그 양식에 그대로 값을 채워 드립니다. 양식이 없으면 가이드 화면과 같은 머리글로 새로 만듭니다.'
    )
    with st.expander('📌 쿠팡이 정한 요청 조건 (눌러서 보기)'):
        for _note in _cpc.RULE_NOTES:
            st.markdown(f'- {_note}')
        st.markdown('---')
        st.markdown(
            'D-7이 지났거나 발주서에 락이 걸려 날짜가 선택되지 않으면 인스탁 매니저에게 '
            '메일로 요청해야 합니다. 대량 컨테이너를 같은 날짜로 넣어야 할 때도 메일입니다.'
        )

    # ── 1단계 : 발주서 목록 엑셀 ───────────────────────
    st.divider()
    st.subheader('1️⃣ 발주서 목록 엑셀 올리기')
    st.caption('쿠팡에서 받은 발주서 목록을 올리면 발주번호 · 납품센터 · 입고예정일 열을 알아서 찾습니다.')
    po_file = st.file_uploader('발주서 목록 (.xlsx)', type=['xlsx'], key='pochg_src')

    _rows, _info = [], {}
    if po_file:
        try:
            _rows, _info = _cpc.read_po_list(po_file.getvalue())
        except Exception as e:
            st.error(f'엑셀을 읽지 못했습니다: {e}')
        if not _rows:
            st.warning('발주번호를 찾지 못했습니다. 발주번호가 들어 있는 엑셀인지 확인해주세요. '
                       f'(읽은 시트: {_info.get("sheet", "?")})')
        else:
            _found = _info.get('columns', {})
            _label = {'po': '발주번호', 'center': '납품센터', 'eta': '입고예정일'}
            _desc = ' · '.join(
                f'{_label[k]}={get_column_letter(v)}열' for k, v in _found.items() if k in _label)
            st.success(f'발주서 {len(_rows)}건을 읽었습니다. (시트 "{_info.get("sheet")}" · {_desc})')
            for _need, _msg in [('center', '납품센터'), ('eta', '입고예정일')]:
                if _need not in _found:
                    st.warning(f'{_msg} 열을 못 찾았습니다 — 필수 항목이라 아래 표에서 직접 채워주세요.')

    # ── 2단계 : 변경 내용 ─────────────────────────────
    st.divider()
    st.subheader('2️⃣ 변경할 내용 (전체 공통)')
    _c1, _c2, _c3 = st.columns(3)
    with _c1:
        _new_center = st.text_input('변경 납품센터', value='', key='pochg_center',
                                    help='센터만 바꾸면 날짜는 비워 두세요. 예: 인천4, 덕평1')
    with _c2:
        _use_date = st.checkbox('입고예정일도 변경', value=False, key='pochg_usedate')
        _new_eta = st.date_input('변경 입고예정일', key='pochg_date') if _use_date else None
    with _c3:
        _reason = st.text_input('요청사유', value='생산지연', key='pochg_reason',
                                help='Supplier Hub 화면의 요청사유 목록과 똑같이 적어야 반려되지 않습니다. '
                                     '예: ' + ', '.join(_cpc.REASON_PRESETS))

    # ── 3단계 : 확인하고 고치기 ────────────────────────
    _edited = None
    if _rows:
        for _r in _rows:
            _r.new_center = _new_center.strip()
            _r.new_eta = _new_eta if _use_date else None
            _r.reason = _reason.strip()
        _cpc.validate(_rows)

        st.divider()
        st.subheader('3️⃣ 확인하고 고치기')
        st.caption('행마다 다르게 넣어야 하면 표에서 바로 고치세요. '
                   '위 공통값을 바꾸면 표가 다시 채워지니, 공통값을 먼저 정하고 표를 고치는 순서로 하세요.')
        _df = _pd_chg.DataFrame([{
            '발주 번호': _r.po,
            '기존 납품센터': _r.old_center,
            '변경 납품센터': _r.new_center,
            '기존 입고예정일': _cpc.fmt_date(_r.old_eta),
            '변경 입고예정일': _cpc.fmt_date(_r.new_eta),
            '요청사유': _r.reason,
            '확인 필요': ' / '.join(_r.issues),
        } for _r in _rows])
        _edited = st.data_editor(
            _df, use_container_width=True, hide_index=True, num_rows='dynamic',
            key='pochg_editor',
            disabled=['확인 필요'],
            column_config={'확인 필요': st.column_config.TextColumn(
                '확인 필요', help='쿠팡 조건에 걸릴 수 있는 항목입니다. 표를 고치면 아래 버튼을 눌러 다시 검증하세요.',
                width='large')},
        )
        _bad = sum(1 for _r in _rows if _r.issues)
        if _bad:
            st.warning(f'{_bad}건이 쿠팡 조건에 걸릴 수 있습니다. "확인 필요" 열을 봐주세요. '
                       '그래도 엑셀은 만들 수 있습니다 — 최종 판단은 쿠팡 화면에서 합니다.')
        else:
            st.success('조건에 걸리는 건이 없습니다.')

    # ── 4단계 : 엑셀 만들기 ───────────────────────────
    st.divider()
    st.subheader('4️⃣ 업로드용 엑셀 만들기')
    tpl_file = st.file_uploader(
        'Supplier Hub에서 [엑셀 다운로드]로 받은 양식 (.xlsx, 선택)', type=['xlsx'], key='pochg_tpl',
        help='올리면 그 양식에 값만 채워 드립니다. 서식과 드롭다운이 그대로 남아 가장 안전합니다. '
             '안 올리면 가이드 화면과 같은 머리글로 새로 만듭니다.')

    if st.button('📄 변경 요청 엑셀 만들기', type='primary', key='pochg_btn'):
        if _edited is None or len(_edited) == 0:
            st.warning('⚠️ 발주서 목록을 먼저 올려주세요.')
        else:
            # 표에서 고친 값을 그대로 반영해 다시 만든다
            _final = []
            for _, _row in _edited.iterrows():
                _po = str(_row.get('발주 번호', '') or '').strip()
                if not _po:
                    continue
                _final.append(_cpc.ChangeRow(
                    po=_po,
                    old_center=str(_row.get('기존 납품센터', '') or '').strip(),
                    new_center=str(_row.get('변경 납품센터', '') or '').strip(),
                    old_eta=_cpc.parse_date(_row.get('기존 입고예정일')),
                    new_eta=_cpc.parse_date(_row.get('변경 입고예정일')),
                    reason=str(_row.get('요청사유', '') or '').strip(),
                ))
            _cpc.validate(_final)
            try:
                _buf = _cpc.build_excel(_final, tpl_file.getvalue() if tpl_file else None)
            except Exception as e:
                st.error(f'엑셀 생성 실패: {e}')
            else:
                _issues = [f'{r.po}: {" / ".join(r.issues)}' for r in _final if r.issues]
                _made_from = '올려주신 양식' if tpl_file else '기본 양식'
                st.success(f'🎉 {len(_final)}건 · {_made_from}으로 만들었습니다.')
                if _issues:
                    st.warning(f'조건에 걸릴 수 있는 {len(_issues)}건:')
                    for _t in _issues[:20]:
                        st.caption(f'· {_t}')
                    if len(_issues) > 20:
                        st.caption(f'· … 외 {len(_issues) - 20}건')
                _today = datetime.now().strftime('%Y%m%d')
                st.download_button(
                    '⬇️ 엑셀 다운로드', _buf,
                    file_name=f'발주서변경요청_{_today}_{kit_config.company_name() or "업체"}.xlsx',
                    mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    key='pochg_dl')
                st.caption('Supplier Hub → 물류 → 상품 공급상태 관리 → + 요청 등록 → '
                           '발주서 변경 > 입고예정일 & FC 동시 변경 → [대량 가져오기]에 올리세요.')
