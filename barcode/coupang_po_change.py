# ╔══════════════════════════════════════════════════════╗
# ║  쿠팡 발주서 변경 요청 엑셀 만들기               ║
# ║  (입고예정일 & FC(납품센터) 동시 변경 · 대량 업로드용)  ║
# ╚══════════════════════════════════════════════════════╝
#
# 무엇을 하는 파일인가?
#   Supplier Hub → 물류 → 상품 공급상태 관리 → + 요청 등록 →
#   '발주서 변경 > 입고예정일 & FC(물류센터) 동시 변경' 에서
#   [대량 업로드] 로 올릴 엑셀을 만들어 준다.
#
#   한 건씩 캘린더로 찍으면 발주서 수십 개를 일일이 눌러야 하는데,
#   쿠팡이 내려준 발주서 목록 엑셀만 있으면 여기서 한 번에 만들어 낸다.
#
# 양식 (SRMS 가이드 Ver9.1 p.27 화면 그대로)
#   발주 번호* | 기존 납품센터* | 변경 납품센터 |
#   기존 입고예정일(YYYY-MM-DD)* | 변경 입고예정일(YYYY-MM-DD) | 요청사유*
#   (* 는 필수. '삭제' 열은 화면 버튼이라 엑셀에는 넣지 않는다)
#
# 구조
#   1. 양식 정의 · 규칙
#   2. 날짜 다루기
#   3. 발주서 목록 엑셀 읽기 (열 자동 찾기)
#   4. 규칙 검증
#   5. 업로드용 엑셀 만들기

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 이 기능을 고칠 때마다 함께 올린다 — 앱 화면에 찍어서 배포 반영 여부를 확인한다.
VERSION = '2026-09-17a · 발주 변경 요청서'

# ══════════════════════════════════════════════════════
# 1. 양식 정의 · 규칙
# ══════════════════════════════════════════════════════
HEADERS = [
    '발주 번호',
    '기존 납품센터',
    '변경 납품센터',
    '기존 입고예정일(YYYY-MM-DD)',
    '변경 입고예정일(YYYY-MM-DD)',
    '요청사유',
]

# 쿠팡 안내 메일 + SRMS 가이드 p.27 에 적힌 조건들.
# 숫자를 코드 곳곳에 흩어놓지 않고 여기 모아 둔다 (기준이 바뀌면 여기만 고친다).
RULE_LEAD_DAYS = 7      # 기존 입고예정일 기준 D-7 이전까지만 셀프 변경 가능
RULE_MAX_DELAY = 14     # 변경일은 최초 입고예정일로부터 14일 이내 (앞당기기는 제한 없음)

RULE_NOTES = [
    '입고예정일 기준 **D-7 이전**까지만 이 기능으로 변경할 수 있습니다. '
    '그 이후에는 인스탁 매니저에게 메일로 요청해야 합니다.',
    '변경 입고예정일은 **최초 입고예정일로부터 14일 이내**만 가능합니다. '
    '일정을 앞당기는 것은 제한이 없습니다.',
    '입고예정일 당일은 선택할 수 없습니다.',
    '로켓프레시 · 벤더플렉스(VF) · Private Label(PL) 발주서는 요청이 불가능합니다.',
    '로켓설치 SKU는 이 기능으로 납품센터(FC)를 바꿀 수 없습니다.',
    '요청사유가 실제 발주서 상태와 맞지 않으면 시스템에서 자동 반려됩니다.',
    'OOS 위험 SKU가 포함되면 담당자 승인이 필요해 하루 두 번 처리됩니다.',
    '선택이 안 되는 날짜는 물류센터 수용량이 꽉 찬 것이라 메일로도 불가합니다.',
]

# 요청사유 예시 — Supplier Hub 화면에 있는 값과 똑같이 적어야 반려되지 않는다.
REASON_PRESETS = ['생산지연', '수입지연', '출고지연', '물류사 사정', '재고부족', '기타']


# ══════════════════════════════════════════════════════
# 2. 날짜 다루기
# ══════════════════════════════════════════════════════
def parse_date(value):
    """엑셀에서 나온 온갖 날짜 표기를 date 로 바꾼다. 못 읽으면 None.

    받아주는 형태: datetime/date, '2026-09-21', '2026.09.21', '2026/09/21',
                  '20260921', 20260921(숫자)
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s:
        return None
    # 엑셀에서 숫자로 읽힌 20260921 같은 값
    if isinstance(value, (int, float)) and float(value).is_integer():
        s = str(int(value))
    s = s.split(' ')[0]                       # '2026-09-21 00:00:00' 대응
    digits = re.sub(r'\D', '', s)
    if len(digits) == 8:
        try:
            return date(int(digits[:4]), int(digits[4:6]), int(digits[6:]))
        except ValueError:
            return None
    for fmt in ('%Y-%m-%d', '%Y.%m.%d', '%Y/%m/%d', '%y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def fmt_date(d):
    """양식이 요구하는 YYYY-MM-DD 문자열로."""
    return d.strftime('%Y-%m-%d') if isinstance(d, date) else ''


# ══════════════════════════════════════════════════════
# 3. 발주서 목록 엑셀 읽기
# ══════════════════════════════════════════════════════
# 쿠팡에서 내려받는 목록은 열 순서가 화면마다 조금씩 다르다.
# 그래서 열 번호를 외우게 하지 않고, 머리글 이름을 보고 알아서 찾는다.
COLUMN_HINTS = {
    'po': ['발주서번호', '발주 번호', '발주번호', 'po번호', 'po 번호', 'purchase order'],
    'center': ['납품센터', '입고센터', '물류센터', 'fc', '센터'],
    'eta': ['입고예정일', '입고 예정일', '예정일', 'eta'],
}


def _norm(text):
    """머리글 비교용 — 공백·괄호·별표를 없애고 소문자로."""
    return re.sub(r'[\s()\[\]*·:]+', '', str(text or '')).lower()


def find_header_row(ws, max_scan=15):
    """머리글이 있는 행과 {역할: 열번호} 를 찾는다.

    위쪽에 안내문이 몇 줄 붙어 있는 파일이 많아서, 앞 15줄을 훑어
    '발주번호 비슷한 것'이 들어 있는 행을 머리글로 본다.
    """
    best = (None, {})
    for r in range(1, min(ws.max_row, max_scan) + 1):
        found = {}
        for c in range(1, min(ws.max_column, 60) + 1):
            cell = _norm(ws.cell(r, c).value)
            if not cell:
                continue
            for role, hints in COLUMN_HINTS.items():
                if role in found:
                    continue
                if any(_norm(h) in cell for h in hints):
                    found[role] = c
        if 'po' in found and len(found) > len(best[1]):
            best = (r, found)
    return best


@dataclass
class ChangeRow:
    po: str                                   # 발주서번호
    old_center: str = ''                      # 기존 납품센터
    new_center: str = ''                      # 변경 납품센터
    old_eta: Optional[date] = None            # 기존 입고예정일
    new_eta: Optional[date] = None            # 변경 입고예정일
    reason: str = ''                          # 요청사유
    issues: list = field(default_factory=list)   # 검증에서 걸린 내용


def read_po_list(xlsx_bytes, sheet_name=None, col_map=None, header_row=None):
    """발주서 목록 엑셀 → (ChangeRow 목록, 찾아낸 정보 dict).

    col_map / header_row 를 주면 그대로 쓰고, 안 주면 알아서 찾는다.
    """
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    if header_row is None or not col_map:
        header_row, col_map = find_header_row(ws)
    info = {'sheet': ws.title, 'header_row': header_row, 'columns': dict(col_map or {}),
            'sheets': wb.sheetnames}
    if not header_row or not col_map or 'po' not in col_map:
        return [], info

    rows, seen = [], set()
    for r in range(header_row + 1, ws.max_row + 1):
        po_raw = ws.cell(r, col_map['po']).value
        if po_raw is None or not str(po_raw).strip():
            continue
        po = str(po_raw).strip()
        if isinstance(po_raw, float) and po_raw.is_integer():
            po = str(int(po_raw))
        # 같은 발주서가 SKU 단위로 여러 줄 있을 수 있다 — 변경 요청은 발주서 단위다
        if po in seen:
            continue
        seen.add(po)
        rows.append(ChangeRow(
            po=po,
            old_center=str(ws.cell(r, col_map['center']).value or '').strip()
            if 'center' in col_map else '',
            old_eta=parse_date(ws.cell(r, col_map['eta']).value) if 'eta' in col_map else None,
        ))
    info['duplicates_skipped'] = True
    return rows, info


# ══════════════════════════════════════════════════════
# 4. 규칙 검증
# ══════════════════════════════════════════════════════
def validate(rows, today=None):
    """쿠팡 조건에 걸리는 행을 찾아 issues 에 적어 둔다.

    막지는 않는다 — 실제 가능 여부는 쿠팡 화면이 최종 판단하므로,
    여기서는 '이건 반려될 가능성이 높다'고 알려주기만 한다.
    """
    today = today or date.today()
    for row in rows:
        row.issues = []
        if not row.po:
            row.issues.append('발주번호 없음')
        if not row.old_center:
            row.issues.append('기존 납품센터 비어 있음 (필수)')
        if row.old_eta is None:
            row.issues.append('기존 입고예정일 비어 있음 (필수)')
        if not row.reason:
            row.issues.append('요청사유 비어 있음 (필수)')
        if not row.new_center and row.new_eta is None:
            row.issues.append('변경할 내용이 없음 (센터·날짜 중 하나는 입력)')

        if row.old_eta:
            deadline = row.old_eta - timedelta(days=RULE_LEAD_DAYS)
            if today > deadline:
                row.issues.append(
                    f'D-7 지남 — {fmt_date(deadline)}까지만 셀프 변경 가능, 메일 요청 필요')
        if row.new_eta:
            if row.new_eta <= today:
                row.issues.append('변경일이 오늘이거나 지난 날짜')
            if row.old_eta:
                if row.new_eta == row.old_eta:
                    row.issues.append('변경일이 기존 입고예정일과 같음')
                # 뒤로 미루는 경우만 14일 제한 (앞당기기는 제한 없음)
                delay = (row.new_eta - row.old_eta).days
                if delay > RULE_MAX_DELAY:
                    row.issues.append(
                        f'{delay}일 연기 — 최초 입고예정일로부터 {RULE_MAX_DELAY}일 이내만 가능')
        if row.new_center and row.old_center and row.new_center == row.old_center:
            row.issues.append('변경 납품센터가 기존과 같음')
    return rows


# ══════════════════════════════════════════════════════
# 5. 업로드용 엑셀 만들기
# ══════════════════════════════════════════════════════
def _row_values(row):
    return [row.po, row.old_center, row.new_center,
            fmt_date(row.old_eta), fmt_date(row.new_eta), row.reason]


def build_excel(rows, template_bytes=None):
    """업로드용 엑셀 (BytesIO).

    template_bytes 를 주면 Supplier Hub 에서 받은 그 양식에 값만 채운다.
    서식·드롭다운·숨은 시트가 그대로 남아 가장 안전하다.
    양식이 없으면 가이드 화면과 같은 머리글로 새로 만든다.
    """
    if template_bytes:
        out = _fill_template(rows, template_bytes)
        if out is not None:
            return out
    return _build_fresh(rows)


def _fill_template(rows, template_bytes):
    """받은 양식의 머리글을 찾아 그 아래에 값을 채운다. 못 찾으면 None."""
    try:
        wb = load_workbook(io.BytesIO(template_bytes))
    except Exception:
        return None
    ws = wb.active
    header_row, col_map = find_header_row(ws, max_scan=20)
    if not header_row or 'po' not in (col_map or {}):
        return None

    # 양식의 열 순서를 그대로 따르려고, 머리글 글자를 보고 6개 열을 다시 짚는다.
    targets = {}
    for c in range(1, min(ws.max_column, 60) + 1):
        h = _norm(ws.cell(header_row, c).value)
        if not h:
            continue
        if '발주' in h and 'po' not in targets:
            targets['po'] = c
        elif '기존' in h and '센터' in h:
            targets['old_center'] = c
        elif '변경' in h and '센터' in h:
            targets['new_center'] = c
        elif '기존' in h and '입고예정일' in h:
            targets['old_eta'] = c
        elif '변경' in h and '입고예정일' in h:
            targets['new_eta'] = c
        elif '사유' in h:
            targets['reason'] = c
    if 'po' not in targets:
        return None

    r = header_row + 1
    for row in rows:
        vals = {'po': row.po, 'old_center': row.old_center, 'new_center': row.new_center,
                'old_eta': fmt_date(row.old_eta), 'new_eta': fmt_date(row.new_eta),
                'reason': row.reason}
        for key, col in targets.items():
            ws.cell(r, col).value = vals.get(key, '')
        r += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_fresh(rows):
    """양식 파일이 없을 때 — 가이드 화면과 같은 머리글로 새 엑셀을 만든다."""
    wb = Workbook()
    ws = wb.active
    ws.title = '발주서 변경 요청'

    head_fill = PatternFill('solid', fgColor='E8F0FE')
    head_font = Font(bold=True)
    for c, name in enumerate(HEADERS, start=1):
        cell = ws.cell(1, c, name)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(_row_values(row), start=1):
            cell = ws.cell(r, c, v)
            # 날짜는 반드시 글자로 — 엑셀이 날짜 서식으로 바꿔 'YYYY-MM-DD' 가 깨지는 것 방지
            if c in (1, 4, 5):
                cell.number_format = '@'

    widths = [16, 16, 16, 24, 24, 20]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
